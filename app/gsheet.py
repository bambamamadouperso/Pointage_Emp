"""Source Google Sheets : téléchargement du classeur, lecture des onglets et typage des colonnes.

Deux modes d'accès :
- « public » : classeur partagé « Tous les utilisateurs disposant du lien » (aucune clé nécessaire) ;
- « service_account » : classeur partagé avec l'e-mail d'un compte de service Google
  (clé JSON saisie dans la connexion, API Google Drive activée dans le projet Google Cloud).

Le classeur est exporté au format Excel (.xlsx) puis lu avec openpyxl : on obtient ainsi tous les
onglets avec des valeurs typées (nombres, dates, booléens), indépendamment de la langue du classeur.
"""
import base64
import io
import json
import re
import time as _time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from sqlalchemy import (
    BigInteger,
    Boolean,
    Column,
    Date,
    DateTime,
    Float,
    Integer,
    MetaData,
    Numeric,
    Table,
    Text,
    Time,
)
from sqlalchemy.types import TypeEngine

from .crypto import decrypt

AUTH_PUBLIC = "public"
AUTH_SERVICE_ACCOUNT = "service_account"
AUTH_LABELS = {
    AUTH_PUBLIC: "Lien public (« Tous les utilisateurs disposant du lien »)",
    AUTH_SERVICE_ACCOUNT: "Compte de service Google (classeur privé)",
}

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
TIMEOUT = 60


class SheetError(Exception):
    pass


# --------------------------------------------------------------------------- identifiants


def parse_spreadsheet_id(value: str) -> str:
    """Accepte l'URL complète du classeur ou son identifiant."""
    value = (value or "").strip()
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", value) or re.search(r"/d/([a-zA-Z0-9_-]+)", value)
    if match:
        return match.group(1)
    match = re.search(r"[?&]id=([a-zA-Z0-9_-]+)", value)
    if match:
        return match.group(1)
    return value


def normalize_identifier(name: Any, index: int = 0) -> str:
    """Nom d'onglet ou d'en-tête -> identifiant PostgreSQL simple (minuscules, sans accents)."""
    text = unicodedata.normalize("NFKD", str(name if name is not None else "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^0-9a-zA-Z]+", "_", text).strip("_").lower()[:60]
    if not text:
        text = f"col_{index + 1}"
    if text[0].isdigit():
        text = f"c_{text}"
    return text


def service_account_email(conn) -> str:
    try:
        return json.loads(decrypt(conn.password_enc)).get("client_email", "")
    except (ValueError, TypeError):
        return ""


# --------------------------------------------------------------------------- téléchargement


def _http_get(url: str, headers: Optional[dict] = None) -> tuple[int, str, bytes]:
    request = urllib.request.Request(url, headers={"User-Agent": "mariadb-pg-sync", **(headers or {})})
    try:
        with urllib.request.urlopen(request, timeout=TIMEOUT) as response:
            return response.status, response.headers.get("Content-Type", ""), response.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.headers.get("Content-Type", "") if exc.headers else "", exc.read()


def _b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b"=").decode()


def access_token(info: dict) -> str:
    """Jeton OAuth d'un compte de service (JWT signé RS256, sans bibliothèque Google)."""
    for field in ("client_email", "private_key"):
        if not info.get(field):
            raise SheetError(f"Clé JSON invalide : champ « {field} » manquant.")
    token_uri = info.get("token_uri") or "https://oauth2.googleapis.com/token"
    now = int(_time.time())
    header = _b64(json.dumps({"alg": "RS256", "typ": "JWT"}).encode())
    claims = _b64(json.dumps({
        "iss": info["client_email"], "scope": DRIVE_SCOPE, "aud": token_uri, "iat": now, "exp": now + 3600,
    }).encode())
    signing_input = f"{header}.{claims}".encode()
    try:
        key = serialization.load_pem_private_key(info["private_key"].encode(), password=None)
    except (ValueError, TypeError) as exc:
        raise SheetError("Clé privée du compte de service illisible.") from exc
    signature = key.sign(signing_input, padding.PKCS1v15(), hashes.SHA256())
    assertion = f"{header}.{claims}.{_b64(signature)}"
    body = urllib.parse.urlencode({
        "grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer", "assertion": assertion,
    }).encode()
    request = urllib.request.Request(token_uri, data=body,
                                     headers={"Content-Type": "application/x-www-form-urlencoded"})
    try:
        with urllib.request.urlopen(request, timeout=TIMEOUT) as response:
            return json.loads(response.read())["access_token"]
    except urllib.error.HTTPError as exc:
        raise SheetError(f"Authentification Google refusée : {exc.read()[:300].decode(errors='replace')}") from exc


def _google_error(payload: bytes) -> str:
    try:
        error = json.loads(payload).get("error", {})
        return error.get("message", "") if isinstance(error, dict) else str(error)
    except ValueError:
        return ""


def download_workbook(conn) -> bytes:
    """Télécharge le classeur au format .xlsx."""
    sheet_id = parse_spreadsheet_id(conn.database)
    if not sheet_id:
        raise SheetError("Identifiant du classeur manquant.")

    if conn.username == AUTH_SERVICE_ACCOUNT:
        try:
            info = json.loads(decrypt(conn.password_enc))
        except ValueError as exc:
            raise SheetError("Clé JSON du compte de service invalide.") from exc
        headers = {"Authorization": f"Bearer {access_token(info)}"}
        base = f"https://www.googleapis.com/drive/v3/files/{sheet_id}"
        status, _, payload = _http_get(f"{base}/export?mimeType={urllib.parse.quote(XLSX_MIME)}", headers)
        if status in (400, 403) and "export" in _google_error(payload).lower():
            # Fichier Excel déposé sur Drive (pas un Google Sheets natif) : téléchargement direct.
            status, _, payload = _http_get(f"{base}?alt=media&supportsAllDrives=true", headers)
        if status == 404:
            raise SheetError(
                f"Classeur introuvable : partagez-le (lecteur) avec {info.get('client_email')}."
            )
        message = _google_error(payload).lower()
        if status == 403 and ("has not been used" in message or "disabled" in message):
            raise SheetError("Activez l'API « Google Drive » dans le projet Google Cloud du compte de service.")
        if status != 200:
            raise SheetError(f"Erreur Google ({status}) : {_google_error(payload) or payload[:200]!r}")
    else:
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
        status, _, payload = _http_get(url)
        if status == 404:
            raise SheetError("Classeur introuvable : vérifiez le lien.")
        if status != 200 or not payload.startswith(b"PK"):
            raise SheetError(
                "Classeur inaccessible : partagez-le avec « Tous les utilisateurs disposant du lien » "
                "(lecteur), ou utilisez un compte de service."
            )
    if not payload.startswith(b"PK"):
        raise SheetError("Le fichier reçu n'est pas un classeur valide.")
    return payload


# --------------------------------------------------------------------------- lecture


@dataclass
class SheetData:
    title: str
    columns: list[str]        # noms normalisés
    headers: list[str]        # en-têtes d'origine
    rows: list[list[Any]]     # valeurs brutes (une liste par ligne, alignée sur columns)


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.replace("\x00", "").strip()
        return value or None
    return value


def read_workbook(payload: bytes) -> dict[str, SheetData]:
    """Lit tous les onglets : la première ligne non vide contient les en-têtes."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheets: dict[str, SheetData] = {}
    try:
        for ws in workbook.worksheets:
            raw_rows = [[_clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
            raw_rows = [r for r in raw_rows if any(v is not None for v in r)]
            if not raw_rows:
                sheets[ws.title] = SheetData(ws.title, [], [], [])
                continue
            header, body = raw_rows[0], raw_rows[1:]
            width = max(len(r) for r in raw_rows)
            header = list(header) + [None] * (width - len(header))
            # Colonnes conservées : en-tête renseigné, ou au moins une valeur.
            keep = [i for i in range(width)
                    if header[i] is not None or any(i < len(r) and r[i] is not None for r in body)]
            seen: set[str] = set()
            columns, headers = [], []
            for i in keep:
                name = base = normalize_identifier(header[i], i)
                n = 2
                while name in seen:
                    name = f"{base}_{n}"
                    n += 1
                seen.add(name)
                columns.append(name)
                headers.append("" if header[i] is None else str(header[i]))
            rows = [[r[i] if i < len(r) else None for i in keep] for r in body]
            sheets[ws.title] = SheetData(ws.title, columns, headers, rows)
    finally:
        workbook.close()
    return sheets


def load_sheets(conn) -> dict[str, SheetData]:
    return read_workbook(download_workbook(conn))


# --------------------------------------------------------------------------- types


def infer_type(values: list[Any]) -> TypeEngine:
    present = [v for v in values if v is not None]
    if not present:
        return Text()
    if all(isinstance(v, bool) for v in present):
        return Boolean()
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        if all(isinstance(v, int) or float(v).is_integer() for v in present) and \
                all(abs(v) < 2 ** 63 for v in present):
            return BigInteger()
        return Numeric()
    if all(isinstance(v, (datetime, date)) for v in present):
        if all(not isinstance(v, datetime) or v.time() == time(0) for v in present):
            return Date()
        return DateTime()
    if all(isinstance(v, time) for v in present):
        return Time()
    return Text()


def build_table(name: str, data: SheetData) -> Table:
    """Table SQLAlchemy (en mémoire) décrivant l'onglet, avec des types déduits des valeurs."""
    columns = [
        Column(col, infer_type([row[i] for row in data.rows]))
        for i, col in enumerate(data.columns)
    ]
    return Table(name, MetaData(), *columns)


_NUMBER_SPACES = re.compile(r"[\s  ']")
_DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y")
_TIME_FORMATS = ("", " %H:%M", " %H:%M:%S")


def _parse_number(text: str) -> Decimal:
    s = _NUMBER_SPACES.sub("", text).replace("%", "")
    if "," in s and "." in s:
        # Le dernier séparateur est le séparateur décimal.
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".") if s.count(",") == 1 else s.replace(",", "")
    return Decimal(s)


def _parse_datetime(text: str) -> datetime:
    text = text.strip().replace("T", " ")
    for fmt in _DATE_FORMATS:
        for tfmt in _TIME_FORMATS:
            try:
                return datetime.strptime(text, fmt + tfmt)
            except ValueError:
                continue
    raise ValueError(text)


def _as_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") if value.time() != time(0) else value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


_TRUE = {"true", "vrai", "oui", "yes", "1", "x", "o", "y"}
_FALSE = {"false", "faux", "non", "no", "0", ""}


def convert(value: Any, target: TypeEngine) -> Any:
    """Convertit une valeur de cellule vers le type de la colonne cible. Lève ValueError si impossible."""
    if value is None:
        return None
    if isinstance(target, Boolean):
        if isinstance(value, bool):
            return value
        text = _as_text(value).strip().lower()
        if text in _TRUE:
            return True
        if text in _FALSE:
            return False
        raise ValueError(value)
    if isinstance(target, (Integer, BigInteger)):
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float)):
            if float(value).is_integer():
                return int(value)
            raise ValueError(value)
        number = _parse_number(str(value))
        if number != number.to_integral_value():
            raise ValueError(value)
        return int(number)
    if isinstance(target, (Numeric, Float)):
        if isinstance(value, bool):
            return Decimal(int(value))
        if isinstance(value, (int, float)):
            return Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
        try:
            return _parse_number(str(value))
        except InvalidOperation as exc:
            raise ValueError(value) from exc
    if isinstance(target, DateTime):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time(0))
        return _parse_datetime(str(value))
    if isinstance(target, Date):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return _parse_datetime(str(value)).date()
    if isinstance(target, Time):
        if isinstance(value, time):
            return value
        if isinstance(value, datetime):
            return value.time()
        if isinstance(value, timedelta):
            return (datetime.min + value).time()
        return datetime.strptime(str(value).strip(), "%H:%M:%S" if str(value).count(":") == 2 else "%H:%M").time()
    return _as_text(value)


def describe(conn) -> str:
    """Test de connexion : renvoie un résumé du classeur."""
    sheets = load_sheets(conn)
    names = ", ".join(list(sheets)[:8]) + ("…" if len(sheets) > 8 else "")
    return f"classeur accessible : {len(sheets)} onglet(s) ({names})"
