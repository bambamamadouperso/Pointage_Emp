"""Espace d'administration (rôle admin) : paramètres horaires historisés, source des pointages, jours fériés,
utilisateurs et rôles, journal d'audit, traitements planifiés."""
import re
import secrets
import threading
import time
from datetime import date, datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import JSONResponse, Response
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import auth, pointage, scheduler
from ..database import get_db
from ..errors import friendly
from ..models import ROLES, AuditEntry, Connection, PointageConfig, SyncJob, User, utcnow
from ..sync import is_running, make_engine
from ..web import flash, redirect, render, require_login

router = APIRouter(prefix="/admin", dependencies=[Depends(require_login)])


def _config(db: Session) -> PointageConfig:
    cfg = db.query(PointageConfig).order_by(PointageConfig.id).first()
    if cfg is None:
        cfg = PointageConfig(data="{}")
        db.add(cfg)
        db.commit()
    return cfg


def _mapping_ready(db: Session):
    from .suivi import load_config

    cfg = _config(db)
    _, mapping = load_config(db)  # réinstalle les calculs s'ils datent d'une version précédente
    return cfg, mapping


# --------------------------------------------------------------------------- accueil


@router.get("")
def index(request: Request, db: Session = Depends(get_db)):
    cfg = _config(db)
    jobs = db.scalars(select(SyncJob).order_by(SyncJob.name)).all()
    return render(
        request, "admin/index.html", cfg=cfg, mapping=pointage.Mapping.from_json(cfg.data), jobs=jobs,
        running={j.id: is_running(j.id) for j in jobs}, next_runs={j.id: scheduler.next_run_time(j.id) for j in jobs},
        users=db.scalar(select(func.count(User.id))), audits=db.scalar(select(func.count(AuditEntry.id))),
    )


# --------------------------------------------------------------------------- source des pointages

_MAPPING_FIELDS = [f for f in pointage.Mapping.__dataclass_fields__]
# Champs à choix multiples (plusieurs colonnes « Date aller » / « Date retour ») : « a, b ».
_MULTI_FIELDS = ("mission_start_col", "mission_end_col")


def _field(values, key: str) -> str:
    if key in _MULTI_FIELDS:
        return ", ".join(pointage.multi([str(v) for v in values.getlist(key)]))
    return str(values.get(key, "") or "").strip()


@router.get("/pointage")
def pointage_config(request: Request, db: Session = Depends(get_db)):
    cfg = _config(db)
    params = request.query_params
    exploring = "conn_id" in params
    mapping = pointage.Mapping.from_json(cfg.data)
    conn_id = cfg.conn_id
    if exploring:  # rechargement du formulaire (choix d'une connexion, d'un schéma ou d'une table)
        mapping = pointage.Mapping(**{k: _field(params, k) for k in _MAPPING_FIELDS})
        conn_id = int(params["conn_id"]) if params["conn_id"].isdigit() else None
    connections = db.scalars(select(Connection).where(Connection.kind == "postgresql").order_by(Connection.name)).all()
    conn = db.get(Connection, conn_id) if conn_id else (connections[0] if connections else None)
    schemas, tables, punch_cols, emp_cols, service_cols, error, installed = [], [], {}, {}, {}, None, False
    person_cols, hier_cols, leave_cols, tw_cols, diag, stale = {}, {}, {}, {}, None, False
    cat_cols, cat_values, mission_cols, mission_match = {}, [], {}, None
    dims = [{"name": name, "prefix": prefix, "label": label, "cols": {}, "values": [],
             "hint": "ex. Usine, Plateau" if name == "site" else "colonne « direction » du personnel"}
            for name, (prefix, label) in pointage.DIMENSIONS.items()]
    if conn is not None:
        engine = make_engine(conn, **pointage.WEB_LIMITS)
        try:
            schemas = pointage.list_schemas(engine)
            if mapping.schema not in schemas and schemas:
                mapping.schema = "public" if "public" in schemas else schemas[0]
            tables = pointage.list_tables(engine, mapping.schema)
            names = [t for t in tables]
            mapping.punch_table = mapping.punch_table or pointage.guess("punch_table", names)
            mapping.emp_table = mapping.emp_table or pointage.guess("emp_table", names, (mapping.punch_table,))
            punch_cols = pointage.column_types(engine, mapping.schema, mapping.punch_table)
            emp_cols = pointage.column_types(engine, mapping.schema, mapping.emp_table)
            service_cols = pointage.column_types(engine, mapping.schema, mapping.service_table)
            person_cols = pointage.column_types(engine, mapping.schema, mapping.person_table)
            hier_cols = pointage.column_types(engine, mapping.schema, mapping.hier_table)
            leave_cols = pointage.column_types(engine, mapping.schema, mapping.leave_table)
            tw_cols = pointage.column_types(engine, mapping.schema, mapping.tw_table)
            mission_cols = pointage.column_types(engine, mapping.schema, mapping.mission_table)
            cat_cols = pointage.column_types(engine, mapping.schema, mapping.cat_table)
            for d in dims:  # direction, site : colonnes de la table des libellés proposées
                d["cols"] = pointage.column_types(engine, mapping.schema, getattr(mapping, f"{d['prefix']}_table"))
                for key in (f"{d['prefix']}_key_col", f"{d['prefix']}_label_col"):
                    if d["cols"] and not getattr(mapping, key):
                        setattr(mapping, key, pointage.guess(key, list(d["cols"])))
            for key in ("cat_key_col", "cat_label_col"):
                if cat_cols and not getattr(mapping, key):
                    setattr(mapping, key, pointage.guess(key, list(cat_cols)))
            mission_emp_guessed = bool(mission_cols) and not mapping.mission_emp_col
            # Colonnes des tables facultatives : proposées dès que la table est choisie.
            for key, cols in (("person_key_col", person_cols), ("person_nom_col", person_cols),
                              ("person_prenom_col", person_cols), ("hier_emp_col", hier_cols),
                              ("hier_manager_col", hier_cols), ("leave_emp_col", leave_cols),
                              ("leave_start_col", leave_cols), ("leave_end_col", leave_cols),
                              ("leave_state_col", leave_cols), ("leave_type_col", leave_cols),
                              ("tw_emp_col", tw_cols), ("tw_start_col", tw_cols), ("tw_end_col", tw_cols),
                              ("tw_state_col", tw_cols), ("mission_emp_col", mission_cols),
                              ("mission_state_col", mission_cols)):
                if cols and not getattr(mapping, key):
                    setattr(mapping, key, pointage.guess(key, list(cols)))
            if mission_cols:
                # Toutes les colonnes de dates d'aller / de retour (ex. une par étape du voyage).
                dates = [c for c, t in mission_cols.items() if "date" in t or "timestamp" in t] or list(mission_cols)
                for key in ("mission_start_col", "mission_end_col"):
                    if not getattr(mapping, key):
                        setattr(mapping, key, pointage.guess_all(key, dates) or pointage.guess(key, dates))
                if mission_emp_guessed:
                    # Pas de matricule dans la feuille (ex. Smartsheet) : colonne d'adresses e-mail → lien par e-mail.
                    emails = pointage.email_columns(engine, mapping.schema, mapping.mission_table, mission_cols)
                    if emails:
                        mapping.mission_emp_col, mapping.mission_ref = emails[0], "email"
                    elif re.search(r"mail|courriel", mapping.mission_emp_col or "", re.I):
                        mapping.mission_ref = "email"
            if person_cols and not mapping.emp_person_col:
                mapping.emp_person_col = pointage.guess("emp_person_col", list(emp_cols))
            pc, ec = list(punch_cols), list(emp_cols)
            if not mapping.punch_emp_col:
                mapping.punch_emp_col = pointage.guess("punch_emp_col", pc)
            if not mapping.punch_ts_col:
                mapping.punch_ts_col = pointage.guess("punch_ts_col", pc, (mapping.punch_emp_col,))
            for key in ("emp_key_col", "emp_nom_col", "emp_prenom_col", "emp_service_col"):
                if not getattr(mapping, key) and not exploring:
                    setattr(mapping, key, pointage.guess(key, ec))
            if not mapping.emp_matricule_col and not exploring:
                mapping.emp_matricule_col = pointage.guess("emp_matricule_col", ec)
            if not exploring and cfg.installed_at is None and not mapping.hier_table:
                # Première configuration : hiérarchie souvent dans la table des employés (colonne « responsable »).
                boss = pointage.guess("hier_manager_col", ec)
                if boss:
                    mapping.hier_table, mapping.hier_manager_col = mapping.emp_table, boss
                    mapping.hier_emp_col = mapping.emp_key_col
                    hier_cols = emp_cols
            if not mapping.email_col and not exploring:
                src = person_cols if mapping.email_in != "emp" and mapping.person_table else emp_cols
                mapping.email_col = pointage.guess("email_col", list(src))
            for d in dims:  # suggestion : colonne « direction » / « site » de Personnel (enregistrée si l'on valide)
                px = d["prefix"]
                if not getattr(mapping, f"{px}_col") and not exploring:
                    for where, src in (("person", person_cols if mapping.person_table else {}), ("emp", emp_cols)):
                        found = pointage.guess(f"{px}_col", list(src))
                        if found:
                            setattr(mapping, f"{px}_in", where)
                            setattr(mapping, f"{px}_col", found)
                            break
            if not mapping.cat_col and not exploring:  # suggestion (enregistrée seulement si l'on valide)
                src = person_cols if mapping.person_table else emp_cols
                mapping.cat_in = "person" if mapping.person_table else "emp"
                mapping.cat_col = pointage.guess("cat_col", list(src))
            installed = cfg.installed_at is not None and pointage.is_installed(engine, mapping)
            for d in dims:
                if installed and getattr(mapping, f"{d['prefix']}_col"):
                    try:
                        d["values"] = pointage.dimension_values(engine, mapping, d["name"])[:20]
                    except Exception:
                        d["values"] = []
            if installed and mapping.cat_col:
                try:
                    cat_values = pointage.categories(engine, mapping)[:15]
                except Exception:
                    cat_values = []
            if installed and mapping.mission_table:
                try:
                    mission_match = pointage.unmatched_requests(engine, mapping, "mission")
                except Exception:
                    mission_match = None
            if installed and not exploring:
                diag = pointage.diagnostics(engine, mapping)
                last = diag["dernier_pointage"]
                stale = last is not None and (date.today() - last.date()).days > 3
        except Exception as exc:
            error = friendly(exc)
        finally:
            engine.dispose()
    return render(
        request, "admin/pointage.html", cfg=cfg, m=mapping, conn=conn, connections=connections, schemas=schemas,
        tables=tables, punch_cols=punch_cols, emp_cols=emp_cols, service_cols=service_cols, error=error,
        person_cols=person_cols, hier_cols=hier_cols, leave_cols=leave_cols, diag=diag, stale=stale,
        leave_guess=pointage.guess("leave_table", tables) if not mapping.leave_table else "",
        tw_cols=tw_cols, cat_cols=cat_cols, cat_values=cat_values, tw_guess=pointage.guess("tw_table", tables) if not mapping.tw_table else "",
        installed=installed, exploring=exploring, mission_cols=mission_cols, dims=dims, mission_match=mission_match,
        mission_guess=pointage.guess("mission_table", tables) if not mapping.mission_table else "",
    )


@router.post("/pointage")
async def pointage_save(request: Request, db: Session = Depends(get_db)):
    form = await request.form()
    cfg = _config(db)
    conn_id = str(form.get("conn_id", ""))
    conn = db.get(Connection, int(conn_id)) if conn_id.isdigit() else None
    if conn is None or conn.kind != "postgresql":
        flash(request, "Choisissez la base PostgreSQL qui contient les pointages.", "err")
        return redirect("/admin/pointage")
    mapping = pointage.Mapping(**{k: _field(form, k) for k in _MAPPING_FIELDS})
    before = cfg.data
    engine = make_engine(conn, **pointage.WEB_LIMITS)
    try:
        pointage.install(engine, mapping, request.session.get("user", ""))
    except pointage.PointageError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/pointage?" + _query(conn.id, mapping))
    except Exception as exc:
        flash(request, f"Installation impossible : {friendly(exc)}", "err")
        return redirect("/admin/pointage?" + _query(conn.id, mapping))
    finally:
        engine.dispose()
    cfg.conn_id, cfg.data, cfg.installed_at = conn.id, mapping.to_json(), utcnow()
    cfg.sql_version = pointage.SQL_VERSION
    db.commit()
    auth.audit(request, "Source des pointages configurée", f"{conn.name} · {mapping.schema}",
               f"avant : {before}\naprès : {cfg.data}")
    flash(request, f"Configuration enregistrée : fonction {mapping.objs}.f_pointage_journalier et vue "
                   f"{mapping.objs}.v_pointage_journalier installées dans {conn.database}.", "ok")
    return redirect("/admin/pointage")


def _query(conn_id: int, m: pointage.Mapping) -> str:
    from urllib.parse import urlencode

    return urlencode({"conn_id": conn_id, **{k: getattr(m, k) for k in _MAPPING_FIELDS}})


# --------------------------------------------------------------------------- paramètres horaires


@router.get("/parametres")
def params_page(request: Request, db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    context = dict(mapping=mapping, params=pointage.PARAMS, weekdays=pointage.WEEKDAYS, today=date.today(),
                   values=dict(pointage.PARAM_DEFAULTS), history=[], holidays=[], error=None, upcoming=[])
    if mapping is not None:
        engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
        try:
            context["values"] = pointage.params_at(engine, mapping, date.today())
            context["history"] = pointage.params_history(engine, mapping)
            context["holidays"] = pointage.holidays(engine, mapping)
            context["upcoming"] = [h for h in context["history"] if h.date_effet > date.today()]
        except Exception as exc:
            context["error"] = friendly(exc)
        finally:
            engine.dispose()
    return render(request, "admin/parametres.html", **context)


@router.post("/parametres")
async def params_save(request: Request, db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    form = await request.form()
    effective = pointage.parse_day(str(form.get("date_effet", "")))
    if effective is None:
        flash(request, "Indiquez la date d'effet.", "err")
        return redirect("/admin/parametres")
    try:
        values = {}
        current = None
        for key, _, kind, _ in pointage.PARAMS:
            if kind != "days" and key not in form:  # page ouverte avant l'ajout de ce paramètre : inchangé
                if current is None:
                    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
                    try:
                        current = pointage.params_at(engine, mapping, effective)
                    finally:
                        engine.dispose()
                values[key] = current.get(key, pointage.PARAM_DEFAULTS[key])
                continue
            raw = ",".join(form.getlist("jours_ouvres")) if kind == "days" else str(form.get(key, ""))
            values[key] = pointage.validate_param(key, raw)
        problem = pointage.check_consistency(values)
        if problem:
            raise pointage.PointageError(problem)
    except pointage.PointageError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/parametres")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        changes = pointage.save_params(engine, mapping, values, effective, request.session.get("user", ""))
    except Exception as exc:
        flash(request, f"Enregistrement impossible : {friendly(exc)}", "err")
        return redirect("/admin/parametres")
    finally:
        engine.dispose()
    if not changes:
        flash(request, "Aucun paramètre modifié.", "info")
        return redirect("/admin/parametres")
    details = "\n".join(f"{pointage.PARAM_LABELS[k]} : {old} → {new}" for k, (old, new) in changes.items())
    auth.audit(request, "Paramètres horaires modifiés", f"date d'effet {effective:%d/%m/%Y}", details)
    when = "à partir d'aujourd'hui" if effective == date.today() else f"à partir du {effective:%d/%m/%Y}"
    flash(request, f"{len(changes)} paramètre(s) modifié(s), appliqué(s) {when}. Les jours antérieurs gardent "
                   f"les valeurs de l'époque.", "ok")
    return redirect("/admin/parametres")


@router.post("/feries")
def holiday_add(request: Request, jour: str = Form(...), libelle: str = Form(""), db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    day = pointage.parse_day(jour)
    if mapping is None or day is None:
        flash(request, "Date invalide.", "err")
        return redirect("/admin/parametres")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        pointage.add_holiday(engine, mapping, day, libelle.strip(), request.session.get("user", ""))
    finally:
        engine.dispose()
    auth.audit(request, "Jour férié ajouté", f"{day:%d/%m/%Y}", libelle.strip())
    flash(request, f"Jour férié du {day:%d/%m/%Y} enregistré : personne n'y sera compté absent.", "ok")
    return redirect("/admin/parametres#feries")


@router.post("/feries/delete")
def holiday_delete(request: Request, jour: str = Form(...), db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    day = pointage.parse_day(jour)
    if mapping is None or day is None:
        return redirect("/admin/parametres")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        pointage.delete_holiday(engine, mapping, day)
    finally:
        engine.dispose()
    auth.audit(request, "Jour férié supprimé", f"{day:%d/%m/%Y}")
    flash(request, f"Jour férié du {day:%d/%m/%Y} supprimé.", "ok")
    return redirect("/admin/parametres#feries")


# --------------------------------------------------------------------------- agents terrain


@router.get("/terrain")
def field_page(request: Request, db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    context = dict(mapping=mapping, entries=[], services=[], error=None, types=pointage.FIELD_TYPES,
                   duree=pointage.PARAM_DEFAULTS["duree_terrain"])
    if mapping is not None:
        engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
        try:
            context["entries"] = pointage.field_entries(engine, mapping)
            context["services"] = pointage.services(engine, mapping)
            context["duree"] = pointage.params_at(engine, mapping, date.today())["duree_terrain"]
        except Exception as exc:
            context["error"] = friendly(exc)
        finally:
            engine.dispose()
    return render(request, "admin/terrain.html", **context)


@router.post("/terrain")
def field_add(request: Request, type: str = Form(...), service: str = Form(""), matricule: str = Form(""),
              libelle: str = Form(""), db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    value = service if type == "service" else matricule
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        pointage.add_field(engine, mapping, type, value, libelle, request.session.get("user", ""))
    except pointage.PointageError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/terrain")
    finally:
        engine.dispose()
    what = f"service « {value.strip()} »" if type == "service" else f"employé {value.strip()}"
    auth.audit(request, "Agent terrain ajouté", what, libelle.strip())
    flash(request, f"{what[0].upper()}{what[1:]} déclaré agent terrain : ses journées ne seront plus comptées "
                   f"en absence ni en retard.", "ok")
    return redirect("/admin/terrain")


MAX_IMPORT_MB = pointage.MAX_IMPORT_MB
MAX_IMPORT_BYTES = MAX_IMPORT_MB * 1024 * 1024
MAX_IMPORT_FILES = 30  # classeurs de planning importés en une fois
# Colonne du matricule, par ordre de préférence (« Nom employé » ne doit pas être pris pour le matricule).
_MAT_HEADERS = [re.compile(p, re.I) for p in (r"matric", r"badge", r"^n[°o]\s*(d.)?employ", r"^code", r"^id")]
_LABEL_HEADER = re.compile(r"motif|libell|fonction|poste|comment|zone|remarque", re.I)


def read_field_file(name: str, data: bytes) -> list[tuple[str, str]]:
    """Lignes (matricule, motif) d'un fichier Excel (.xlsx) ou CSV. En-tête facultatif ; sinon colonne A = matricule,
    colonne B = motif. Lève ValueError avec un message lisible."""
    lower = (name or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        import io

        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValueError("Fichier Excel illisible (enregistrez-le au format .xlsx).") from exc
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif lower.endswith((".csv", ".txt")):
        import csv

        for encoding in ("utf-8-sig", "cp1252"):
            try:
                content = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        delimiter = ";" if content.count(";") >= content.count(",") else ","
        rows = [r for r in csv.reader(content.splitlines(), delimiter=delimiter)]
    elif lower.endswith(".xls"):
        raise ValueError("Ancien format Excel (.xls) : enregistrez le fichier au format .xlsx puis réessayez.")
    else:
        raise ValueError("Format non pris en charge : choisissez un fichier Excel (.xlsx) ou CSV.")
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)][:20000]
    if not rows:
        raise ValueError("Le fichier est vide.")
    header = [str(c or "").strip() for c in rows[0]]
    mat_col = next((i for pattern in _MAT_HEADERS for i, h in enumerate(header) if pattern.search(h)), None)
    label_col = next((i for i, h in enumerate(header) if i != mat_col and _LABEL_HEADER.search(h)), None)
    if mat_col is not None:
        rows = rows[1:]
    else:
        mat_col, label_col = 0, 1
    out = []
    for r in rows:
        mat = pointage.clean_matricule(r[mat_col] if mat_col < len(r) else None)
        label = r[label_col] if label_col is not None and label_col < len(r) else ""
        if mat:
            out.append((mat, str(label or "").strip()))
    if not out:
        raise ValueError("Aucun matricule trouvé (colonne « Matricule » ou première colonne).")
    return out


@router.post("/terrain/import")
async def field_import(request: Request, fichier: UploadFile = File(...), mode: str = Form("ajouter"),
                       db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    data = await fichier.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        flash(request, f"Fichier trop volumineux ({MAX_IMPORT_MB} Mo maximum).", "err")
        return redirect("/admin/terrain")
    try:
        rows = read_field_file(fichier.filename or "", data)
    except ValueError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/terrain")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        result = pointage.import_fields(engine, mapping, rows, request.session.get("user", ""), replace=mode == "remplacer")
    except Exception as exc:  # noqa: BLE001
        flash(request, f"Import impossible : {friendly(exc)}", "err")
        return redirect("/admin/terrain")
    finally:
        engine.dispose()
    auth.audit(request, "Agents terrain importés", fichier.filename or "",
               f"{result['lus']} matricule(s) lus, {result['ajoutes']} ajouté(s), {result['mis_a_jour']} mis à jour, "
               f"{result['retires']} retiré(s), {len(result['inconnus'])} inconnu(s)")
    flash(request, f"Import de « {fichier.filename} » : {result['ajoutes']} agent(s) ajouté(s), {result['mis_a_jour']} "
                   f"déjà présent(s) mis à jour" + (f", {result['retires']} retiré(s) (absents du fichier)"
                                                    if mode == "remplacer" else "") + ".", "ok")
    if result["inconnus"]:
        shown = ", ".join(result["inconnus"][:30]) + (" …" if len(result["inconnus"]) > 30 else "")
        flash(request, f"{len(result['inconnus'])} matricule(s) introuvable(s) dans la liste des employés, ignoré(s) : "
                       f"{shown}", "warn")
    return redirect("/admin/terrain")


def _xlsx(headers: list[str], rows: list[list], name: str) -> Response:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Agents terrain"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, _ in enumerate(headers, 1):
        ws.cell(1, i).font = Font(bold=True)
        ws.column_dimensions[chr(64 + i)].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/terrain/modele.xlsx")
def field_template():
    return _xlsx(["Matricule", "Motif"], [["590394", "Commercial zone Nord"], ["590412", "Force de vente"]],
                 "modele_agents_terrain.xlsx")


@router.get("/terrain/export.xlsx")
def field_export(db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        entries = pointage.field_entries(engine, mapping)
    finally:
        engine.dispose()
    rows = [[e["valeur"], e["libelle"] or "", e["nom"] or ""] for e in entries if e["type"] == "employe"]
    return _xlsx(["Matricule", "Motif", "Nom"], rows, "agents_terrain.xlsx")


@router.post("/terrain/{entry_id}/delete")
def field_delete(request: Request, entry_id: int, db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        removed = pointage.delete_field(engine, mapping, entry_id)
    finally:
        engine.dispose()
    if removed:
        auth.audit(request, "Agent terrain retiré", f"{removed[0]} « {removed[1]} »")
        flash(request, f"« {removed[1]} » n'est plus agent terrain.", "ok")
    return redirect("/admin/terrain")


# --------------------------------------------------------------------------- horaires postés (planning)


# Au-delà, la grille devient illisible et lourde : la période affichée est ramenée à un an.
MAX_GRID_DAYS = 366


def _month_bounds(day: date) -> tuple[date, date]:
    first = day.replace(day=1)
    return first, (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)


def _planning_shortcuts(start: date, end: date, first: Optional[date], last: Optional[date]) -> dict:
    """Raccourcis de période : mois, période planifiée, période précédente / suivante de même durée."""
    today = date.today()
    span = end - start + timedelta(days=1)
    this_month = _month_bounds(today)
    items = [("Mois en cours", *this_month),
             ("Mois précédent", *_month_bounds(this_month[0] - timedelta(days=1))),
             ("Mois suivant", *_month_bounds(this_month[1] + timedelta(days=1))),
             ("4 dernières semaines", today - timedelta(days=27), today)]
    if first and last:
        items.append(("Tout le planning", first, last))
    return {"items": items, "prev": (start - span, start - timedelta(days=1)),
            "next": (end + timedelta(days=1), end + span)}


def _planning_back(du=None, au=None) -> str:
    return f"/admin/planning?du={du.isoformat()}&au={au.isoformat()}" if du and au else "/admin/planning"


@router.get("/planning")
def planning_page(request: Request, du: str = "", au: str = "", db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    context = dict(mapping=mapping, postes=[], grid=None, error=None, types=pointage.POSTE_TYPES,
                   marge=pointage.PARAM_DEFAULTS["marge_poste"], du=None, au=None, nuit=set(),
                   tronque=False, raccourcis=None, max_jours=MAX_GRID_DAYS,
                   duree_repos=pointage.PARAM_DEFAULTS["duree_repos"])
    if mapping is not None:
        engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
        try:
            context["postes"] = pointage.postes(engine, mapping)
            context["nuit"] = {p["code"] for p in context["postes"] if p["type"] == "travail" and p["fin"] <= p["debut"]}
            current = pointage.params_at(engine, mapping, date.today())
            context["marge"], context["duree_repos"] = current["marge_poste"], current["duree_repos"]
            start, end = pointage.parse_day(du), pointage.parse_day(au)
            if not (start and end):
                bounds = pointage.planning_grid(engine, mapping, date.today(), date.today())
                if bounds["dernier"]:  # par défaut : la dernière période planifiée (5 semaines au plus)
                    end = bounds["dernier"]
                    start = max(bounds["premier"], end - timedelta(days=34))
                else:
                    start, end = date.today().replace(day=1), date.today()
            if start > end:
                start, end = end, start
            context["tronque"] = (end - start).days + 1 > MAX_GRID_DAYS
            end = min(end, start + timedelta(days=MAX_GRID_DAYS - 1))
            grid = pointage.planning_grid(engine, mapping, start, end)
            context.update(du=start, au=end, grid=grid,
                           raccourcis=_planning_shortcuts(start, end, grid["premier"], grid["dernier"]))
        except Exception as exc:
            context["error"] = friendly(exc)
        finally:
            engine.dispose()
    return render(request, "admin/planning.html", **context)


async def _read_uploads(fichier: list[UploadFile]) -> tuple[list[tuple[str, bytes]], list[tuple[str, str]]]:
    """Fichiers envoyés (au plus MAX_IMPORT_FILES, MAX_IMPORT_MB chacun) et messages sur ceux qui sont écartés."""
    files, too_big, notes = [], [], []
    for upload in fichier[:MAX_IMPORT_FILES]:
        if not upload.filename:
            continue
        data = await upload.read(MAX_IMPORT_BYTES + 1)
        if len(data) > MAX_IMPORT_BYTES:
            too_big.append(upload.filename)
        else:
            files.append((upload.filename, data))
    if len(fichier) > MAX_IMPORT_FILES:
        notes.append(("warn", f"{MAX_IMPORT_FILES} fichiers au plus par import : les {len(fichier) - MAX_IMPORT_FILES} "
                              f"suivant(s) ont été ignorés."))
    if too_big:
        notes.append(("err", f"Fichier trop volumineux ({MAX_IMPORT_MB} Mo maximum), ignoré : "
                             + ", ".join(f"« {n} »" for n in too_big) + "."))
    if not files and not too_big:
        notes.append(("err", "Choisissez au moins un classeur Excel (.xlsx)."))
    return files, notes


def _run_planning_import(engine, mapping, files: list[tuple[str, bytes]], author: str, progress=None) -> dict:
    """Lecture puis enregistrement ; renvoie {messages, url, audit} (les erreurs deviennent des messages)."""
    try:
        plan = pointage.read_planning_files(files, progress)
        result = pointage.import_planning(engine, mapping, plan, author, progress)
    except pointage.PointageError as exc:
        return {"messages": [("err", str(exc))], "url": "/admin/planning", "audit": None}
    except Exception as exc:  # noqa: BLE001
        return {"messages": [("err", f"Import impossible : {friendly(exc)}")], "url": "/admin/planning", "audit": None}
    finally:
        engine.dispose()
    return {"messages": _planning_messages(plan, result, files), "url": _planning_back(result["du"], result["au"]),
            "audit": (", ".join(n for n, _ in files)[:500],
                      f"du {result['du']:%d/%m/%Y} au {result['au']:%d/%m/%Y} : {result['employes']} employé(s), "
                      f"{result['jours']} jour(s) planifié(s)")}


def _planning_messages(plan: dict, result: dict, files: list) -> list[tuple[str, str]]:
    out = []
    label = files[0][0] if len(files) == 1 else f"{len(files)} classeurs"
    out.append(("ok", f"Planning « {label} » importé du {result['du']:%d/%m/%Y} au {result['au']:%d/%m/%Y} : "
                      f"{result['employes']} employé(s), {result['jours']} jour(s) planifié(s). "
                      f"Les calculs du suivi en tiennent compte immédiatement."))
    sheets = plan.get("feuilles", [])
    if len(sheets) > 1:
        by_file: dict[str, list] = {}
        for x in sheets:
            by_file.setdefault(x["fichier"], []).append(x)
        summary = " ; ".join(
            f"« {name} » : {len(xs)} feuille(s), du {min(x['du'] for x in xs):%d/%m/%Y} au "
            f"{max(x['au'] for x in xs):%d/%m/%Y}" for name, xs in by_file.items())
        out.append(("ok", f"{len(sheets)} feuilles importées — {summary}."))
    if plan.get("fichiers_illisibles"):
        out.append(("warn", "Fichier(s) ignoré(s) : " + " ".join(plan["fichiers_illisibles"])))
    if plan.get("feuilles_ignorees"):
        out.append(("warn", "Feuilles ignorées (pas au format du planning : ligne « MATRICULE » puis une ligne par "
                            "jour) : " + ", ".join(f"« {x} »" for x in plan["feuilles_ignorees"]) + "."))
    if plan.get("conflits"):
        sample = "; ".join(f"{mat} le {day:%d/%m/%Y} : {c1} (« {f1} ») remplacé par {c2} (« {f2} »)"
                           for mat, day, f1, c1, f2, c2 in plan["conflits"][:2])
        out.append(("warn", f"{len(plan['conflits'])} jour(s) planifié(s) dans plusieurs feuilles avec des codes "
                            f"différents : la dernière feuille (dans l'ordre des fichiers puis des feuilles) l'emporte. "
                            f"Ex. {sample}."))
    if result["surcharges"]:
        out.append(("warn", f"Planning chargé (plus de {pointage.PLANNING_WEEKLY_MAX} h par semaine en moyenne) : "
                            + ", ".join(f"{mat} — {h['heures']:.0f} h sur la période, soit {h['hebdo']:.1f} h/semaine"
                                        for mat, h in result["surcharges"])
                            + ". Vérifiez qu'il ne s'agit pas d'une erreur de saisie."))
    if result["codes_crees"]:
        out.append(("warn", "Nouveaux codes créés depuis la légende du fichier (à vérifier ci-dessous) : "
                            + ", ".join(result["codes_crees"]) + "."))
    if result["codes_inconnus"]:
        out.append(("warn", "Codes inconnus, jours ignorés : " + ", ".join(
            f"« {k} » ({n} j)" for k, n in sorted(result["codes_inconnus"].items()))
            + ". Ajoutez ces codes puis réimportez le fichier."))
    if result["matricules_inconnus"]:
        out.append(("warn", f"{len(result['matricules_inconnus'])} matricule(s) introuvable(s) dans la liste des "
                            f"employés, ignoré(s) : " + ", ".join(result["matricules_inconnus"][:15])
                            + (" …" if len(result["matricules_inconnus"]) > 15 else "")))
    return out


def _finish(request: Request, outcome: dict):
    for category, message in outcome["messages"]:
        flash(request, message, category)
    if outcome.get("audit"):
        auth.audit(request, "Planning importé", *outcome["audit"])
    return redirect(outcome["url"])


@router.post("/planning/import")
async def planning_import(request: Request, fichier: list[UploadFile] = File(...), db: Session = Depends(get_db)):
    """Un ou plusieurs classeurs à la fois (formulaire classique, sans suivi de progression)."""
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    files, notes = await _read_uploads(fichier)
    for category, message in notes:
        flash(request, message, category)
    if not files:
        return redirect("/admin/planning")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    return _finish(request, _run_planning_import(engine, mapping, files, request.session.get("user", "")))


# Imports suivis en direct : traitement en arrière-plan, avancement lu par la page toutes les demi-secondes.
_IMPORTS: dict[str, dict] = {}
_IMPORTS_LOCK = threading.Lock()


@router.post("/planning/import/demarrer")
async def planning_import_start(request: Request, fichier: list[UploadFile] = File(...),
                                db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return JSONResponse({"error": "Configurez d'abord la source des pointages."}, status_code=400)
    files, notes = await _read_uploads(fichier)
    user = request.session.get("user", "")
    token = secrets.token_urlsafe(16)
    job = {"user": user, "state": "running", "percent": 0.0, "label": "Préparation de l'import", "notes": notes,
           "outcome": None, "started": time.monotonic()}
    with _IMPORTS_LOCK:
        for key in [k for k, v in _IMPORTS.items() if time.monotonic() - v["started"] > 3600]:
            del _IMPORTS[key]  # imports anciens (page fermée avant la fin)
        _IMPORTS[token] = job
    if not files:
        job.update(state="done", percent=100.0, outcome={"messages": [], "url": "/admin/planning", "audit": None})
        return {"job": token}
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)

    def progress(percent: float, label: str) -> None:
        job["percent"], job["label"] = round(max(job["percent"], min(float(percent), 99.0)), 1), label

    def work() -> None:
        outcome = _run_planning_import(engine, mapping, files, user, progress)
        job.update(outcome=outcome, percent=100.0, label="Import terminé", state="done")

    threading.Thread(target=work, daemon=True, name="import-planning").start()
    return {"job": token}


def _job(request: Request, token: str) -> Optional[dict]:
    job = _IMPORTS.get(token)
    return job if job is not None and job["user"] == request.session.get("user", "") else None


@router.get("/planning/import/{token}")
def planning_import_status(request: Request, token: str):
    job = _job(request, token)
    if job is None:
        return JSONResponse({"error": "Import introuvable (déjà terminé ou expiré)."}, status_code=404)
    return {"state": job["state"], "percent": job["percent"], "label": job["label"],
            "elapsed": round(time.monotonic() - job["started"])}


@router.get("/planning/import/{token}/fin")
def planning_import_finish(request: Request, token: str):
    job = _job(request, token)
    if job is None or job["state"] != "done":
        return redirect("/admin/planning")
    with _IMPORTS_LOCK:
        _IMPORTS.pop(token, None)
    for category, message in job["notes"]:
        flash(request, message, category)
    return _finish(request, job["outcome"])


@router.post("/planning/postes")
def poste_save(request: Request, code: str = Form(...), libelle: str = Form(""), type: str = Form(...),
               debut: str = Form(""), fin: str = Form(""), pause: str = Form(""), duree: str = Form(""),
               db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        saved = pointage.save_poste(engine, mapping, request.session.get("user", ""), code, libelle, type,
                                    debut, fin, pause, duree)
    except pointage.PointageError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/planning#postes")
    finally:
        engine.dispose()
    auth.audit(request, "Code de poste enregistré", saved, f"{pointage.POSTE_TYPES.get(type, type)} {debut}-{fin}".strip())
    flash(request, f"Code « {saved} » enregistré.", "ok")
    return redirect("/admin/planning#postes")


@router.post("/planning/postes/{code}/delete")
def poste_delete(request: Request, code: str, db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    if mapping is None:
        return redirect("/admin/pointage")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        pointage.delete_poste(engine, mapping, code)
    except pointage.PointageError as exc:
        flash(request, str(exc), "err")
        return redirect("/admin/planning#postes")
    finally:
        engine.dispose()
    auth.audit(request, "Code de poste supprimé", code)
    flash(request, f"Code « {code} » supprimé.", "ok")
    return redirect("/admin/planning#postes")


@router.post("/planning/supprimer")
def planning_delete(request: Request, du: str = Form(...), au: str = Form(...), matricule: str = Form(""),
                    db: Session = Depends(get_db)):
    cfg, mapping = _mapping_ready(db)
    start, end = pointage.parse_day(du), pointage.parse_day(au)
    if mapping is None or not (start and end):
        return redirect("/admin/planning")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        n = pointage.delete_planning(engine, mapping, start, end, matricule.strip())
    finally:
        engine.dispose()
    who = f" de {matricule.strip()}" if matricule.strip() else ""
    auth.audit(request, "Planning supprimé", f"du {start:%d/%m/%Y} au {end:%d/%m/%Y}{who}", f"{n} jour(s)")
    flash(request, f"{n} jour(s) planifié(s){who} supprimé(s) : ces jours suivent de nouveau l'horaire de bureau.", "ok")
    return redirect(_planning_back(start, end))


@router.get("/planning/export.xlsx")
def planning_export(du: str = "", au: str = "", db: Session = Depends(get_db)):
    """Planning de la période au format de l'import (réimportable après modification)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    cfg, mapping = _mapping_ready(db)
    start, end = pointage.parse_day(du), pointage.parse_day(au)
    if mapping is None or not (start and end):
        return redirect("/admin/planning")
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        grid = pointage.planning_grid(engine, mapping, start, min(end, start + timedelta(days=MAX_GRID_DAYS - 1)))
        codes = pointage.postes(engine, mapping)
    finally:
        engine.dispose()
    wb = Workbook()
    ws = wb.active
    ws.title = "PLANNING"
    ws.append(["", "Code", "Heure Début", "Heure Fin", "Libellé"])
    for p in codes:
        ws.append(["", p["code"], p["debut"] if p["type"] == "travail" else p["libelle"].upper(),
                   p["fin"] if p["type"] == "travail" else None, p["libelle"]])
    ws.append([])
    ws.append(["MATRICULE"] + [e["matricule"] for e in grid["employes"]])
    ws.append(["NOM"] + [e["nom"] for e in grid["employes"]])
    ws.append(["DATE"])
    for day in grid["jours"]:
        ws.append([day] + [grid["grille"].get(e["matricule"], {}).get(day, ("",))[0] for e in grid["employes"]])
        ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in ("Code", "MATRICULE", "NOM", "DATE"):
                cell.font = Font(bold=True)
            if hasattr(cell.value, "hour") and not hasattr(cell.value, "year"):
                cell.number_format = "HH:MM"
            if isinstance(cell.value, str) and cell.value.isdigit():
                cell.number_format = "@"
    ws.column_dimensions["A"].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="planning_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"'})


# --------------------------------------------------------------------------- utilisateurs


_PASSWORD_ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # sans 0/O, 1/l/I


def _employee_accounts(db: Session) -> dict:
    """Employés actifs de la liste et comptes existants : qui n'a pas encore de compte."""
    cfg, mapping = _mapping_ready(db)
    out = {"configured": mapping is not None, "hierarchy": bool(mapping and mapping.hier_table), "employees": [],
           "missing": [], "error": None}
    if mapping is None:
        return out
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        out["employees"] = [e for e in pointage.employee_directory(engine, mapping) if e["matricule"]]
    except Exception as exc:  # noqa: BLE001
        out["error"] = friendly(exc)
        return out
    finally:
        engine.dispose()
    users = db.scalars(select(User)).all()
    taken = {u.username.lower() for u in users} | {(u.emp_matricule or "").lower() for u in users if u.emp_matricule}
    out["missing"] = [e for e in out["employees"] if e["matricule"].lower() not in taken]
    return out


@router.get("/users")
def users(request: Request, db: Session = Depends(get_db)):
    return render(request, "admin/users.html", users=db.scalars(select(User).order_by(User.username)).all(),
                  roles=ROLES, rescue=auth.settings.admin_username, accounts=_employee_accounts(db))


@router.post("/users/comptes-employes")
def create_employee_accounts(request: Request, mode: str = Form("aleatoire"), mot_de_passe: str = Form(""),
                             db: Session = Depends(get_db)):
    """Crée un compte pour chaque employé actif qui n'en a pas : identifiant = matricule, rôle lecteur, périmètre
    « son équipe » (lui-même et tous ceux qui sont sous lui), mot de passe provisoire à changer à la connexion.
    Renvoie le fichier Excel des identifiants et mots de passe provisoires (affichés une seule fois)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    info = _employee_accounts(db)
    if not info["configured"] or info["error"]:
        flash(request, info["error"] or "Configurez d'abord la source des pointages.", "err")
        return redirect("/admin/users")
    if not info["hierarchy"]:
        flash(request, "La hiérarchie (responsable N+1) n'est pas configurée : chaque employé ne pourrait pas voir son "
                       "équipe. Configurez-la dans Source des pointages → Hiérarchie, puis recommencez.", "err")
        return redirect("/admin/users")
    if mode == "commun" and auth.password_problem(mot_de_passe):
        flash(request, f"Mot de passe provisoire commun : {auth.password_problem(mot_de_passe)}", "err")
        return redirect("/admin/users")
    if not info["missing"]:
        flash(request, "Tous les employés actifs ont déjà un compte.", "ok")
        return redirect("/admin/users")
    created = []
    for e in info["missing"]:
        password = mot_de_passe if mode == "commun" else "".join(secrets.choice(_PASSWORD_ALPHABET) for _ in range(10))
        name = " ".join(x for x in (e["nom"], e["prenom"]) if x)
        db.add(User(username=e["matricule"], full_name=name[:200], role="lecteur", scope="equipe",
                    emp_matricule=e["matricule"], email=(e.get("email") or None), active=True,
                    password_hash=auth.hash_password(password), must_change_password=True))
        created.append((e, name, password))
    db.commit()
    auth.audit(request, "Comptes employés créés", f"{len(created)} compte(s)",
               f"identifiant = matricule, rôle lecteur, périmètre son équipe, mot de passe "
               f"{'commun' if mode == 'commun' else 'aléatoire'} à changer à la connexion")
    flash(request, f"{len(created)} compte(s) employé créé(s). Le fichier Excel des identifiants et mots de passe "
                   f"provisoires vient d'être téléchargé : transmettez-les aux employés (ils devront choisir leur "
                   f"mot de passe à la première connexion).", "ok")
    wb = Workbook()
    ws = wb.active
    ws.title = "Comptes employés"
    ws.append(["Matricule", "Nom", "Service", "Identifiant", "Mot de passe provisoire", "Adresse de connexion"])
    base = str(request.base_url).rstrip("/") + "/login"
    for e, name, password in created:
        ws.append([e["matricule"], name, e.get("service") or "", e["matricule"], password, base])
    for col, width in zip("ABCDEF", (12, 30, 24, 14, 22, 34)):
        ws.column_dimensions[col].width = width
        ws[f"{col}1"].font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row[:5]:
            cell.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="comptes_employes_{date.today():%Y%m%d}.xlsx"',
                             "Cache-Control": "no-store"})


@router.post("/users/save")
def user_save(
    request: Request,
    user_id: int = Form(0),
    username: str = Form(""),
    full_name: str = Form(""),
    role: str = Form("lecteur"),
    active: bool = Form(False),
    password: str = Form(""),
    emp_matricule: str = Form(""),
    scope: str = Form("tous"),
    sick_leave_hr: bool = Form(False),
    email: str = Form(""),
    db: Session = Depends(get_db),
):
    me = request.session.get("user", "")
    if role not in ROLES:
        flash(request, "Rôle inconnu.", "err")
        return redirect("/admin/users")
    scope = scope if scope in ("tous", "equipe") else "tous"
    from ..mails import valid_email

    email = email.strip()
    if email and not valid_email(email):
        flash(request, f"Adresse e-mail invalide : « {email} ».", "err")
        return redirect("/admin/users")
    if scope == "equipe" and not emp_matricule.strip():
        flash(request, "Pour limiter un compte à son équipe, indiquez le matricule de l'employé correspondant.", "err")
        return redirect("/admin/users")
    user = db.get(User, user_id) if user_id else User()
    if user is None:
        return redirect("/admin/users")
    if not user_id:
        username = username.strip()
        if not username or auth.is_rescue_admin(username):
            flash(request, "Identifiant vide ou réservé au compte de secours.", "err")
            return redirect("/admin/users")
        problem = auth.password_problem(password)
        if problem:
            flash(request, problem, "err")
            return redirect("/admin/users")
        user.username, user.password_hash = username, auth.hash_password(password)
    elif user.username == me and (role != user.role or not active):
        flash(request, "Vous ne pouvez pas modifier votre propre rôle ni désactiver votre compte.", "err")
        return redirect("/admin/users")
    def describe(u: User) -> str:
        return (f"rôle {u.role}, {'actif' if u.active else 'inactif'}, matricule {u.emp_matricule or '—'}, "
                f"périmètre {'équipe' if u.scope == 'equipe' else 'tous'}, "
                f"RH arrêts maladie {'oui' if u.sick_leave_hr else 'non'}, e-mail {u.email or '—'}")

    before = describe(user) if user_id else "création"
    user.full_name, user.role, user.active = full_name.strip(), role, active if user_id else True
    user.emp_matricule, user.scope = emp_matricule.strip() or None, scope
    user.sick_leave_hr = sick_leave_hr
    user.email = email or None
    if not user_id:
        db.add(user)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        flash(request, f"L'identifiant « {username} » existe déjà.", "err")
        return redirect("/admin/users")
    auth.audit(request, "Utilisateur créé" if not user_id else "Utilisateur modifié", user.username,
               f"{before} → {describe(user)}")
    flash(request, f"Utilisateur « {user.username} » enregistré.", "ok")
    return redirect("/admin/users")


@router.post("/users/{user_id}/password")
def user_password(request: Request, user_id: int, password: str = Form(...), db: Session = Depends(get_db)):
    user = db.get(User, user_id)
    problem = auth.password_problem(password)
    if user is None or problem:
        flash(request, problem or "Utilisateur introuvable.", "err")
        return redirect("/admin/users")
    user.password_hash = auth.hash_password(password)
    user.must_change_password = True  # mot de passe connu de l'administrateur : à changer à la connexion
    db.commit()
    auth.audit(request, "Mot de passe réinitialisé", user.username)
    flash(request, f"Mot de passe de « {user.username} » réinitialisé.", "ok")
    return redirect("/admin/users")


@router.post("/users/{user_id}/delete")
def user_delete(request: Request, user_id: int, db: Session = Depends(get_db)):
    user = db.get(User, user_id)
    if user is None:
        return redirect("/admin/users")
    if user.username == request.session.get("user"):
        flash(request, "Vous ne pouvez pas supprimer votre propre compte.", "err")
        return redirect("/admin/users")
    name = user.username
    db.delete(user)
    db.commit()
    auth.audit(request, "Utilisateur supprimé", name)
    flash(request, f"Utilisateur « {name} » supprimé.", "ok")
    return redirect("/admin/users")


# --------------------------------------------------------------------------- journal d'audit


@router.get("/audit")
def audit_log(request: Request, user: str = "", q: str = "", du: str = "", au: str = "", page: int = 1,
              db: Session = Depends(get_db)):
    size = 50
    stmt = select(AuditEntry)
    if user:
        stmt = stmt.where(AuditEntry.username == user)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(AuditEntry.action.ilike(like), AuditEntry.target.ilike(like),
                              AuditEntry.details.ilike(like)))
    d1, d2 = pointage.parse_day(du), pointage.parse_day(au)
    if d1:
        stmt = stmt.where(AuditEntry.created_at >= datetime.combine(d1, datetime.min.time()))
    if d2:
        stmt = stmt.where(AuditEntry.created_at < datetime.combine(d2 + timedelta(days=1), datetime.min.time()))
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    page = max(page, 1)
    entries = db.scalars(stmt.order_by(AuditEntry.id.desc()).offset((page - 1) * size).limit(size)).all()
    usernames = [u for u in db.scalars(select(AuditEntry.username).distinct().order_by(AuditEntry.username)) if u]
    return render(request, "admin/audit.html", entries=entries, total=total, page=page,
                  pages=max((total + size - 1) // size, 1), usernames=usernames,
                  filters={"user": user, "q": q, "du": du, "au": au})


# --------------------------------------------------------------------------- vérifier un employé


@router.get("/verifier")
def check_employee(request: Request, q: str = "", du: str = "", au: str = "", db: Session = Depends(get_db)):
    """Explique, jour par jour, pourquoi un employé apparaît (ou non) absent."""
    cfg, mapping = _mapping_ready(db)
    today = date.today()
    d_au = pointage.parse_day(au) or today
    d_du = pointage.parse_day(du) or (d_au - timedelta(days=13))
    if d_du > d_au:
        d_du, d_au = d_au, d_du
    d_du = max(d_du, d_au - timedelta(days=92))
    context = dict(q=q, du=d_du, au=d_au, mapping=mapping, results=None, others=[], error=None, calendar=[],
                   last=None, hhmm=pointage.hhmm, statuts=pointage.STATUTS)
    if mapping is not None and q.strip():
        engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
        try:
            context["results"], context["others"] = pointage.inspect_employee(engine, mapping, q, d_du, d_au)
            context["last"] = pointage.last_punch(engine, mapping)
            context["workdays"] = pointage.params_at(engine, mapping, d_au)["jours_ouvres"].split(",")
        except Exception as exc:
            context["error"] = friendly(exc)
        finally:
            engine.dispose()
        days, d = [], d_du
        while d <= min(d_au, today):
            days.append(d)
            d += timedelta(days=1)
        context["calendar"] = days
    return render(request, "admin/verifier.html", **context)
