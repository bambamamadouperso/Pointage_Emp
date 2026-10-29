"""Écran principal : suivi journalier des pointages (filtres dans l'URL, détail, exports Excel et PDF)."""
import io
import threading
import time
from datetime import date, datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, Response
from sqlalchemy.orm import Session

from .. import pointage
from ..auth import audit
from ..database import get_db
from ..joblog import write_log
from ..errors import friendly
from ..models import VALIDEE_DISPLAY, PointageConfig, SuiviSettings, User
from ..sync import make_engine
from ..web import back_url, redirect, render, require_login

router = APIRouter(prefix="/suivi", dependencies=[Depends(require_login)])

PAGE_SIZES = (50, 100, 250, 500)
MAX_DAYS = 366


def load_config(db: Session) -> tuple[Optional[PointageConfig], Optional[pointage.Mapping]]:
    cfg = db.query(PointageConfig).order_by(PointageConfig.id).first()
    if cfg is None or cfg.conn is None or cfg.installed_at is None:
        return cfg, None
    mapping = pointage.Mapping.from_json(cfg.data)
    if (cfg.sql_version or 0) < pointage.SQL_VERSION and _upgrade_allowed():
        # Nouvelle version de l'application : fonctions et vues PostgreSQL réinstallées une fois.
        # Une seule tentative à la fois (les autres pages continuent avec les objets actuels) ;
        # après un échec, nouvel essai dans 10 minutes seulement.
        try:
            engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
            try:
                pointage.install(engine, mapping, "mise à jour automatique")
                cfg.sql_version = pointage.SQL_VERSION
                db.commit()
                audit(None, "Calculs du pointage mis à jour automatiquement", f"version {pointage.SQL_VERSION}",
                      username="système")
            except Exception as exc:
                _upgrade_state["failed_at"] = time.monotonic()
                write_log("ERROR", f"Mise à jour automatique des calculs du pointage impossible : {friendly(exc)}. "
                                   "Nouvel essai dans 10 minutes, ou réinstallez-les depuis Administration → "
                                   "Source des pointages.")
            finally:
                engine.dispose()
        finally:
            _upgrade_state["lock"].release()
    return cfg, mapping


_upgrade_state = {"lock": threading.Lock(), "failed_at": None}


def _upgrade_allowed() -> bool:
    failed = _upgrade_state["failed_at"]
    if failed is not None and time.monotonic() - failed < 600:
        return False
    return _upgrade_state["lock"].acquire(blocking=False)


def default_categories(db: Session, mapping=None) -> list[str]:
    """Statuts du personnel cochés par défaut (Paramètres horaires → Affichage du suivi), ex. Cadre et Non Cadre ;
    aucun si la colonne du statut du personnel n'est pas configurée (sinon personne ne serait affiché)."""
    if mapping is None or not mapping.cat_col:
        return []
    raw = suivi_settings(db).default_categories
    return pointage.multi((raw if raw is not None else "Cadre, Non Cadre").split(","))


def chosen_categories(params, defaults: list[str]) -> list[str]:
    """Statuts choisis dans l'URL ; sans aucun choix encore fait (pas de « categorie_set »), ceux par défaut."""
    chosen = pointage.multi(params.getlist("categorie"))
    return chosen if chosen or "categorie_set" in params else list(defaults)


def _filters(request: Request, defaults: Optional[list[str]] = None) -> tuple[pointage.Filters, Optional[str]]:
    p = request.query_params
    today = date.today()
    day = pointage.parse_day(p.get("date"))
    du, au = pointage.parse_day(p.get("du")), pointage.parse_day(p.get("au"))
    warning = None
    if day:
        du = au = day
    du = du or au or today
    au = au or du
    if au < du:
        du, au = au, du
    if (au - du).days >= MAX_DAYS:
        au = du + timedelta(days=MAX_DAYS - 1)
        warning = f"Période limitée à {MAX_DAYS} jours : affichage du {du:%d/%m/%Y} au {au:%d/%m/%Y}."
    return pointage.Filters(
        du=du, au=au, q=p.get("q", "").strip(), service=pointage.multi(p.getlist("service")),
        statuts=[s for s in p.getlist("statut") if s in pointage.STATUTS],
        sort=p.get("sort", "nom") if p.get("sort", "nom") in pointage.SORTABLE else "nom",
        desc=p.get("dir") == "desc", team=p.get("equipe", ""), directs=p.get("directs") == "1" and bool(p.get("equipe")),
        population=p.get("pop", "") if p.get("pop") in ("liste", "hors") else "",
        # Badges hors liste : sans statut du personnel, le filtre par défaut les masquerait tous.
        categorie=chosen_categories(p, [] if p.get("pop") == "hors" else (defaults or [])),
        direction=pointage.multi(p.getlist("direction")), site=pointage.multi(p.getlist("site")),
        horaires=p.get("horaires", "") if p.get("horaires") in pointage.HORAIRES else "",
    ), warning


def objectif_minutes(engine, mapping, day) -> int:
    """Objectif de durée standard (paramètre historisé, 8h par défaut) en minutes."""
    try:
        h, m = pointage.params_at(engine, mapping, day)["objectif_duree"].split(":")
        return int(h) * 60 + int(m)
    except Exception:  # noqa: BLE001
        return 480


def suivi_settings(db: Session) -> SuiviSettings:
    row = db.query(SuiviSettings).order_by(SuiviSettings.id).first()
    if row is None:
        row = SuiviSettings(validee_display="rh")
        db.add(row)
        db.commit()
    return row


def validee_allowed(db: Session, role: Optional[str]) -> bool:
    """La règle d'affichage (Paramètres horaires) autorise-t-elle ce rôle à voir la durée standard ?"""
    mode = suivi_settings(db).validee_display
    if mode not in VALIDEE_DISPLAY:
        mode = "rh"  # ancien réglage « visible pour tous » : désormais réservé aux RH et aux administrateurs
    return mode == "rh" and role in ("admin", "rh")


def show_validee(request: Request, db: Session) -> tuple[bool, bool]:
    """(autorisée par la règle, affichée) : chacun peut en plus la masquer pour lui-même (bouton du suivi)."""
    allowed = validee_allowed(db, request.session.get("role"))
    return allowed, allowed and not request.session.get("hide_validee")


class ScopeError(Exception):
    pass


def dimension_filters(mapping, engine=None, scope_root=None) -> list[dict]:
    """Filtres d'organisation (direction, site) : libellé, valeurs proposées, colonne configurée ou non."""
    out = []
    for name, (prefix, label) in pointage.DIMENSIONS.items():
        configured = bool(mapping and getattr(mapping, f"{prefix}_col"))
        values = pointage.dimension_values(engine, mapping, name, scope_root) if engine is not None and configured else []
        out.append({"name": name, "label": label, "values": values, "configured": configured})
    return out


def apply_scope(request: Request, db: Session, engine, mapping: pointage.Mapping, f: pointage.Filters) -> Optional[str]:
    """Compte limité à son équipe : impose sa hiérarchie comme périmètre. Renvoie le libellé du périmètre."""
    user = db.query(User).filter(User.username == request.session.get("user", "")).one_or_none()
    if user is None or not user.team_only:
        return None
    if not mapping.hier_table:
        raise ScopeError("Votre compte est limité à votre équipe, mais la hiérarchie n'est pas configurée : "
                         "contactez un administrateur.")
    found = pointage.resolve_employee(engine, mapping, user.emp_matricule or "") if user.emp_matricule else None
    if found is None:
        raise ScopeError("Votre compte est limité à votre équipe mais n'est rattaché à aucun employé "
                         "(matricule manquant ou introuvable) : contactez un administrateur.")
    f.scope_root = found[0]
    if f.team and not pointage.in_scope(engine, mapping, f.scope_root, f.team):
        f.team = ""
    return found[1]


def _pending_requests(db: Session, engine, mapping, f, rows) -> dict:
    """Absences couvertes par une demande non validée (congé, télétravail, mission) : {(employé, jour): libellé}."""
    from .. import validations

    keys = sorted({r["emp_key"] for r in rows if r["statut"] == "ABSENT"})
    if not keys:
        return {}
    try:
        return validations.pending_map(engine, mapping, validations.get_settings(db), f.du, f.au, keys)
    except Exception as exc:  # noqa: BLE001 - indication facultative : ne doit pas empêcher l'affichage
        write_log("WARNING", f"Demandes en attente de validation non lues : {friendly(exc)}")
        return {}


@router.get("/affichage")
def display_toggle(request: Request, validee: str = "1"):
    """Masque ou réaffiche la durée standard pour l'utilisateur connecté (mémorisé pour sa session)."""
    request.session["hide_validee"] = validee == "0"
    return redirect(back_url(request, "/suivi"))


@router.get("")
def suivi(request: Request, db: Session = Depends(get_db)):
    cfg, mapping = load_config(db)
    f, warning = _filters(request, default_categories(db, mapping))
    validee_ok, validee_shown = show_validee(request, db)
    try:
        page = max(int(request.query_params.get("page", 1)), 1)
    except ValueError:
        page = 1
    try:
        size = int(request.query_params.get("size", 100))
    except ValueError:
        size = 100
    size = size if size in PAGE_SIZES else 100
    context = dict(f=f, warning=warning, today=date.today(), statuts=pointage.STATUTS, page=page, size=size, page_sizes=PAGE_SIZES,
                   hhmm=pointage.hhmm, configured=mapping is not None, data=None, error=None, services=[],
                   managers=[], categories=[], directions=[], dims=dimension_filters(mapping), has_direction=bool(mapping and mapping.dir_col), objectif_min=480, scope_label=None, population=None, last_punch=None, stale=False, has_hierarchy=bool(mapping and mapping.hier_table),
                   single_day=f.du == f.au, mode="jour" if f.du == f.au else "periode",
                   validee_ok=validee_ok, show_validee=validee_shown, horaires=pointage.HORAIRES, pending={}, hors_liste=[])
    if mapping is None:
        return render(request, "suivi.html", **context)
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        context["scope_label"] = apply_scope(request, db, engine, mapping, f)
        context["data"] = pointage.daily(engine, mapping, f, page, size)
        context["pending"] = _pending_requests(db, engine, mapping, f, context["data"]["rows"])
        # Badges sans fiche : pas pour un compte limité à son équipe, ni quand on cherche un employé ou les seuls
        # employés de la liste.
        if context["scope_label"] is None and not f.q and f.population != "liste":
            context["hors_liste"] = pointage.hors_liste(engine, mapping, f.du, f.au)
        context["services"] = pointage.services(engine, mapping, f.scope_root)
        context["managers"] = pointage.managers(engine, mapping, f.scope_root)
        context["categories"] = pointage.categories(engine, mapping, f.scope_root)
        f.categorie = pointage.match_values(f.categorie, context["categories"])
        context["directions"] = pointage.directions(engine, mapping, f.scope_root)
        context["dims"] = dimension_filters(mapping, engine, f.scope_root)
        context["objectif_min"] = objectif_minutes(engine, mapping, f.au)
        try:
            context["population"] = pointage.population(engine, mapping)
        except Exception as exc:  # contrôle facultatif : ne doit pas empêcher l'affichage
            write_log("WARNING", f"Contrôle de la liste des employés impossible : {friendly(exc)}")
        last = pointage.last_punch(engine, mapping)
        context["last_punch"] = last
        # Pointages plus anciens que la période affichée : synchronisation probablement arrêtée.
        context["stale"] = last is not None and last.date() < min(f.au, date.today()) - timedelta(days=1)
    except ScopeError as exc:
        context["error"] = str(exc)
    except Exception as exc:
        context["error"] = friendly(exc)
    finally:
        engine.dispose()
    total = context["data"]["total"] if context["data"] else 0
    context["pages"] = max((total + size - 1) // size, 1)
    return render(request, "suivi.html", **context)


@router.get("/detail", response_class=HTMLResponse)
def detail(request: Request, key: str, jour: str, db: Session = Depends(get_db)):
    cfg, mapping = load_config(db)
    day = pointage.parse_day(jour)
    if mapping is None or day is None:
        return HTMLResponse('<p class="empty">Détail indisponible.</p>', status_code=400)
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        f = pointage.Filters(du=day, au=day)
        apply_scope(request, db, engine, mapping, f)
        if not pointage.in_scope(engine, mapping, f.scope_root, key):
            raise ScopeError("Cet employé ne fait pas partie de votre équipe.")
        columns, rows = pointage.detail(engine, mapping, key, day)
        error = None
    except ScopeError as exc:
        columns, rows, error = [], [], str(exc)
    except Exception as exc:
        columns, rows, error = [], [], friendly(exc)
    finally:
        engine.dispose()
    return render(request, "_suivi_detail.html", columns=columns, rows=rows, error=error, day=day, key=key)


def _export_data(request: Request, db: Session):
    """Lignes et contexte communs aux exports Excel et PDF (ou une réponse d'erreur)."""
    cfg, mapping = load_config(db)
    if mapping is None:
        return Response("Module de pointage non configuré.", status_code=400)
    f, _ = _filters(request, default_categories(db, mapping))
    _, with_validee = show_validee(request, db)
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        apply_scope(request, db, engine, mapping, f)
        rows = pointage.export_rows(engine, mapping, f)
        objectif = objectif_minutes(engine, mapping, f.au)
    except ScopeError as exc:
        return Response(str(exc), status_code=403, media_type="text/plain; charset=utf-8")
    except Exception as exc:
        return Response(f"Export impossible : {friendly(exc)}", status_code=400, media_type="text/plain; charset=utf-8")
    finally:
        engine.dispose()
    return f, rows, objectif, with_validee


def _averages(rows) -> dict:
    """Ligne « Moyenne » (présents au bureau), comme à l'écran."""
    def avg(values):
        values = [v for v in values if v is not None]
        return sum(values, timedelta()) / len(values) if values else None

    def clock(td):
        return (datetime.min + td).time().replace(microsecond=0) if td is not None else None

    bureau = [r for r in rows if r["statut"] in ("A_L_HEURE", "RETARD", "INCOMPLET")]
    complete = [r for r in bureau if r["statut"] != "INCOMPLET"]
    bureau = [r for r in bureau if not r.get("poste")]  # heures moyennes : horaire de bureau seulement
    first = avg([timedelta(hours=r["premier_pointage"].hour, minutes=r["premier_pointage"].minute,
                           seconds=r["premier_pointage"].second) for r in bureau if r["premier_pointage"]])
    last = avg([timedelta(hours=r["dernier_pointage"].hour, minutes=r["dernier_pointage"].minute,
                          seconds=r["dernier_pointage"].second)
                for r in bureau if r["dernier_pointage"] and r["nb_pointages"] >= 2 and r["statut"] != "INCOMPLET"])
    return {"premier": clock(first), "dernier": clock(last), "validee": avg([r["duree_validee"] for r in complete]),
            "effective": avg([r["duree_effective"] for r in complete])}


def _filter_items(request: Request, f: pointage.Filters, objectif: int, with_validee: bool, count: int) -> list:
    return [
        ("Du", f.du.strftime("%d/%m/%Y")), ("Au", f.au.strftime("%d/%m/%Y")), ("Recherche", f.q or "—"),
        ("Objectif de durée standard", f"{objectif // 60}h{objectif % 60:02d} (vert si atteint, rouge sinon)"),
        ("Service", ", ".join(f.service) or "Tous"), ("Statut du personnel", ", ".join(f.categorie) or "Tous"),
        ("Direction", ", ".join(f.direction) or "Toutes"), ("Site", ", ".join(f.site) or "Tous"),
        ("Personnes", {"liste": "employés de la liste", "hors": "hors liste"}.get(f.population, "toutes")),
        ("Équipe", (f.team + (" (directs)" if f.directs else "")) if f.team else "Toutes"),
        ("Périmètre", "équipe du compte" if f.scope_root is not None else "tout le personnel"),
        ("Statuts", ", ".join(pointage.STATUTS[s][0] for s in f.statuts) or "Tous"),
        ("Horaires", pointage.HORAIRES.get(f.horaires, "Tous (bureau et postés)")),
        ("Durée standard", "affichée" if with_validee else "masquée"),
        ("Lignes", count), ("Exporté par", request.session.get("user", "")),
    ]


def _export_name(f: pointage.Filters, ext: str) -> str:
    return f"suivi_{f.du:%Y%m%d}" + ("" if f.du == f.au else f"_{f.au:%Y%m%d}") + f".{ext}"


@router.get("/export.xlsx")
def export(request: Request, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = _export_data(request, db)
    if isinstance(data, Response):
        return data
    f, rows, objectif, with_validee = data

    wb = Workbook()
    ws = wb.active
    ws.title = "Suivi journalier"
    headers = ["Date", "Matricule", "Nom", "Prénom", "Service", "Statut du personnel", "1er pointage", "Statut", "Dernier pointage",
               "Nb pointages", "Durée standard", "Durée effective", "Durée standard (min)", "Durée effective (min)",
               "Responsable", "Dans la liste des employés", "Poste planifié", "Direction", "Site", "Allaitement",
               "Site du 1er pointage", "Site du dernier pointage"]
    ws.append(headers)
    fills = {"A_L_HEURE": "DCFCE7", "RETARD": "FFEDD5", "ABSENT": "FEE2E2", "INCOMPLET": "E5E7EB", "NON_OUVRE": "E0F2FE",
             "CONGE_ANNUEL": "E4F5D3", "CONGE_EXCEP": "E4F5D3", "CONGE_MATERNITE": "FCE7F3",
             "TELETRAVAIL": "CDEEE7", "TERRAIN": "DBEAFE",
             "ARRET_MALADIE": "E7E5E4", "REPOS": "E0F2FE", "FORMATION": "EDE9FE",
             "MISSION": "F3E8FF"}
    for r in rows:
        ws.append([
            r["jour"], r["matricule"], r["nom"], r["prenom"], r["service"], r["categorie"],
            r["premier_pointage"].time() if r["premier_pointage"] else None, r["statut_libelle"],
            r["dernier_pointage"].time() if r["dernier_pointage"] else None,
            r["nb_pointages"], r["duree_validee"], r["duree_effective"],
            float(r["duree_validee_min"]) if r["duree_validee_min"] is not None else None,
            float(r["duree_effective_min"]) if r["duree_effective_min"] is not None else None,
            r["responsable"], "Non" if r["hors_liste"] else "Oui", r.get("poste"), r.get("direction"), r.get("site"),
            r.get("allaitement"), r.get("site_premier"), r.get("site_dernier") if r["nb_pointages"] > 1 else None,
        ])
        row = ws.max_row
        ws.cell(row, 1).number_format = "DD/MM/YYYY"
        for col in (7, 9):
            ws.cell(row, col).number_format = "HH:MM"
        for col in (11, 12, 20):
            ws.cell(row, col).number_format = "[H]:MM"
        ws.cell(row, 8).fill = PatternFill("solid", fgColor=fills.get(r["statut"], "FFFFFF"))
        for col, key in ((11, "duree_validee_min"), (12, "duree_effective_min")):  # objectif : vert atteint, rouge sinon
            if r[key] is not None:
                ok = float(r[key]) >= objectif
                ws.cell(row, col).fill = PatternFill("solid", fgColor="C6EFCE" if ok else "FFC7CE")
                ws.cell(row, col).font = Font(color="006100" if ok else "9C0006", bold=True)
    for i, h in enumerate(headers, 1):
        ws.cell(1, i).font = Font(bold=True)
        ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 2)
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["O"].width = 26
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    moy = _averages(rows)
    ws.append([])
    ws.append(["Moyenne", "", "présents au bureau", "", "", "", moy["premier"], "", moy["dernier"], "",
               moy["validee"], moy["effective"]])
    row = ws.max_row
    for col in range(1, 13):
        ws.cell(row, col).font = Font(bold=True)
    for col in (7, 9):
        ws.cell(row, col).number_format = "HH:MM"
    for col in (11, 12):
        ws.cell(row, col).number_format = "[H]:MM"
    if not with_validee:
        # Durée standard masquée (règle d'affichage ou choix de l'utilisateur) : colonnes retirées de l'export.
        for col in (13, 11):
            ws.delete_cols(col)
        for i, cell in enumerate(ws[1], 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(cell.value or "")) + 2)
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["E"].width = 22
        ws.column_dimensions["M"].width = 26
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{len(rows) + 1}"
    info = wb.create_sheet("Filtres")
    for label, value in _filter_items(request, f, objectif, with_validee, len(rows)):
        info.append([label, value])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_export_name(f, "xlsx")}"'})


PDF_MAX_ROWS = 5000


PPTX_MAX_ROWS = 300  # PowerPoint : au-delà, l'export Excel ou PDF


def _suivi_content(request: Request, db: Session, data: tuple, max_rows: int) -> dict:
    """Contenu commun aux exports PDF et PowerPoint du suivi : indicateurs et tableau (statuts et durées en couleur)."""
    from .. import pdf

    f, rows, objectif, with_validee = data
    hhmm = pointage.hhmm
    period = f"du {f.du:%d/%m/%Y}" + ("" if f.du == f.au else f" au {f.au:%d/%m/%Y}")
    _, mapping = load_config(db)
    has_dir = bool(mapping and mapping.dir_col)  # colonne Direction quand elle est configurée
    headers = ["Date", "Matricule", "Nom", *(["Direction"] if has_dir else []), "Service", "Statut du personnel",
               "1er pointage", "Statut", "Dernier", "Nb", *(["Durée standard"] if with_validee else []), "Durée effective",
               "Responsable"]
    widths = [21, 16, 34, 28, *([28] if has_dir else []), 20, 15, 31, 15, 8, *([17] if with_validee else []), 17, 30]
    align = ["C", "L", "L", "L", *(["L"] if has_dir else []), "L", "C", "L", "C", "C", *(["C"] if with_validee else []),
             "C", "L"]
    status_col, first_duration = 6 + has_dir, 9 + has_dir
    count = lambda *codes: sum(1 for r in rows if r["statut"] in codes)  # noqa: E731
    moy = _averages(rows)
    kpis = [
        ("À l'heure", count("A_L_HEURE"), "journées", pdf.GREEN),
        ("En retard", count("RETARD"), "journées", pdf.ORANGE),
        ("Absents", count("ABSENT"), "sans justificatif", pdf.RED),
        ("Pointages incomplets", count("INCOMPLET"), "un seul pointage", pdf.GREY),
        ("Autres situations", len(rows) - count("A_L_HEURE", "RETARD", "ABSENT", "INCOMPLET"),
         "congés, télétravail, terrain, missions...", pdf.BLUE),
        ("Durée effective moyenne", hhmm(moy["effective"]) if moy["effective"] else "—",
         f"objectif {objectif // 60}h{objectif % 60:02d} · présents au bureau", pdf.TEAL),
    ]
    body, pills, colors = [], {}, {}
    for i, r in enumerate(rows[:max_rows]):
        colors[(i, 1)] = colors[(i, len(headers) - 1)] = pdf.MUTED
        body.append([r["jour"].strftime("%d/%m/%Y"), r["matricule"], " ".join(filter(None, [r["nom"], r["prenom"]])),
                     *([r.get("direction")] if has_dir else []), r["service"], r["categorie"],
                     hhmm(r["premier_pointage"]) if r["premier_pointage"] else "",
                     r["statut_libelle"] + (" (allaitement)" if r.get("allaitement") else ""),
                     hhmm(r["dernier_pointage"]) if r["dernier_pointage"] else "", r["nb_pointages"] or "",
                     *([hhmm(r["duree_validee"]) if r["duree_validee"] is not None else ""] if with_validee else []),
                     hhmm(r["duree_effective"]) if r["duree_effective"] is not None else "", r["responsable"]])
        pills[(i, status_col)] = pdf.STATUS_PILLS.get(r["statut"], pdf.STATUS_PILLS["INCOMPLET"])
        for col, key in ((first_duration, "duree_validee_min") if with_validee else (None, None),
                         (first_duration + with_validee, "duree_effective_min")):
            if col is not None and r[key] is not None:  # objectif : vert atteint, rouge sinon
                pills[(i, col)] = pdf.GOOD if float(r[key]) >= objectif else pdf.BAD
    body.append(["Moyenne", "", "présents au bureau", "", *([""] if has_dir else []), "",
                 hhmm(moy["premier"]) if moy["premier"] else "", "",
                 hhmm(moy["dernier"]) if moy["dernier"] else "", "",
                 *([hhmm(moy["validee"]) if moy["validee"] else ""] if with_validee else []),
                 hhmm(moy["effective"]) if moy["effective"] else "", ""])
    return {"title": "Suivi journalier des pointages", "subtitle": f"{period} · {len(rows)} ligne(s)",
            "user": request.session.get("user", ""),
            "filters": [(k, v) for k, v in _filter_items(request, f, objectif, with_validee, len(rows))
                        if k not in ("Du", "Au", "Lignes", "Exporté par") and v not in ("—", "Tous", "Toutes", "toutes")],
            "kpis": kpis, "headers": headers, "rows": body, "widths": widths, "align": align, "pills": pills,
            "colors": colors, "hidden": max(len(rows) - max_rows, 0), "f": f}


@router.get("/export.pdf")
def export_pdf(request: Request, db: Session = Depends(get_db)):
    from .. import pdf

    data = _export_data(request, db)
    if isinstance(data, Response):
        return data
    c = _suivi_content(request, db, data, PDF_MAX_ROWS)
    doc = pdf.Document(c["title"], c["subtitle"], c["user"])
    doc.filters(c["filters"])
    doc.kpis(c["kpis"])
    doc.table(c["headers"], c["rows"], c["widths"], c["align"], c["pills"], c["colors"], bold_last=True)
    if c["hidden"]:
        doc.section("Export tronqué", f"{c['hidden']} ligne(s) au-delà des {PDF_MAX_ROWS} premières ne figurent pas "
                                      "dans le PDF : affinez les filtres ou utilisez l'export Excel.")
    return Response(doc.output_bytes(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_export_name(c["f"], "pdf")}"'})


@router.get("/export.pptx")
def export_pptx(request: Request, db: Session = Depends(get_db)):
    from .. import ppt

    data = _export_data(request, db)
    if isinstance(data, Response):
        return data
    c = _suivi_content(request, db, data, PPTX_MAX_ROWS)
    deck = ppt.Deck(c["title"], c["subtitle"], c["user"], c["filters"])
    deck.kpis("Indicateurs du suivi", c["kpis"])
    note = c["subtitle"] + (f" — {PPTX_MAX_ROWS} premières lignes ; {c['hidden']} de plus dans l'export Excel ou PDF"
                            if c["hidden"] else "")
    deck.table("Détail des journées", note, c["headers"], c["rows"], c["widths"], c["align"], c["pills"], c["colors"])
    return Response(deck.output_bytes(), media_type=ppt.MEDIA_TYPE,
                    headers={"Content-Disposition": f'attachment; filename="{_export_name(c["f"], "pptx")}"'})
