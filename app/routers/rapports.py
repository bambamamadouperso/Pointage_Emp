"""Rapports d'analyse : indicateurs, tendances, habitudes d'arrivée, services, employés à suivre, exports Excel et PDF."""
import io
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session

from .. import auth, pointage, rapports
from ..database import get_db
from ..errors import friendly
from ..sync import make_engine
from ..web import render, require_login
from .suivi import (ScopeError, apply_scope, chosen_categories, default_categories, dimension_filters, load_config,
                    validee_allowed)

router = APIRouter(prefix="/rapports", dependencies=[Depends(require_login)])

SORTS = {"absences": "Absences", "retards": "Retards", "retard_min": "Minutes de retard", "incomplets": "Pointages partiels",
         "presence": "Taux de présence (croissant)", "heures": "Heures standard", "nom": "Nom"}


# Durée standard (heures standard, moyenne par jour) : réservée aux RH et aux administrateurs (Paramètres horaires).
STD_HEADERS = {"Heures standard", "Moyenne heures / jour au bureau", "Moy. / jour au bureau"}


def _sorts(std: bool) -> dict:
    return SORTS if std else {k: v for k, v in SORTS.items() if k != "heures"}


def _filters(request: Request, defaults: Optional[list[str]] = None) -> tuple[pointage.Filters, str]:
    p = request.query_params
    presets = p.getlist("p")
    preset = "perso" if "perso" in presets else (presets[0] if presets else "30j")
    du, au = pointage.parse_day(p.get("du")), pointage.parse_day(p.get("au"))
    if du and au and preset == "perso":
        du, au = min(du, au), max(du, au)
        au = min(au, du + timedelta(days=365))
    else:
        preset = preset if preset in dict(rapports.PRESETS) else "30j"
        du, au = rapports.period(preset)
    team = p.get("equipe", "")
    return pointage.Filters(du=du, au=au, service=pointage.multi(p.getlist("service")), team=team, categorie=chosen_categories(p, defaults or []),
                            direction=pointage.multi(p.getlist("direction")), site=pointage.multi(p.getlist("site")),
                            directs=p.get("directs") == "1" and bool(team),
                            horaires=p.get("horaires", "") if p.get("horaires") in pointage.HORAIRES else ""), preset


@router.get("")
def page(request: Request, db: Session = Depends(get_db)):
    cfg, mapping = load_config(db)
    f, preset = _filters(request, default_categories(db, mapping))
    std = validee_allowed(db, request.session.get("role"))
    sorts = _sorts(std)
    sort = request.query_params.get("tri", "absences")
    sort = sort if sort in sorts else "absences"
    context = dict(f=f, preset=preset, presets=rapports.PRESETS, sorts=sorts, sort=sort, configured=mapping is not None, std=std,
                   has_direction=bool(mapping and mapping.dir_col),
                   data=None, error=None, services=[], managers=[], categories=[], directions=[], scope_label=None,
                   hhmm=pointage.hhmm, dims=dimension_filters(mapping),
                   series=rapports.SERIES, has_hierarchy=bool(mapping and mapping.hier_table), today=date.today(),
                   horaires=pointage.HORAIRES)
    if mapping is None:
        return render(request, "rapports.html", **context)
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        context["scope_label"] = apply_scope(request, db, engine, mapping, f)
        data = rapports.build(engine, mapping, f, sort)
        params = pointage.params_at(engine, mapping, f.au)
        data["trend"] = rapports.trend_chart(data["jours"])
        data["arrival"] = rapports.arrival_chart(data["arrivees"], params["seuil_retard"], params["debut_journee"])
        context["data"] = data
        context["services"] = pointage.services(engine, mapping, f.scope_root)
        context["managers"] = pointage.managers(engine, mapping, f.scope_root)
        context["categories"] = pointage.categories(engine, mapping, f.scope_root)
        f.categorie = pointage.match_values(f.categorie, context["categories"])
        context["directions"] = pointage.directions(engine, mapping, f.scope_root)
        context["dims"] = dimension_filters(mapping, engine, f.scope_root)
    except ScopeError as exc:
        context["error"] = str(exc)
    except Exception as exc:
        context["error"] = friendly(exc)
    finally:
        engine.dispose()
    return render(request, "rapports.html", **context)


def _export_data(request: Request, db: Session):
    """Données communes aux exports Excel et PDF (ou une réponse d'erreur)."""
    cfg, mapping = load_config(db)
    if mapping is None:
        return Response("Module de pointage non configuré.", status_code=400)
    f, _ = _filters(request, default_categories(db, mapping))
    engine = make_engine(cfg.conn, **pointage.WEB_LIMITS)
    try:
        apply_scope(request, db, engine, mapping, f)
        sort = request.query_params.get("tri", "absences")
        sort = sort if sort in _sorts(validee_allowed(db, request.session.get("role"))) else "absences"
        data = rapports.build(engine, mapping, f, sort, compare=False)
    except ScopeError as exc:
        return Response(str(exc), status_code=403, media_type="text/plain; charset=utf-8")
    except Exception as exc:
        return Response(f"Export impossible : {friendly(exc)}", status_code=400, media_type="text/plain; charset=utf-8")
    finally:
        engine.dispose()
    return f, data


def _filter_items(request: Request, f: pointage.Filters) -> list:
    return [("Du", f.du.strftime("%d/%m/%Y")), ("Au", f.au.strftime("%d/%m/%Y")), ("Service", ", ".join(f.service) or "Tous"),
            ("Statut du personnel", ", ".join(f.categorie) or "Tous"),
            ("Direction", ", ".join(f.direction) or "Toutes"), ("Site", ", ".join(f.site) or "Tous"),
            ("Équipe", (f.team + (" (directs)" if f.directs else "")) if f.team else "Toutes"),
            ("Horaires", pointage.HORAIRES.get(f.horaires, "Tous (bureau et postés)")),
            ("Exporté par", request.session.get("user", ""))]


def _without_std(title: str, headers: list, rows: list, pct_cols: set) -> tuple:
    """Feuille sans les colonnes de durée standard (numéros des colonnes en pourcentage recalculés)."""
    keep = [i for i, h in enumerate(headers) if h not in STD_HEADERS]
    if len(keep) == len(headers):
        return title, headers, rows, pct_cols
    renum = {old + 1: new + 1 for new, old in enumerate(keep)}
    return (title, [headers[i] for i in keep], [[row[i] for i in keep] for row in rows],
            {renum[c] for c in pct_cols if c in renum})


@router.get("/export.xlsx")
def export(request: Request, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    result = _export_data(request, db)
    if isinstance(result, Response):
        return result
    f, data = result

    def pct(v):
        return round(v / 100, 4) if v is not None else None

    def hours(v):
        return round(float(v), 2) if v is not None else None

    wb = Workbook()
    sheets = [
        ("Employés", ["Matricule", "Nom", "Prénom", "Service", "Statut du personnel", "Responsable", "Agent terrain", "Jours attendus",
                      "Présences", "Taux de présence", "À l'heure", "Retards", "Minutes de retard", "Taux de ponctualité",
                      "Absences", "dont lundi/vendredi", "Pointages incomplets", "Congés", "Télétravail", "Sur le terrain",
                      "Heures standard", "Direction", "Site"],
         [[e["matricule"], e["nom"], e["prenom"], e["service"], e["categorie"], e["responsable"], "Oui" if e["terrain"] else "Non",
           e["attendus"], e["presents"], pct(e["taux_presence"]), e["a_l_heure"], e["retards"], float(e["retard_min_total"]),
           pct(e["taux_ponctualite"]), e["absences"], e["absences_lun_ven"], e["incomplets"], e["conges"], e["teletravail"],
           e["terrain_jours"], hours(e["heures_validees"]), e.get("direction"), e.get("site")]
          for e in data["employes"]], {10, 14}),
        ("Services", ["Service", "Employés", "Taux de présence", "Taux d'absentéisme", "Taux de ponctualité", "Retards",
                      "Minutes de retard", "Absences", "Pointages incomplets", "Heures standard",
                      "Moyenne heures / jour au bureau"],
         [[s["service"], s["employes"], pct(s["taux_presence"]), pct(s["taux_absence"]), pct(s["taux_ponctualite"]),
           s["retards"], float(s["retard_min_total"]), s["absences"], s["incomplets"], hours(s["heures_validees"]),
           hours(s["heures_moy_bureau"])] for s in data["services"]], {3, 4, 5}),
        *([("Statuts du personnel", ["Statut du personnel", "Employés", "Taux de présence", "Taux d'absentéisme", "Taux de ponctualité",
                            "Retards", "Minutes de retard", "Absences", "Pointages incomplets", "Heures standard",
                            "Moyenne heures / jour au bureau"],
             [[g["groupe"], g["employes"], pct(g["taux_presence"]), pct(g["taux_absence"]), pct(g["taux_ponctualite"]),
               g["retards"], float(g["retard_min_total"]), g["absences"], g["incomplets"], hours(g["heures_validees"]),
               hours(g["heures_moy_bureau"])] for g in data["categories"]], {3, 4, 5})] if data["categories"] else []),
        *([("Directions", ["Direction", "Employés", "Taux de présence", "Taux d'absentéisme", "Taux de ponctualité",
                            "Retards", "Minutes de retard", "Absences", "Pointages incomplets", "Heures standard",
                            "Moyenne heures / jour au bureau"],
             [[g["groupe"], g["employes"], pct(g["taux_presence"]), pct(g["taux_absence"]), pct(g["taux_ponctualite"]),
               g["retards"], float(g["retard_min_total"]), g["absences"], g["incomplets"], hours(g["heures_validees"]),
               hours(g["heures_moy_bureau"])] for g in data["directions"]], {3, 4, 5})] if data.get("directions") else []),
        *([("Sites", ["Site", "Employés", "Taux de présence", "Taux d'absentéisme", "Taux de ponctualité",
                      "Retards", "Minutes de retard", "Absences", "Pointages incomplets", "Heures standard",
                      "Moyenne heures / jour au bureau"],
             [[g["groupe"], g["employes"], pct(g["taux_presence"]), pct(g["taux_absence"]), pct(g["taux_ponctualite"]),
               g["retards"], float(g["retard_min_total"]), g["absences"], g["incomplets"], hours(g["heures_validees"]),
               hours(g["heures_moy_bureau"])] for g in data["sites"]], {3, 4, 5})] if data.get("sites") else []),
        ("Par jour" if data["granularite"] == "jour" else "Par semaine",
         ["Date"] + [label for _, label, _ in rapports.SERIES],
         [[j["jour"]] + [j[k] for k, _, _ in rapports.SERIES] for j in data["jours"]], set()),
        # Employés à suivre : réservé aux administrateurs.
        *([("Alertes", ["Motif", "Matricule", "Nom", "Service", "Responsable", "Détail"],
            [[a["motif"], a["matricule"], a["nom"], a["service"], a["responsable"], a["detail"]] for a in data["alertes"]],
            set())] if auth.has_role(request.session.get("role"), "admin") else []),
    ]
    if not validee_allowed(db, request.session.get("role")):
        sheets = [_without_std(*sheet) for sheet in sheets]
    for i, (title, headers, rows, pct_cols) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for row in rows:
            ws.append(row)
            if title.startswith("Par"):
                ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
            for col in pct_cols:
                ws.cell(ws.max_row, col).number_format = "0.0%"
        for col, h in enumerate(headers, 1):
            ws.cell(1, col).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(col)].width = max(12, min(len(h) + 2, 34))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Filtres")
    for row in _filter_items(request, f):
        info.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    name = f"rapport_{f.du:%Y%m%d}_{f.au:%Y%m%d}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _report_content(request: Request, db: Session, f: pointage.Filters, data: dict) -> dict:
    """Contenu du rapport commun aux exports PDF et PowerPoint : indicateurs, tableaux, couleurs des cellules.
    Respecte les droits : durée standard (RH, administrateurs), employés à suivre (administrateurs)."""
    from .. import pdf

    k = data["kpi"]
    role = request.session.get("role")
    std, admin = validee_allowed(db, role), auth.has_role(role, "admin")
    _, mapping = load_config(db)
    has_dir = bool(mapping and mapping.dir_col)
    cut = None if std else -1  # durée standard (dernières colonnes) réservée aux RH et aux administrateurs

    def pct(v):
        return f"{float(v):.1f} %".replace(".", ",") if v is not None else "—"

    def hours(v):
        return pointage.hhmm(timedelta(hours=float(v))) if v is not None else "—"

    def num(v):
        return f"{int(v or 0):,}".replace(",", " ")

    out = {"title": "Rapport de présence", "subtitle": f"du {f.du:%d/%m/%Y} au {f.au:%d/%m/%Y} · employés de la liste",
           "user": request.session.get("user", ""), "empty": not k.get("attendus"), "admin": admin, "std": std,
           "filters": [(label, value) for label, value in _filter_items(request, f)
                       if label not in ("Du", "Au", "Exporté par") and value not in ("Tous", "Toutes")],
           "tables": [], "name": f"rapport_{f.du:%Y%m%d}_{f.au:%Y%m%d}"}
    if out["empty"]:
        return out
    out["kpis"] = [
        ("Taux de présence", pct(k["taux_presence"]), f"{num(k['presents'])} journées sur {num(k['disponibles'])} (hors congés)", pdf.GREEN),
        ("Taux d'absentéisme", pct(k["taux_absence"]), f"{num(k['absences'])} journée(s) d'absence", pdf.RED),
        ("Taux de ponctualité", pct(k["taux_ponctualite"]), "à l'heure / (à l'heure + retards)", pdf.BLUE),
        ("Retards", num(k["retards"]), f"{pointage.hhmm(timedelta(minutes=float(k['retard_min_total'] or 0)))} cumulées", pdf.ORANGE),
        ("Heures standard", f"{num(round(float(k['heures_validees'] or 0)))} h",
         f"{hours(k['heures_moy_bureau'])} par jour au bureau" if k.get("heures_moy_bureau") else "", pdf.TEAL) if std else
        ("Journées attendues", num(k["attendus"]), f"{num(k['employes'])} employé(s)", pdf.TEAL),
        ("Pointages partiels", num(k["incomplets"]), "journées avec un seul pointage", pdf.GREY),
    ]
    out["volumes"] = [("Employés", num(k["employes"])), ("Journées attendues", num(k["attendus"])),
                      ("En congé", num(k["conges"])), ("Arrêt maladie", num(k["maladies"])),
                      ("Télétravail", num(k["teletravail"])), ("Sur le terrain", num(k["terrain"])),
                      *([("En mission", num(k["missions"]))] if k.get("missions") else [])]

    def rates(row, i, cols, pills):
        """Pastilles des taux : présence et ponctualité (plus haut = mieux), absentéisme (plus bas = mieux)."""
        for col, key, good, warn, up in cols:
            pill = pdf.rate_pill(row[key], good, warn, up)
            if pill:
                pills[(i, col)] = pill

    group_headers = ["Employés", "Présence", "Absentéisme", "Ponctualité", "Retards", "Min. de retard", "Absences",
                     "Pointages partiels", "Heures standard", "Moy. / jour au bureau"]
    group_widths = [16, 18, 20, 20, 15, 20, 16, 20, 20, 26]
    if not std:
        group_headers, group_widths = group_headers[:-2], group_widths[:-2]
    for title, key, field in (("Par service", "services", "service"), ("Par statut du personnel", "categories", "groupe"),
                              ("Par direction", "directions", "groupe"), ("Par site", "sites", "groupe")):
        groups = data.get(key) or []
        if not groups:
            continue
        pills, colors = {}, {}
        for i, g in enumerate(groups):
            rates(g, i, ((2, "taux_presence", 95, 85, True), (3, "taux_absence", 95, 90, False),
                         (4, "taux_ponctualite", 95, 85, True)), pills)
            if g["absences"]:
                colors[(i, 7)] = pdf.RED
        out["tables"].append({
            "title": title, "note": f"{len(groups)} groupe(s), du plus fort au plus faible absentéisme",
            "headers": [title[4:].capitalize()] + group_headers,
            "rows": [[g[field], num(g["employes"]), pct(g["taux_presence"]), pct(g["taux_absence"]),
                      pct(g["taux_ponctualite"]), num(g["retards"]), num(g["retard_min_total"]), num(g["absences"]),
                      num(g["incomplets"]), hours(g["heures_validees"]), hours(g["heures_moy_bureau"])][:None if std else -2]
                     for g in groups],
            "widths": [50] + group_widths, "align": ["L"] + ["R"] * len(group_headers), "pills": pills, "colors": colors})

    if data["alertes"] and admin:  # réservé aux administrateurs
        out["tables"].append({
            "title": "Employés à suivre", "note": f"{len(data['alertes'])} alerte(s)",
            "headers": ["Motif", "Matricule", "Nom", "Service", "Responsable", "Détail"],
            "rows": [[a["motif"], a["matricule"], a["nom"], a["service"], a["responsable"], a["detail"]] for a in data["alertes"]],
            "widths": [50, 18, 44, 36, 40, 80], "align": None,
            "pills": {(i, 0): pdf.BAD if a["gravite"] >= 3 else pdf.WARN if a["gravite"] == 2 else pdf.STATUS_PILLS["INCOMPLET"]
                      for i, a in enumerate(data["alertes"])},
            "colors": {(i, 1): pdf.MUTED for i in range(len(data["alertes"]))}})

    pills, colors = {}, {}
    for i, e in enumerate(data["employes"]):
        rates(e, i, ((6, "taux_presence", 95, 85, True), (9, "taux_ponctualite", 95, 85, True)), pills)
        colors[(i, 0)] = pdf.MUTED
        if e["absences"]:
            colors[(i, 10)] = pdf.RED
        if e["retards"]:
            colors[(i, 7)] = pdf.ORANGE
    # Colonne Direction (si configurée) avant le service : les numéros de colonnes colorées sont décalés d'autant.
    if has_dir:
        pills = {(i, j + (j >= 2)): v for (i, j), v in pills.items()}
        colors = {(i, j + (j >= 2)): v for (i, j), v in colors.items()}
    out["tables"].append({
        "title": "Par employé", "employees": True,
        "note": f"{len(data['employes'])} employé(s), tri : "
                f"{_sorts(std).get(request.query_params.get('tri', 'absences'), SORTS['absences']).lower()}",
        "headers": ["Matricule", "Nom", *(["Direction"] if has_dir else []), "Service", "Statut du personnel", "Attendus",
                    "Présences", "Présence", "Retards", "Min. de retard", "Ponctualité", "Absences", "Pointages partiels",
                    "Congés", "Heures standard"][:cut],
        "rows": [[e["matricule"], " ".join(filter(None, [e["nom"], e["prenom"]])),
                  *([e.get("direction")] if has_dir else []), e["service"], e["categorie"], num(e["attendus"]),
                  num(e["presents"]), pct(e["taux_presence"]), num(e["retards"]), num(e["retard_min_total"]),
                  pct(e["taux_ponctualite"]), num(e["absences"]), num(e["incomplets"]), num(e["conges"]),
                  hours(e["heures_validees"])][:cut] for e in data["employes"]],
        "widths": [18, 40, 30, *([28] if has_dir else []), 22, 16, 17, 18, 15, 18, 19, 16, 14, 14, 19][:cut],
        "align": (["L"] * (4 + has_dir) + ["R"] * 10)[:cut], "pills": pills, "colors": colors})
    return out


@router.get("/export.pdf")
def export_pdf(request: Request, db: Session = Depends(get_db)):
    from .. import pdf

    result = _export_data(request, db)
    if isinstance(result, Response):
        return result
    f, data = result
    c = _report_content(request, db, f, data)
    doc = pdf.Document(c["title"], c["subtitle"], c["user"])
    doc.filters(c["filters"])
    if c["empty"]:
        doc.section("Aucune journée à analyser sur cette période avec ces filtres.")
    else:
        doc.kpis(c["kpis"])
        doc.filters(c["volumes"])
        for t in c["tables"]:
            doc.section(t["title"], t["note"])
            doc.table(t["headers"], t["rows"], t["widths"], t["align"], t["pills"], t["colors"])
    return Response(doc.output_bytes(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{c["name"]}.pdf"'})


@router.get("/export.pptx")
def export_pptx(request: Request, db: Session = Depends(get_db)):
    from .. import ppt

    result = _export_data(request, db)
    if isinstance(result, Response):
        return result
    f, data = result
    c = _report_content(request, db, f, data)
    deck = ppt.Deck(c["title"], c["subtitle"], c["user"], c["filters"])
    if c["empty"]:
        deck.message("Aucune journée à analyser sur cette période avec ces filtres.")
    else:
        deck.kpis("Indicateurs clés", c["kpis"], c["volumes"])
        # Graphiques natifs PowerPoint (modifiables) : répartition des statuts dans le temps, jour de la semaine.
        if data["jours"]:
            fmt = (lambda d: f"{d:%d/%m}") if data["granularite"] == "jour" else (lambda d: f"sem. {d:%d/%m}")
            deck.chart(f"Présence {'par jour' if data['granularite'] == 'jour' else 'par semaine'}",
                       "Répartition des journées-employé par statut.",
                       [fmt(j["jour"]) for j in data["jours"]],
                       [(label, [j[key] for j in data["jours"]], color) for key, label, color in rapports.SERIES],
                       kind="stacked")
        if data["semaine"]:
            deck.chart("Selon le jour de la semaine", "Taux de présence, d'absentéisme et de retard (%).",
                       [w["jour"] for w in data["semaine"]],
                       [("Taux de présence", [round(w["taux_presence"] or 0, 1) for w in data["semaine"]], "#10b981"),
                        ("Taux d'absentéisme", [round(w["taux_absence"] or 0, 1) for w in data["semaine"]], "#e11d48"),
                        ("Taux de retard", [round(w["taux_retard"] or 0, 1) for w in data["semaine"]], "#f97316")],
                       kind="bars", percent=True)
        if c["admin"] and data.get("arrivees"):
            arr = [a for a in data["arrivees"] if a["n"]]
            deck.chart("Heures d'arrivée au bureau", "Premiers pointages par quart d'heure (réservé à l'administrateur).",
                       [f"{a['minute'] // 60:02d}h{a['minute'] % 60:02d}" for a in arr],
                       [("Arrivées", [a["n"] for a in arr], "#10b981")], kind="columns")
        for t in c["tables"]:
            deck.table(t["title"], t["note"], t["headers"], t["rows"], t["widths"], t["align"], t["pills"], t["colors"],
                       max_rows=ppt.EMPLOYEE_MAX_ROWS if t.get("employees") else None)
    return Response(deck.output_bytes(), media_type=ppt.MEDIA_TYPE,
                    headers={"Content-Disposition": f'attachment; filename="{c["name"]}.pptx"'})
