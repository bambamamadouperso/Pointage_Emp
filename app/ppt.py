"""Exports PowerPoint (rapports, suivi journalier) : diapositives 16:9 aux couleurs de l'application — titre et
filtres, cartes d'indicateurs, graphiques natifs (modifiables dans PowerPoint), tableaux paginés avec statuts et taux
en couleur."""
import io
from datetime import datetime
from typing import Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
ROWS_PER_SLIDE = 14
EMPLOYEE_MAX_ROWS = 300  # au-delà : export Excel (une présentation de dizaines de diapositives n'est plus lisible)

NAVY, TEAL, INK, MUTED = (27, 79, 138), (14, 124, 110), (11, 26, 43), (122, 136, 152)
HEAD_BG, ZEBRA, CARD_BORDER, PAGE = (235, 241, 248), (248, 250, 252), (230, 236, 243), (255, 255, 255)
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)


def _rgb(color) -> RGBColor:
    if isinstance(color, str):
        return RGBColor.from_string(color.lstrip("#").upper())
    return RGBColor(*color)


class Deck:
    def __init__(self, title: str, subtitle: str = "", user: str = "", filters: Optional[list] = None):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.title, self.subtitle, self.user = title, subtitle, user
        self.stamp = datetime.now()
        self._cover(filters or [])

    # ------------------------------------------------------------------ éléments communs

    def _text(self, slide, x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, _rgb(color)
        run.font.name = "Segoe UI"
        return box

    def _rect(self, slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fill)
        if line:
            s.line.color.rgb = _rgb(line)
            s.line.width = Pt(0.75)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _slide(self, title: str, note: str = ""):
        """Diapositive de contenu : bandeau bleu, titre, sous-titre, pied de page."""
        slide = self.prs.slides.add_slide(self.blank)
        self._rect(slide, 0, 0, W, Inches(0.95), NAVY)
        self._rect(slide, 0, Inches(0.95), W, Inches(0.06), TEAL)
        self._text(slide, MARGIN, Inches(0.16), Inches(9), Inches(0.5), title, 24, True, (255, 255, 255))
        self._text(slide, MARGIN, Inches(0.58), Inches(6.8), Inches(0.3), note or self.subtitle, 11, False, (199, 219, 240))
        self._text(slide, Inches(9.6), Inches(0.22), Inches(3.25), Inches(0.3), "POINTAGE", 11, True, (255, 255, 255),
                   PP_ALIGN.RIGHT)
        self._text(slide, Inches(7.35), Inches(0.52), Inches(5.5), Inches(0.3), self.subtitle if note else "", 10, False,
                   (199, 219, 240), PP_ALIGN.RIGHT)
        n = len(self.prs.slides)
        self._text(slide, MARGIN, H - Inches(0.4), Inches(8), Inches(0.25),
                   f"{self.title}  ·  exporté le {self.stamp:%d/%m/%Y à %H:%M}" + (f" par {self.user}" if self.user else ""),
                   9, False, MUTED)
        self._text(slide, W - MARGIN - Inches(1.5), H - Inches(0.4), Inches(1.5), Inches(0.25), str(n), 9, False, MUTED,
                   PP_ALIGN.RIGHT)
        return slide

    def _cover(self, filters: list) -> None:
        slide = self.prs.slides.add_slide(self.blank)
        self._rect(slide, 0, 0, W, H, NAVY)
        self._rect(slide, 0, Inches(4.55), W, Inches(0.08), TEAL)
        self._text(slide, Inches(0.8), Inches(0.7), Inches(6), Inches(0.4), "POINTAGE", 14, True, (199, 219, 240))
        self._text(slide, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2), self.title, 44, True, (255, 255, 255))
        self._text(slide, Inches(0.8), Inches(3.45), Inches(11.5), Inches(0.6), self.subtitle, 20, False, (199, 219, 240))
        lines = [f"{label} : {value}" for label, value in filters]
        if lines:
            self._text(slide, Inches(0.8), Inches(4.9), Inches(11.5), Inches(1.6), "   ·   ".join(lines), 13, False,
                       (220, 232, 245))
        self._text(slide, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.3),
                   f"Exporté le {self.stamp:%d/%m/%Y à %H:%M}" + (f" par {self.user}" if self.user else ""), 11, False,
                   (160, 185, 215))

    # ------------------------------------------------------------------ contenus

    def message(self, text: str) -> None:
        slide = self._slide("Résultat")
        self._text(slide, MARGIN, Inches(3.2), W - 2 * MARGIN, Inches(1), text, 20, False, MUTED, PP_ALIGN.CENTER)

    def kpis(self, title: str, items: list, volumes: Optional[list] = None) -> None:
        """Cartes d'indicateurs (libellé, valeur, précision, couleur), 3 par ligne, et volumes en bas."""
        slide = self._slide(title)
        cols, gap = 3, Inches(0.3)
        cw = (W - 2 * MARGIN - gap * (cols - 1)) / cols
        ch = Inches(2.05)
        for n, (label, value, hint, color) in enumerate(items):
            x = MARGIN + (cw + gap) * (n % cols)
            y = Inches(1.35) + (ch + gap) * (n // cols)
            card = self._rect(slide, int(x), int(y), int(cw), int(ch), PAGE, CARD_BORDER, MSO_SHAPE.ROUNDED_RECTANGLE)
            card.adjustments[0] = 0.06
            self._rect(slide, int(x), int(y + Inches(0.35)), Inches(0.09), int(ch - Inches(0.7)), color)
            self._text(slide, int(x + Inches(0.35)), int(y + Inches(0.3)), int(cw - Inches(0.5)), Inches(0.3),
                       label.upper(), 11, True, MUTED)
            self._text(slide, int(x + Inches(0.35)), int(y + Inches(0.68)), int(cw - Inches(0.5)), Inches(0.75),
                       value, 36, True, INK)
            self._text(slide, int(x + Inches(0.35)), int(y + Inches(1.5)), int(cw - Inches(0.5)), Inches(0.4),
                       hint, 11, False, MUTED)
        if volumes:
            self._text(slide, MARGIN, Inches(6.25), W - 2 * MARGIN, Inches(0.5),
                       "   ·   ".join(f"{label} : {value}" for label, value in volumes), 12, False, INK)

    def chart(self, title: str, note: str, categories: list, series: list, kind: str = "columns",
              percent: bool = False) -> None:
        """Graphique natif : « stacked » (colonnes empilées), « columns », « bars » (barres horizontales groupées)."""
        if not categories:
            return
        slide = self._slide(title, note)
        data = CategoryChartData()
        data.categories = categories
        for name, values, _ in series:
            data.add_series(name, values)
        chart_type = {"stacked": XL_CHART_TYPE.COLUMN_STACKED, "columns": XL_CHART_TYPE.COLUMN_CLUSTERED,
                      "bars": XL_CHART_TYPE.BAR_CLUSTERED}[kind]
        frame = slide.shapes.add_chart(chart_type, MARGIN, Inches(1.3), W - 2 * MARGIN, Inches(5.6), data)
        ch = frame.chart
        ch.font.size, ch.font.name = Pt(11), "Segoe UI"
        ch.has_title = False  # le titre est celui de la diapositive
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        for s, (_, _, color) in zip(ch.series, series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb(color)
        plot = ch.plots[0]
        plot.gap_width = 60 if kind != "bars" else 40
        if kind == "stacked":
            plot.overlap = 100
        if kind == "bars":
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format, labels.number_format_is_linked = '0.0" %"' if percent else "0", False
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            labels.font.size = Pt(10)
            ch.category_axis.reverse_order = True  # lundi en haut
        value_axis = ch.value_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = _rgb(CARD_BORDER)
        value_axis.tick_labels.font.color.rgb = _rgb(MUTED)
        if percent:
            value_axis.maximum_scale = 100
            value_axis.tick_labels.number_format, value_axis.tick_labels.number_format_is_linked = '0" %"', False
        ch.category_axis.tick_labels.font.color.rgb = _rgb(INK)

    def table(self, title: str, note: str, headers: list, rows: list, widths: list, align: Optional[list] = None,
              pills: Optional[dict] = None, colors: Optional[dict] = None, max_rows: Optional[int] = None) -> None:
        """Tableau paginé (en-tête répété sur chaque diapositive) ; `pills` : {(ligne, colonne): (fond, texte)} ;
        `colors` : {(ligne, colonne): texte}."""
        pills, colors = pills or {}, colors or {}
        align = align or ["L"] * len(headers)
        shown = rows[:max_rows] if max_rows else rows
        pages = [shown[i:i + ROWS_PER_SLIDE] for i in range(0, len(shown), ROWS_PER_SLIDE)] or [[]]
        total_w = W - 2 * MARGIN
        scale = total_w / sum(widths)
        font = 11 if len(headers) <= 8 else 9 if len(headers) <= 12 else 8
        for p, page in enumerate(pages):
            suffix = f" ({p + 1}/{len(pages)})" if len(pages) > 1 else ""
            extra = f" — {len(rows) - len(shown)} ligne(s) de plus dans l'export Excel" if max_rows and len(rows) > len(shown) else ""
            slide = self._slide(title + suffix, note + extra)
            shape = slide.shapes.add_table(len(page) + 1, len(headers), MARGIN, Inches(1.3), int(total_w),
                                           Inches(0.4) * (len(page) + 1))
            tbl = shape.table
            tbl.first_row = True
            for j, w in enumerate(widths):
                tbl.columns[j].width = Emu(int(w * scale))
            for j, h in enumerate(headers):
                self._cell(tbl.cell(0, j), str(h).upper(), font - 1, True, NAVY, HEAD_BG, align[j])
            for i, row in enumerate(page):
                index = p * ROWS_PER_SLIDE + i
                for j, value in enumerate(row):
                    pill = pills.get((index, j))
                    bg = pill[0] if pill else (ZEBRA if i % 2 else PAGE)
                    fg = pill[1] if pill else colors.get((index, j), INK)
                    self._cell(tbl.cell(i + 1, j), "" if value is None else str(value), font, bool(pill), fg, bg, align[j])

    @staticmethod
    def _cell(cell, text, size, bold, color, fill, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = {"R": PP_ALIGN.RIGHT, "C": PP_ALIGN.CENTER}.get(align, PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Segoe UI"
        run.font.color.rgb = _rgb(color)

    def output_bytes(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()
