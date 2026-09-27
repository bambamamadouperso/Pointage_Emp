"""Source Google Sheets : lecture du classeur, typage, authentification et synchronisation."""
import base64
import io
import json
import os
from datetime import date, datetime, time
from decimal import Decimal

import pytest
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding, rsa
from openpyxl import Workbook
from sqlalchemy import BigInteger, Boolean, Date, DateTime, Numeric, Text, create_engine, text

from app import gsheet
from app.crypto import encrypt
from app.database import SessionLocal
from app.models import Connection, JobRun, LogEntry, SyncJob, TableMapping
from app.sync import run_job


def make_xlsx(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


POINTAGE = [
    ["Matricule", "Nom complet", "Date d'arrivée", "Heure", "Présent", "Heures sup.", "Salaire", None],
    ["E001", "Awa Diallo", datetime(2024, 1, 2), datetime(2024, 1, 2, 8, 5), True, 1.5, 450000, None],
    ["E002", "Moussa Ndiaye", datetime(2024, 1, 3), datetime(2024, 1, 3, 7, 55), False, 0, 320000, None],
    [None, None, None, None, None, None, None, None],
    ["E003", "Fatou Sow", datetime(2024, 1, 4), datetime(2024, 1, 4, 9, 0), True, 2.25, 350000, None],
]


def test_parse_spreadsheet_id():
    sid = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    assert gsheet.parse_spreadsheet_id(f"https://docs.google.com/spreadsheets/d/{sid}/edit#gid=0") == sid
    assert gsheet.parse_spreadsheet_id(sid) == sid
    assert gsheet.parse_spreadsheet_id(f"https://drive.google.com/open?id={sid}") == sid


def test_normalize_identifier():
    assert gsheet.normalize_identifier("Date d'arrivée") == "date_d_arrivee"
    assert gsheet.normalize_identifier("  Heures sup. ") == "heures_sup"
    assert gsheet.normalize_identifier("2024") == "c_2024"
    assert gsheet.normalize_identifier(None, 3) == "col_4"
    assert gsheet.normalize_identifier("Feuille 1") == "feuille_1"


def test_read_workbook_and_types():
    sheets = gsheet.read_workbook(make_xlsx({"Pointage": POINTAGE, "Vide": []}))
    assert list(sheets) == ["Pointage", "Vide"]
    data = sheets["Pointage"]
    # Colonne sans en-tête ni valeur supprimée ; ligne vide ignorée.
    assert data.columns == ["matricule", "nom_complet", "date_d_arrivee", "heure", "present", "heures_sup", "salaire"]
    assert len(data.rows) == 3
    table = gsheet.build_table("Pointage", data)
    types = {c.name: type(c.type) for c in table.columns}
    assert types == {
        "matricule": Text, "nom_complet": Text, "date_d_arrivee": Date, "heure": DateTime,
        "present": Boolean, "heures_sup": Numeric, "salaire": BigInteger,
    }
    assert sheets["Vide"].columns == []


def test_duplicate_headers():
    data = gsheet.read_workbook(make_xlsx({"S": [["Nom", "nom", "NOM"], [1, 2, 3]]}))["S"]
    assert data.columns == ["nom", "nom_2", "nom_3"]


def test_convert():
    assert gsheet.convert("1 234,50", Numeric()) == Decimal("1234.50")
    assert gsheet.convert("1,234.50", Numeric()) == Decimal("1234.50")
    assert gsheet.convert(1.5, Numeric()) == Decimal("1.5")
    assert gsheet.convert("12", BigInteger()) == 12
    assert gsheet.convert(12.0, BigInteger()) == 12
    with pytest.raises(ValueError):
        gsheet.convert("12,5", BigInteger())
    assert gsheet.convert("27/09/2026", Date()) == date(2026, 9, 27)
    assert gsheet.convert("2026-09-27 08:30", DateTime()) == datetime(2026, 9, 27, 8, 30)
    assert gsheet.convert("Vrai", Boolean()) is True
    assert gsheet.convert(12.0, Text()) == "12"
    assert gsheet.convert(datetime(2026, 1, 2), Text()) == "2026-01-02"
    with pytest.raises(ValueError):
        gsheet.convert("abc", Date())


def test_access_token_builds_signed_jwt(monkeypatch):
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                            serialization.NoEncryption()).decode()
    captured = {}

    class FakeResponse(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_urlopen(request, timeout):
        captured["url"] = request.full_url
        captured["body"] = request.data.decode()
        return FakeResponse(json.dumps({"access_token": "tok"}).encode())

    monkeypatch.setattr(gsheet.urllib.request, "urlopen", fake_urlopen)
    token = gsheet.access_token({"client_email": "sa@p.iam.gserviceaccount.com", "private_key": pem})
    assert token == "tok"
    assert captured["url"] == "https://oauth2.googleapis.com/token"
    assertion = dict(p.split("=", 1) for p in captured["body"].split("&"))["assertion"]
    header, claims, signature = assertion.split(".")
    pad = lambda s: s + "=" * (-len(s) % 4)  # noqa: E731
    payload = json.loads(base64.urlsafe_b64decode(pad(claims)))
    assert payload["iss"] == "sa@p.iam.gserviceaccount.com" and "drive.readonly" in payload["scope"]
    key.public_key().verify(base64.urlsafe_b64decode(pad(signature)), f"{header}.{claims}".encode(),
                            padding.PKCS1v15(), hashes.SHA256())


def test_public_download_errors(monkeypatch):
    conn = Connection(kind="gsheet", database="abc", username=gsheet.AUTH_PUBLIC, password_enc="")
    monkeypatch.setattr(gsheet, "_http_get", lambda url, headers=None: (200, "text/html", b"<html>login"))
    with pytest.raises(gsheet.SheetError, match="Tous les utilisateurs disposant du lien"):
        gsheet.download_workbook(conn)
    payload = make_xlsx({"A": [["x"], [1]]})
    calls = []
    monkeypatch.setattr(gsheet, "_http_get", lambda url, headers=None: calls.append(url) or (200, "", payload))
    assert gsheet.download_workbook(conn) == payload
    assert calls == ["https://docs.google.com/spreadsheets/d/abc/export?format=xlsx"]


# --------------------------------------------------------------------------- synchronisation réelle

POSTGRES_URL = os.getenv("TEST_POSTGRES_URL")
SCHEMA = "it_gsheet"


@pytest.fixture
def sheet_job(monkeypatch):
    if not POSTGRES_URL:
        pytest.skip("TEST_POSTGRES_URL non définie")
    from sqlalchemy.engine import make_url

    state = {"sheets": {"Pointage": POINTAGE, "Feuille 2": [["Code", "Libellé"], ["A", "Un"], ["B", "Deux"]]}}
    monkeypatch.setattr(gsheet, "download_workbook", lambda conn: make_xlsx(state["sheets"]))
    dst = create_engine(POSTGRES_URL)
    with dst.begin() as c:
        c.execute(text(f"DROP SCHEMA IF EXISTS {SCHEMA} CASCADE"))
    u = make_url(POSTGRES_URL)
    with SessionLocal() as db:
        db.query(SyncJob).filter_by(name="gs-job").delete()
        db.query(Connection).filter(Connection.name.in_(["gs-src", "gs-dst"])).delete()
        db.commit()
        src = Connection(name="gs-src", kind="gsheet", host="docs.google.com", port=443, database="abc",
                         username=gsheet.AUTH_PUBLIC, password_enc="")
        tgt = Connection(name="gs-dst", kind="postgresql", host=u.host, port=u.port or 5432, database=u.database,
                         username=u.username, password_enc=encrypt(u.password or ""))
        db.add_all([src, tgt])
        db.flush()
        job = SyncJob(name="gs-job", source_id=src.id, target_id=tgt.id, target_schema=SCHEMA, interval_seconds=60)
        job.tables = [
            TableMapping(source_table="Pointage", target_table="pointage", mode="incremental",
                         incremental_column="date_d_arrivee", key_columns="matricule"),
            TableMapping(source_table="Feuille 2", target_table="feuille_2", mode="full"),
        ]
        db.add(job)
        db.commit()
        job_id = job.id
    yield job_id, state, dst
    dst.dispose()


def _run(job_id):
    with SessionLocal() as db:
        return db.get(JobRun, run_job(job_id, "manual"))


def test_sheet_sync_end_to_end(sheet_job):
    job_id, state, dst = sheet_job
    run = _run(job_id)
    assert run.status == "success", run.message
    assert run.rows_written == 3 + 2
    with dst.connect() as c:
        rows = c.execute(text(f"SELECT matricule, nom_complet, date_d_arrivee, present, heures_sup, salaire "
                              f"FROM {SCHEMA}.pointage ORDER BY matricule")).all()
        assert rows[0] == ("E001", "Awa Diallo", date(2024, 1, 2), True, Decimal("1.5"), 450000)
        assert c.execute(text(f"SELECT code, libelle FROM {SCHEMA}.feuille_2 ORDER BY code")).all() == [
            ("A", "Un"), ("B", "Deux")]

    # Modification de la feuille : mise à jour, ajout, doublon de clé, valeur illisible, colonne ajoutée.
    state["sheets"]["Pointage"] = [
        POINTAGE[0][:7] + ["Commentaire"],
        ["E003", "Fatou Sow", datetime(2024, 1, 4), datetime(2024, 1, 4, 9, 0), True, 3, 360000, "maj"],
        ["E004", "Ibrahima Fall", datetime(2024, 1, 5), datetime(2024, 1, 5, 8, 0), True, "n/a", 900000, None],
        ["E004", "Ibrahima Fall", datetime(2024, 1, 5), datetime(2024, 1, 5, 8, 0), True, 1, 900000, "doublon"],
    ]
    state["sheets"]["Feuille 2"] = [["Code", "Libellé"], ["C", "Trois"]]
    run = _run(job_id)
    assert run.status == "success", run.message
    with dst.connect() as c:
        rows = c.execute(text(f"SELECT matricule, heures_sup, salaire, commentaire FROM {SCHEMA}.pointage "
                              f"ORDER BY matricule")).all()
        assert rows == [
            ("E001", Decimal("1.5"), 450000, None),
            ("E002", Decimal("0"), 320000, None),
            ("E003", Decimal("3"), 360000, "maj"),
            ("E004", Decimal("1"), 900000, "doublon"),
        ]
        # Mode complet : la feuille remplace entièrement la table.
        assert c.execute(text(f"SELECT code FROM {SCHEMA}.feuille_2")).all() == [("C",)]
    with SessionLocal() as db:
        messages = [l.message for l in db.query(LogEntry).filter_by(run_id=run.id)]
        mapping = db.query(TableMapping).filter_by(job_id=job_id, source_table="Pointage").one()
    assert any("doublon(s) de clé" in m for m in messages)
    assert any("Colonnes ajoutées dans la cible : commentaire" in m for m in messages)
    assert mapping.last_value_display == "2024-01-05"


def test_sheet_errors(sheet_job):
    job_id, state, _ = sheet_job
    with SessionLocal() as db:
        db.add(TableMapping(job_id=job_id, source_table="Inexistant", target_table="x"))
        db.commit()
    run = _run(job_id)
    assert run.status == "partial" and run.tables_failed == 1
    with SessionLocal() as db:
        err = db.query(LogEntry).filter_by(run_id=run.id, level="ERROR").one()
    assert "Onglet « Inexistant » introuvable" in err.message
