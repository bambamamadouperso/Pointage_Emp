"""Module de pointage : calculs PostgreSQL (exemples du cahier des charges), écran, export, rôles, audit.

Nécessite TEST_POSTGRES_URL (voir test_integration.py).
"""
import io
import os
from datetime import date, datetime, timedelta

import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

from app import pointage
from app.crypto import encrypt
from app.database import SessionLocal
from app.models import AuditEntry, Connection, PointageConfig, User

POSTGRES_URL = os.getenv("TEST_POSTGRES_URL")
pytestmark = pytest.mark.skipif(not POSTGRES_URL, reason="TEST_POSTGRES_URL non définie")
SCHEMA = "pt_test"
MON, TUE, WED, SAT = date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23), date(2026, 9, 26)


def at(day: date, hhmm: str) -> datetime:
    h, m = map(int, hhmm.split(":"))
    return datetime(day.year, day.month, day.day, h, m)


@pytest.fixture(scope="module")
def pg():
    engine = create_engine(POSTGRES_URL)
    with engine.begin() as c:
        c.execute(text(f"DROP SCHEMA IF EXISTS {SCHEMA} CASCADE"))
        c.execute(text(f"CREATE SCHEMA {SCHEMA}"))
        c.execute(text(f"CREATE TABLE {SCHEMA}.services (id int PRIMARY KEY, libelle text)"))
        # Noms et prénoms dans une autre table (Personnel) ; hiérarchie dans une table dédiée (par matricule).
        c.execute(text(f"""CREATE TABLE {SCHEMA}."Employes" ("IDEmployes" int PRIMARY KEY, "Matricule" text,
                          "IDPersonnel" int, "IDService" int, "Actif" text)"""))
        c.execute(text(f"""CREATE TABLE {SCHEMA}."Personnel" ("IDPersonnel" int PRIMARY KEY, "Nom" text, "Prenom" text,
                          "IDStatut" int)"""))
        c.execute(text(f"CREATE TABLE {SCHEMA}.statuts (id int PRIMARY KEY, libelle text)"))
        c.execute(text(f"INSERT INTO {SCHEMA}.statuts VALUES (1, 'Cadre'), (2, 'Non cadre')"))
        c.execute(text(f"""CREATE TABLE {SCHEMA}.hierarchie (matricule_employe text, matricule_responsable text)"""))
        c.execute(text(f"CREATE TABLE {SCHEMA}.punchlog (id serial PRIMARY KEY, employe_id int, "
                       f"date_pointage date, heure_pointage text, terminal text)"))
        c.execute(text(f"INSERT INTO {SCHEMA}.services VALUES (1, 'Production'), (2, 'RH')"))
        c.execute(text(f"""INSERT INTO {SCHEMA}."Employes" VALUES
            (1, 'E001', 101, 1, '1'), (2, 'E002', 102, 1, '1'), (3, 'E003', 103, 2, '1'), (4, 'E004', 104, 2, '1'),
            (5, 'E005', 105, 2, '1'), (6, 'E006', 106, 1, '0'), (7, 'E007', 107, 1, '1')"""))
        c.execute(text(f"""INSERT INTO {SCHEMA}."Personnel" VALUES
            (101, 'Diallo', 'Awa', 1), (102, 'Ndiaye', 'Moussa', 2), (103, 'Sow', 'Fatou', 1), (104, 'Fall', 'Ibou', 2),
            (105, 'Ba', 'Khady', 2), (106, 'Gueye', 'Ousmane', 2), (107, 'Diop', 'Aminata', 2),
            (108, 'Kane', 'Omar', 2)"""))  # Kane : dans Personnel, sans fiche dans Employes
        # Diallo encadre Ndiaye, Sow et Gueye ; Sow encadre Fall et Ba ; Diop n'a pas de responsable.
        c.execute(text(f"""INSERT INTO {SCHEMA}.hierarchie VALUES
            ('E002', 'E001'), ('E003', 'E001'), ('E006', 'E001'), ('E004', 'E003'), ('E005', 'E003'), ('E007', NULL)"""))
        punches = [
            (1, MON, "07:20"), (1, MON, "12:00"), (1, MON, "17:10"),   # à l'heure : 7h30 validées, 8h20 effectives
            (2, MON, "08:10"), (2, MON, "16:45"),                      # retard : 6h50 validées
            (3, MON, "07:35"), (3, MON, "15:50"),                      # à l'heure (tolérance) : 6h50 validées
            (4, MON, "07:40"),                                         # un seul pointage : incomplet
            (7, MON, "07:45"), (7, MON, "12:00"),                      # 7h45 pile = retard, pas de pause déduite
            (2, TUE, "08:10"), (2, TUE, "16:30"),                      # seuil passé à 8h15 le mardi : à l'heure
            (1, SAT, "08:00"), (1, SAT, "12:00"),                      # samedi (non ouvré)
            (99, MON, "09:00"), (99, MON, "17:00"),                    # employé inconnu
        ]
        for emp, day, hhmm in punches:
            c.execute(text(f"INSERT INTO {SCHEMA}.punchlog (employe_id, date_pointage, heure_pointage, terminal) "
                           f"VALUES (:e, :d, :h, 'T1')"), {"e": emp, "d": day, "h": hhmm.replace(":", "") + "00"})
    yield engine
    engine.dispose()


@pytest.fixture(scope="module")
def configured(pg):
    """Configure et installe le module via l'administration."""
    from fastapi.testclient import TestClient

    from app.main import app

    u = make_url(POSTGRES_URL)
    with SessionLocal() as db:
        db.query(PointageConfig).delete()
        db.query(Connection).filter(Connection.name == "pt-pg").delete()
        conn = Connection(name="pt-pg", kind="postgresql", host=u.host, port=u.port or 5432, database=u.database,
                          username=u.username, password_enc=encrypt(u.password or ""))
        db.add(conn)
        db.commit()
        conn_id = conn.id
    with TestClient(app) as client:
        client.post("/login", data={"username": "admin", "password": "secret"})
        page = client.get(f"/admin/pointage?conn_id={conn_id}&schema={SCHEMA}&punch_table=punchlog&emp_table=Employes")
        assert page.status_code == 200 and "heure_pointage" in page.text
        r = client.post("/admin/pointage", data={
            "conn_id": conn_id, "schema": SCHEMA, "objects_schema": "", "punch_table": "punchlog",
            "punch_emp_col": "employe_id", "punch_ts_col": "date_pointage", "punch_time_col": "heure_pointage",
            "emp_table": "Employes", "emp_key_col": "IDEmployes", "emp_matricule_col": "Matricule",
            "emp_service_col": "IDService",
            "service_table": "services", "service_key_col": "id", "service_label_col": "libelle",
            "emp_active_col": "Actif", "emp_active_values": "1",
            "person_table": "Personnel", "emp_person_col": "IDPersonnel", "person_key_col": "IDPersonnel",
            "person_nom_col": "Nom", "person_prenom_col": "Prenom",
            "hier_table": "hierarchie", "hier_emp_col": "matricule_employe",
            "hier_manager_col": "matricule_responsable", "hier_ref": "matricule",
            "cat_in": "person", "cat_col": "IDStatut", "cat_table": "statuts", "cat_key_col": "id", "cat_label_col": "libelle",
        }, follow_redirects=True)
        assert "installées" in r.text, r.text[:3000]
        # Seuil de retard porté à 8h15 à partir du mardi ; mercredi férié.
        values = {k: v for k, _, _, v in pointage.PARAMS if k != "jours_ouvres"}
        r = client.post("/admin/parametres", data={**values, "seuil_retard": "08:15", "jours_ouvres": ["1", "2", "3", "4", "5"],
                                                   "date_effet": TUE.isoformat()}, follow_redirects=True)
        assert "1 paramètre(s) modifié(s)" in r.text
        r = client.post("/admin/feries", data={"jour": WED.isoformat(), "libelle": "Fête"}, follow_redirects=True)
        assert "Fête" in r.text
    return conn_id


def rows(pg, du, au=None):
    with pg.connect() as c:
        return {(r.matricule, r.jour): r for r in c.execute(text(
            f"SELECT * FROM {SCHEMA}.f_pointage_journalier(:du, :au)"), {"du": du, "au": au or du}).mappings()}


def hm(value) -> str:
    return pointage.hhmm(value)


def test_calculations_match_specification(configured, pg):
    r = rows(pg, MON)
    awa = r[("E001", MON)]
    assert awa.statut == "A_L_HEURE" and hm(awa.premier_pointage) == "07h20" and hm(awa.dernier_pointage) == "17h10"
    assert hm(awa.duree_validee) == "7h30" and hm(awa.duree_effective) == "8h20" and awa.nb_pointages == 3
    moussa = r[("E002", MON)]
    assert moussa.statut == "RETARD" and hm(moussa.duree_validee) == "6h50" and hm(moussa.duree_effective) == "7h05"
    fatou = r[("E003", MON)]
    assert fatou.statut == "A_L_HEURE" and hm(fatou.duree_validee) == "6h50" and hm(fatou.duree_effective) == "6h45"
    ibou = r[("E004", MON)]
    assert ibou.statut == "INCOMPLET" and ibou.duree_validee is None and ibou.duree_effective is None
    assert r[("E005", MON)].statut == "ABSENT" and r[("E005", MON)].service == "RH"
    assert ("E006", MON) not in r  # inactif : jamais absent
    aminata = r[("E007", MON)]  # 7h45 pile = retard ; départ 12h00 < fin de pause : pas de déduction
    assert aminata.statut == "RETARD" and hm(aminata.duree_validee) == "4h15" and hm(aminata.duree_effective) == "4h15"
    assert r[("99", MON)].nom == "Hors liste" and r[("99", MON)].hors_liste
    assert not awa.hors_liste
    assert r[("E001", MON)].service == "Production" and r[("E001", MON)].statut_libelle == "À l'heure"
    # Noms venant de la table Personnel, responsable venant de la table hierarchie.
    assert (awa.nom, awa.prenom, awa.responsable) == ("Diallo", "Awa", None)
    assert moussa.responsable == "Diallo Awa" and moussa.responsable_key == "1"
    assert r[("E004", MON)].responsable == "Sow Fatou" and aminata.responsable is None


def test_automatic_upgrade_of_sql_objects(configured, pg, logged_client):
    """Après une mise à jour de l'application, les fonctions PostgreSQL sont réinstallées d'elles-mêmes."""
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        cfg.sql_version = 1
        db.commit()
    with pg.begin() as c:  # ancienne version : sans la colonne hors_liste
        c.execute(text(f"DROP VIEW {SCHEMA}.v_pointage_journalier"))
        c.execute(text(f"DROP FUNCTION {SCHEMA}.f_pointage_journalier(date, date)"))
    assert "Diallo" in logged_client.get(f"/suivi?date={MON.isoformat()}").text
    with SessionLocal() as db:
        assert db.query(PointageConfig).one().sql_version == pointage.SQL_VERSION
        assert db.query(AuditEntry).filter(AuditEntry.action.like("Calculs du pointage mis à jour%")).count()


def test_hierarchy_team_function(configured, pg):
    with pg.connect() as c:
        team = dict(c.execute(text(f"SELECT emp_key, niveau FROM {SCHEMA}.f_pointage_equipe('1')")).all())
        assert team == {"1": 0, "2": 1, "3": 1, "6": 1, "4": 2, "5": 2}
        assert {k for (k,) in c.execute(text(f"SELECT emp_key FROM {SCHEMA}.f_pointage_equipe('3')"))} == {"3", "4", "5"}
        emp = c.execute(text(f"SELECT * FROM {SCHEMA}.v_pointage_employes WHERE matricule = 'E005'")).mappings().one()
        assert (emp["nom"], emp["responsable"], emp["service"]) == ("Ba", "Sow Fatou", "RH")


def test_history_holidays_and_weekends(configured, pg):
    week = rows(pg, MON, SAT)
    assert week[("E002", TUE)].statut == "A_L_HEURE"  # seuil 8h15 en vigueur depuis mardi
    assert week[("E002", MON)].statut == "RETARD"     # lundi : seuil 7h45 (valeur de l'époque)
    assert week[("E005", TUE)].statut == "ABSENT"
    # On part de la liste des employés : chacun apparaît chaque jour, même sans pointage.
    assert week[("E005", WED)].statut == "NON_OUVRE"  # jour férié : pas d'absence
    assert week[("E005", SAT)].statut == "NON_OUVRE"  # samedi non ouvré
    assert week[("E005", SAT)].statut_libelle == "Jour non ouvré"
    assert ("E006", MON) not in week                  # inactif sans pointage
    sat = week[("E001", SAT)]
    assert sat.jour_ouvre is False and sat.nb_pointages == 2
    with pg.connect() as c:  # la vue couvre tout l'historique
        n = c.execute(text(f"SELECT count(*) FROM {SCHEMA}.v_pointage_journalier WHERE jour = :d"), {"d": MON}).scalar()
        assert n == len(rows(pg, MON))
        assert c.execute(text(f"SELECT count(*) FROM {SCHEMA}.v_pointage_brut")).scalar() == 16


def test_screen_filters_detail_and_export(configured, logged_client):
    from openpyxl import load_workbook

    page = logged_client.get(f"/suivi?date={MON.isoformat()}")
    assert page.status_code == 200
    html = page.text
    assert "Diallo" in html and "st-late" in html and "Lundi 21/09/2026" in html
    assert "Ponctualité" in html and "Durées moyennes" in html
    only_late = logged_client.get(f"/suivi?date={MON.isoformat()}&statut=RETARD").text
    assert "E002</td>" in only_late and "E001</td>" not in only_late
    search = logged_client.get(f"/suivi?date={MON.isoformat()}&q=E003").text
    assert "E003</td>" in search and "E002</td>" not in search
    service = logged_client.get(f"/suivi?date={MON.isoformat()}&service=RH").text
    assert "E003</td>" in service and "E001</td>" not in service
    hors = logged_client.get(f"/suivi?date={MON.isoformat()}&pop=hors").text
    assert "99</td>" in hors and "E001</td>" not in hors and "Hors liste" in hors
    liste = logged_client.get(f"/suivi?date={MON.isoformat()}&pop=liste").text
    assert "99</td>" not in liste and "E005</td>" in liste
    period = logged_client.get(f"/suivi?du={MON.isoformat()}&au={SAT.isoformat()}&sort=validee&dir=desc").text
    assert "journées-employé" in period
    team = logged_client.get(f"/suivi?date={MON.isoformat()}&equipe=3").text
    assert "E003</td>" in team and "E004</td>" in team and "E005</td>" in team and "E001</td>" not in team
    assert "Responsable" in team and "Sow Fatou · 2 direct(s)" in team.replace(" (E003)", "")
    directs = logged_client.get(f"/suivi?date={MON.isoformat()}&equipe=1&directs=1").text
    assert "E002</td>" in directs and "E004</td>" not in directs
    detail = logged_client.get(f"/suivi/detail?key=1&jour={MON.isoformat()}").text
    assert "3 pointage(s)" in detail and "07:20:00" in detail and "T1" in detail
    xlsx = logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}&statut=RETARD&statut=ABSENT")
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content))["Suivi journalier"]
    statuts = {ws.cell(r, 8).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "Moyenne")}
    assert statuts == {"En retard", "Absent"}
    # Statut entre le premier et le dernier pointage, comme à l'écran.
    assert [ws.cell(1, c).value for c in (7, 8, 9)] == ["1er pointage", "Statut", "Dernier pointage"]
    assert ws.cell(1, 11).value == "Durée standard"


def test_roles_and_audit(configured, client):
    from app import auth

    with SessionLocal() as db:
        db.query(User).filter(User.username.in_(["lect", "mgr"])).delete()
        db.add_all([User(username="lect", role="lecteur", password_hash=auth.hash_password("motdepasse1")),
                    User(username="mgr", role="manager", password_hash=auth.hash_password("motdepasse2"))])
        db.commit()
    r = client.post("/login", data={"username": "lect", "password": "motdepasse1"}, follow_redirects=False)
    assert r.headers["location"] == "/suivi"
    assert client.get(f"/suivi?date={MON.isoformat()}").status_code == 200
    for path in ("/", "/jobs", "/admin", "/connections"):
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/suivi", path
    assert client.post("/jobs/1/run", headers={"accept": "application/json"}).status_code == 403
    client.post("/logout")
    client.post("/login", data={"username": "mgr", "password": "motdepasse2"})
    assert client.get("/jobs").status_code == 200 and client.get("/runs").status_code == 200
    assert client.get("/admin", follow_redirects=False).status_code == 303
    assert client.get("/jobs/new", follow_redirects=False).status_code == 303
    client.post("/logout")
    assert client.post("/login", data={"username": "lect", "password": "mauvais"}, follow_redirects=False) \
        .headers["location"] == "/login"
    # Désactivation : la session en cours est coupée immédiatement.
    client.post("/login", data={"username": "lect", "password": "motdepasse1"})
    with SessionLocal() as db:
        db.query(User).filter_by(username="lect").one().active = False
        db.commit()
    assert client.get("/suivi", follow_redirects=False).headers["location"] == "/login"
    with SessionLocal() as db:
        actions = {a for (a,) in db.query(AuditEntry.action)}
    assert {"Source des pointages configurée", "Paramètres horaires modifiés", "Jour férié ajouté", "Connexion",
            "Échec de connexion"} <= actions


def test_account_limited_to_team(configured, client):
    from openpyxl import load_workbook

    from app import auth

    with SessionLocal() as db:
        db.query(User).filter(User.username.in_(["chef", "perdu"])).delete()
        db.add_all([User(username="chef", role="lecteur", password_hash=auth.hash_password("motdepasse4"),
                         emp_matricule="E003", scope="equipe"),
                    User(username="perdu", role="manager", password_hash=auth.hash_password("motdepasse5"),
                         emp_matricule="X999", scope="equipe")])
        db.commit()
    client.post("/login", data={"username": "chef", "password": "motdepasse4"})
    page = client.get(f"/suivi?date={MON.isoformat()}").text
    assert "équipe de <strong>Sow Fatou</strong>" in page
    assert "badge(s) hors liste" not in page  # badges sans fiche : hors de toute équipe
    assert "E004</td>" in page and "E005</td>" in page and "E001</td>" not in page and "E002</td>" not in page
    assert "E002</td>" not in client.get(f"/suivi?date={MON.isoformat()}&equipe=1").text  # hors périmètre : ignoré
    assert "ne fait pas partie de votre équipe" in client.get(f"/suivi/detail?key=2&jour={MON.isoformat()}").text
    assert "3 pointage(s)" not in client.get(f"/suivi/detail?key=1&jour={MON.isoformat()}").text
    ws = load_workbook(io.BytesIO(client.get(f"/suivi/export.xlsx?date={MON.isoformat()}").content))["Suivi journalier"]
    assert {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)
            if ws.cell(r, 1).value not in (None, "Moyenne")} == {"E003", "E004", "E005"}
    client.post("/logout")
    client.post("/login", data={"username": "perdu", "password": "motdepasse5"})
    assert "rattaché à aucun employé" in client.get(f"/suivi?date={MON.isoformat()}").text


def test_rh_role_sees_everyone_without_sync(configured, client):
    from app import auth

    with SessionLocal() as db:
        db.query(User).filter_by(username="rh_global").delete()
        # Même rattaché à un employé « en équipe », le rôle RH voit tout le personnel.
        db.add(User(username="rh_global", role="rh", password_hash=auth.hash_password("motdepasse6"),
                    emp_matricule="E003", scope="equipe"))
        db.commit()
    r = client.post("/login", data={"username": "rh_global", "password": "motdepasse6"}, follow_redirects=False)
    assert r.headers["location"] == "/suivi"
    page = client.get(f"/suivi?date={MON.isoformat()}").text
    assert "E001</td>" in page and "E002</td>" in page and "E005</td>" in page
    assert "Ressources humaines" in page and "/admin/planning" in page and "/connections" not in page
    admin = client.get("/admin").text
    assert "Horaires postés" in admin and "Source des pointages" not in admin and "Traitements planifiés" not in admin
    for path in ("/admin/parametres", "/admin/planning", "/admin/terrain", "/admin/arrets", "/admin/resumes",
                 "/admin/rappels", "/admin/maternite", "/admin/validations", "/rapports", "/arrets"):
        assert client.get(path, follow_redirects=False).status_code == 200, path
    for path in ("/", "/jobs", "/runs", "/logs", "/data", "/connections", "/connections/new", "/admin/pointage",
                 "/admin/mails", "/admin/users", "/admin/audit"):
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/suivi", path
    assert client.post("/connections/save", data={"name": "x"}, headers={"accept": "application/json"}).status_code == 403
    assert client.post("/jobs/1/run", headers={"accept": "application/json"}).status_code == 403
    client.post("/logout")
    with SessionLocal() as db:
        db.query(User).filter_by(username="rh_global").delete()
        db.commit()


def test_admin_users_and_params_pages(configured, logged_client):
    r = logged_client.post("/admin/users/save", data={"username": "nouveau", "full_name": "N. Ouveau",
                                                      "role": "manager", "password": "motdepasse3"},
                           follow_redirects=True)
    assert "« nouveau » enregistré" in r.text
    with SessionLocal() as db:
        uid = db.query(User).filter_by(username="nouveau").one().id
    r = logged_client.post("/admin/users/save", data={"user_id": uid, "full_name": "X", "role": "lecteur",
                                                      "emp_matricule": "E001", "scope": "equipe"},
                           follow_redirects=True)
    with SessionLocal() as db:
        u = db.get(User, uid)
        assert u.role == "lecteur" and not u.active and u.team_only and u.emp_matricule == "E001"
    r = logged_client.post("/admin/users/save", data={"user_id": uid, "role": "lecteur", "scope": "equipe"},
                           follow_redirects=True)
    assert "indiquez le matricule" in r.text
    assert "Son équipe (E001)" in logged_client.get("/admin/users").text
    assert "au moins 8 caractères" in logged_client.post(f"/admin/users/{uid}/password", data={"password": "court"},
                                                         follow_redirects=True).text
    page = logged_client.get("/admin/parametres").text
    assert "08:15" in page and "Historique des modifications" in page and "Fête" in page
    bad = logged_client.post("/admin/parametres", data={"debut_journee": "09:00", "seuil_retard": "08:00",
                             "debut_pause": "13:00", "fin_pause": "14:00", "fin_journee": "16:30",
                             "duree_pause_deduite": "01:30", "jours_ouvres": ["1"], "date_effet": "2026-09-01"},
                             follow_redirects=True).text
    assert "seuil de retard doit être postérieur" in bad
    audit = logged_client.get("/admin/audit?q=Paramètres").text
    assert "07:45 → 08:15" in audit
    assert "Traitements planifiés" in logged_client.get("/admin").text


def test_diagnostics(configured, pg):
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    d = pointage.diagnostics(pg, m, days=(date.today() - date(2026, 8, 28)).days)  # fenêtre couvrant les données de test
    assert (d["total"], d["actifs"], d["inactifs"]) == (7, 6, 1)
    assert d["hors_liste"] == 1 and d["sans_pointage"] == 1  # badge 99 ; Ba Khady n'a jamais pointé
    assert d["dernier_pointage"] == datetime(2026, 9, 26, 12, 0)


def test_last_punch_with_timestamp_column(configured, pg):
    """Colonne des pointages de type horodatage (date + heure) : « horodatage - 1 » était refusé."""
    import dataclasses

    with pg.begin() as c:
        c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.punch_ts"))
        c.execute(text(f"CREATE TABLE {SCHEMA}.punch_ts (employe_id int, horodatage timestamp)"))
        c.execute(text(f"INSERT INTO {SCHEMA}.punch_ts VALUES (1, '2026-09-27 07:40'), (1, '2026-09-28 16:05'), "
                       f"(2, '2026-09-28 08:10')"))
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    m = dataclasses.replace(m, punch_table="punch_ts", punch_ts_col="horodatage", punch_time_col="")
    assert pointage.last_punch(pg, m, use_cache=False) == datetime(2026, 9, 28, 16, 5)


def test_population_warns_when_most_are_inactive(configured, pg, logged_client, monkeypatch):
    """Colonne « actif » mal réglée (presque personne d'actif) : bandeau d'alerte sur l'écran de suivi."""
    import dataclasses

    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    ok = pointage.population(pg, m)
    assert (ok["total"], ok["actifs"], ok["suspect"]) == (7, 6, False) and ok["valeurs"]
    assert "jamais absentes" not in logged_client.get("/suivi?date=2026-09-25").text

    wrong = dataclasses.replace(m, objects_schema="pt_test_actif", emp_active_values="zzz")
    with pg.begin() as c:
        c.execute(text("DROP SCHEMA IF EXISTS pt_test_actif CASCADE"))
    pointage.install(pg, wrong, "test")
    bad = pointage.population(pg, wrong)
    assert (bad["actifs"], bad["suspect"]) == (0, True)
    assert all(not active for _, active, _ in bad["valeurs"])

    monkeypatch.setattr(pointage, "population", lambda engine, mapping: bad)
    page = logged_client.get("/suivi?date=2026-09-25").text
    assert "jamais absentes" in page and "Valeurs considérées actives" in page


def test_filter_bar_and_contradictory_filters(configured, logged_client):
    """Filtres appliqués affichés (retirables), « directs » sans équipe ignoré, hors liste + service expliqué."""
    page = logged_client.get("/suivi?date=2026-09-25&service=Production&statut=ABSENT&statut=RETARD").text
    assert "Filtres appliqués" in page and "Service : Production" in page and "Statut : Absent, En retard" in page
    assert "Tout retirer" in page

    page = logged_client.get("/suivi?date=2026-09-25&directs=1").text
    assert "directs (N-1)" not in page  # sans responsable choisi, la case n'a pas d'effet
    assert "disabled" in page.split('name="directs"')[1].split(">")[0]

    page = logged_client.get("/suivi?date=2026-09-25&service=Production&pop=hors").text
    assert "Aucune ligne" in page and "n'ont ni service ni responsable" in page


def test_approved_leave_and_remote_work_replace_absence(configured, pg, logged_client):
    """Jour « Absent » couvert par un congé approuvé : « Congé Annuel » / « Congé exceptionnel », 8h validées et effectives."""
    import dataclasses

    with pg.begin() as c:
        c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.demandeconge"))
        c.execute(text(f"""CREATE TABLE {SCHEMA}.demandeconge (matricule text, datedebut timestamp, "Datefin" text,
                          etat text, "Document" text)"""))
        c.execute(text(f"""INSERT INTO {SCHEMA}.demandeconge VALUES
            ('E005', '2026-09-21 00:00', '20260923000000', 'Approuvée', 'CONGE'),        -- lun. → mer. (mer. férié)
            ('E004', '2026-09-21 00:00', '2026-09-21', 'Approuvée', 'CONGE EXCEP'),       -- a badgé : reste incomplet
            ('E007', '2026-09-22 00:00', '2026-09-22', ' APPROUVEE ', 'conge  excep'),    -- casse et espaces ignorés
            ('E003', '2026-09-22 00:00', '2026-09-22', 'En attente', 'CONGE'),            -- non approuvé
            ('E001', '2026-09-22 00:00', '2026-09-22', 'Approuvée', 'MALADIE')            -- autre nature"""))
        c.execute(text(f'DROP TABLE IF EXISTS {SCHEMA}."TdemandedeTeleTravail"'))
        c.execute(text(f"""CREATE TABLE {SCHEMA}."TdemandedeTeleTravail" (matricule text, datedebut timestamp,
                          "Datefin" timestamp, etat text)"""))
        c.execute(text(f"""INSERT INTO {SCHEMA}."TdemandedeTeleTravail" VALUES
            ('E003', '2026-09-22 00:00', '2026-09-22 00:00', 'Approuvée'),   -- télétravail
            ('E005', '2026-09-22 00:00', '2026-09-22 00:00', 'Approuvée'),   -- aussi en congé : le congé l'emporte
            ('E001', '2026-09-22 00:00', '2026-09-22 00:00', 'Refusée'),     -- non approuvé
            ('E002', '2026-09-22 00:00', '2026-09-22 00:00', 'Approuvée'),   -- a badgé : garde son statut
            ('E007', '2026-09-23 00:00', '2026-09-23 00:00', 'Approuvée')    -- mercredi férié : télétravail quand même"""))
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    m = dataclasses.replace(
        m, objects_schema="pt_test_conge", leave_table="demandeconge", leave_emp_col="matricule", leave_ref="matricule",
        leave_start_col="datedebut", leave_end_col="Datefin", leave_state_col="etat", leave_type_col="Document",
        tw_table="TdemandedeTeleTravail", tw_emp_col="matricule", tw_ref="matricule", tw_start_col="datedebut",
        tw_end_col="Datefin", tw_state_col="etat")
    assert (m.leave_state_values, m.leave_annual_values, m.leave_excep_values) == ("Approuvée", "CONGE", "CONGE EXCEP")
    with pg.begin() as c:
        c.execute(text("DROP SCHEMA IF EXISTS pt_test_conge CASCADE"))
    pointage.install(pg, m, "test")
    with pg.begin() as c:
        c.execute(text("INSERT INTO pt_test_conge.pointage_jours_feries (jour, libelle) VALUES (:d, 'Fête')"), {"d": WED})
    with pg.connect() as c:
        r = {(x.matricule, x.jour): x for x in c.execute(text(
            "SELECT * FROM pt_test_conge.f_pointage_journalier(:du, :au)"), {"du": MON, "au": WED}).mappings()}
    khady = r[("E005", MON)]
    assert (khady.statut, khady.statut_libelle) == ("CONGE_ANNUEL", "Congé Annuel")
    assert hm(khady.duree_validee) == "8h00" and hm(khady.duree_effective) == "8h00"
    assert r[("E005", TUE)].statut == "CONGE_ANNUEL" and r[("E005", WED)].statut == "NON_OUVRE"
    assert r[("E004", MON)].statut == "INCOMPLET"
    aminata = r[("E007", TUE)]
    assert (aminata.statut, aminata.statut_libelle) == ("CONGE_EXCEP", "Congé exceptionnel")
    assert r[("E001", TUE)].statut == "ABSENT"
    fatou = r[("E003", TUE)]  # congé non approuvé, mais télétravail approuvé
    assert (fatou.statut, fatou.statut_libelle, hm(fatou.duree_validee), hm(fatou.duree_effective)) == (
        "TELETRAVAIL", "Télétravail", "8h00", "8h00")
    assert r[("E005", TUE)].statut == "CONGE_ANNUEL" and r[("E002", TUE)].statut in ("A_L_HEURE", "RETARD")
    # Le télétravail prime sur les jours non ouvrés (le congé, lui, non) : 8h le mercredi férié.
    assert (r[("E007", WED)].statut, hm(r[("E007", WED)].duree_validee)) == ("TELETRAVAIL", "8h00")
    assert r[("E005", WED)].statut == "NON_OUVRE" and r[("E003", WED)].statut == "NON_OUVRE"

    # Écran : carte « En congé », badge vert clair, moyennes hors congés.
    from app.routers import suivi

    monkey_cfg = suivi.load_config
    cfg, _ = monkey_cfg(SessionLocal())
    suivi.load_config = lambda db: (cfg, m)
    try:
        page = logged_client.get(f"/suivi?date={TUE.isoformat()}").text
    finally:
        suivi.load_config = monkey_cfg
    assert "Autres situations" in page and 'badge st-leave">Congé Annuel' in page and "Congé exceptionnel" in page
    assert 'badge st-remote">Télétravail' in page

    # Administration : la section Congés propose les colonnes de la table choisie.
    admin = logged_client.get(f"/admin/pointage?conn_id={configured}&schema={SCHEMA}&punch_table=punchlog"
                              f"&emp_table=Employes&leave_table=demandeconge&leave_state_values=Approuvée"
                              f"&tw_table=TdemandedeTeleTravail").text
    assert "Table des demandes de congé" in admin and "Datefin (text)" in admin and "Congé exceptionnel" in admin
    tw_section = admin.split("Table des demandes de télétravail")[1]
    assert 'name="tw_start_col"' in tw_section and '<option value="datedebut" selected' in tw_section


def test_reference_whole_personnel(configured, pg, logged_client):
    """Liste de référence = table Personnel, sans colonne « actif » : tout le personnel est attendu."""
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    data = {**pointage.Mapping.from_json(original).__dict__, "reference": "person", "emp_active_col": "",
            "conn_id": conn_id}
    try:
        r = logged_client.post("/admin/pointage", data=data, follow_redirects=True)
        assert "installées" in r.text, r.text[:2000]
        day = rows(pg, MON)
        assert day[("108", MON)].statut == "ABSENT" and day[("108", MON)].nom == "Kane"  # sans fiche Employes
        assert day[("E006", MON)].statut == "ABSENT"  # plus de filtre « actif » : attendu lui aussi
        assert day[("E001", MON)].statut == "A_L_HEURE" and day[("E002", MON)].responsable == "Diallo Awa"
        page = logged_client.get("/admin/pointage").text
        assert "Tout le personnel de la table des noms" in page and "Personnes dans la liste" in page
    finally:
        restore = {**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=restore, follow_redirects=True).text


def test_check_employee_tool(configured, logged_client):
    """« Vérifier un employé » explique jour par jour pourquoi une personne est (ou n'est pas) absente."""
    period = f"&du={MON.isoformat()}&au={SAT.isoformat()}"
    inactive = logged_client.get(f"/admin/verifier?q=E006{period}").text
    assert "Gueye" in inactive and "Inactif : jamais compté absent" in inactive
    assert "Actif = « 0 »" in inactive and "non affiché" in inactive
    active = logged_client.get(f"/admin/verifier?q=Ba Khady{period}").text
    assert "Actif : attendu chaque jour ouvré" in active and active.count(">Absent<") == 4  # lun., mar., jeu., ven.
    assert "Jour non ouvré" in active  # mercredi férié, samedi
    unknown = logged_client.get(f"/admin/verifier?q=99{period}").text
    assert "Badges hors liste" in unknown and "Aucune personne" in unknown
    assert "Vérifier un employé" in logged_client.get(f"/suivi?date={MON.isoformat()}").text


def test_field_agents_never_absent_nor_late(configured, pg, logged_client):
    """Agents terrain (service entier ou employé) : « Sur le terrain », durée standard minimale, jamais absent/retard."""
    page = logged_client.get("/admin/terrain").text
    assert "Aucun agent terrain" in page and "Production" in page
    r = logged_client.post("/admin/terrain", data={"type": "employe", "matricule": "E002", "libelle": "Commercial"},
                           follow_redirects=True)
    assert "déclaré agent terrain" in r.text
    r = logged_client.post("/admin/terrain", data={"type": "service", "service": "RH"}, follow_redirects=True)
    assert "Ndiaye Moussa" in r.text and ">3<" in r.text  # RH : Sow, Fall, Ba
    assert "introuvable" not in r.text
    bad = logged_client.post("/admin/terrain", data={"type": "employe", "matricule": "ZZZ"}, follow_redirects=True).text
    assert "Aucun employé trouvé" in bad

    r = rows(pg, MON)
    moussa = r[("E002", MON)]  # arrivé à 8h10 : pas de retard ; 8h validées minimum, durée effective mesurée
    assert (moussa.statut, moussa.statut_libelle) == ("TERRAIN", "Sur le terrain") and moussa.terrain
    assert hm(moussa.duree_validee) == "8h00" and hm(moussa.duree_effective) == "7h05" and moussa.retard_min is None
    khady = r[("E005", MON)]  # aucun pointage : pas absente
    assert khady.statut == "TERRAIN" and hm(khady.duree_validee) == "8h00" and khady.duree_effective is None
    assert r[("E004", MON)].statut == "TERRAIN"  # un seul pointage : pas « incomplet »
    assert r[("E001", MON)].statut == "A_L_HEURE" and not r[("E001", MON)].terrain
    assert rows(pg, WED)[("E005", WED)].statut == "NON_OUVRE"
    assert r[("E007", MON)].retard_min == 15  # 7h45 pour un début à 7h30

    page = logged_client.get(f"/suivi?date={MON.isoformat()}").text
    assert "Sur le terrain" in page and 'badge st-field">Sur le terrain' in page

    for entry in pointage.field_entries(pg, pointage.Mapping.from_json(_cfg_data())):
        logged_client.post(f"/admin/terrain/{entry['id']}/delete")
    assert rows(pg, MON)[("E005", MON)].statut == "ABSENT"


def _cfg_data() -> str:
    with SessionLocal() as db:
        return db.query(PointageConfig).one().data


def test_reports_page_and_export(configured, logged_client):
    """Rapports : indicateurs, graphiques, services, alertes, employés, export Excel."""
    import io

    from openpyxl import load_workbook

    url = f"/rapports?p=perso&du={MON.isoformat()}&au={(MON + timedelta(days=4)).isoformat()}"
    page = logged_client.get(url)
    assert page.status_code == 200, page.text[:2000]
    html = page.text
    for part in ("Taux de présence", "Taux d'absentéisme", "Présence par jour", "Selon le jour de la semaine",
                 "Heures d'arrivée au bureau", "Par service", "À suivre", "Par employé", "Absences répétées"):
        assert part in html, part
    assert "Ba Khady" in html and "Production" in html and "RH" in html
    assert "<svg class=\"chart\"" in html and "data-tip=" in html
    # Tri, filtre service et période prédéfinie
    assert logged_client.get(url + "&tri=retards&service=RH").status_code == 200
    assert logged_client.get("/rapports?p=mois-1").status_code == 200
    r = logged_client.get("/rapports/export.xlsx?" + url.split("?", 1)[1])
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:5] == ["Employés", "Services", "Statuts du personnel", "Par jour", "Alertes"]
    rows = {row[0]: row for row in wb["Employés"].iter_rows(min_row=2, values_only=True)}
    assert rows["E005"][14] == 4  # Ba Khady : 4 absences (mercredi férié)
    assert any(row[0] == "Absences répétées" for row in wb["Alertes"].iter_rows(min_row=2, values_only=True))


def _pdf_text(content: bytes) -> str:
    """Texte brut des flux (compressés) d'un PDF généré par fpdf2."""
    import re
    import zlib

    assert content.startswith(b"%PDF")
    out = []
    for stream in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", content, re.S):
        try:
            out.append(zlib.decompress(stream).decode("latin-1"))
        except zlib.error:
            out.append(stream.decode("latin-1"))
    return "\n".join(out)


def test_pdf_exports(configured, logged_client):
    """Exports PDF du suivi et des rapports : mêmes filtres que l'écran, durée standard masquable."""
    r = logged_client.get(f"/suivi/export.pdf?date={MON.isoformat()}&statut=RETARD&statut=ABSENT")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert f'suivi_{MON:%Y%m%d}.pdf' in r.headers["content-disposition"]
    body = _pdf_text(r.content)
    assert "Suivi journalier des pointages" in body and "Moyenne" in body and "En retard" in body
    assert "(En retard, Absent) Tj" in body and "(STANDARD) Tj" in body and "Ba Khady" in body
    assert "E001" not in body  # à l'heure : exclu par le filtre des statuts
    assert "Export PDF" in logged_client.get(f"/suivi?date={MON.isoformat()}").text
    try:
        assert logged_client.get("/suivi/affichage?validee=0", follow_redirects=False).status_code in (302, 303)
        hidden = _pdf_text(logged_client.get(f"/suivi/export.pdf?date={MON.isoformat()}").content)
        assert "(STANDARD) Tj" not in hidden and "(EFFECTIVE) Tj" in hidden
    finally:
        logged_client.get("/suivi/affichage?validee=1", follow_redirects=False)

    q = f"p=perso&du={MON.isoformat()}&au={(MON + timedelta(days=4)).isoformat()}"
    assert "Export PDF" in logged_client.get("/rapports?" + q).text
    r = logged_client.get("/rapports/export.pdf?" + q)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    body = _pdf_text(r.content)
    for part in ("Rapport de présence", "TAUX DE PRÉSENCE", "Par service", "Par employé", "Employés à suivre",
                 "Absences répétées", "Ba Khady"):
        assert part in body, part


def test_category_filter(configured, pg, logged_client):
    """Catégorie du personnel (via une table de libellés) : colonne, filtre du suivi et des rapports, exports."""
    import io

    from openpyxl import load_workbook

    r = rows(pg, MON)
    assert r[("E001", MON)].categorie == "Cadre" and r[("E002", MON)].categorie == "Non cadre"
    assert r[("99", MON)].categorie is None  # hors liste
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    assert pointage.categories(pg, m) == ["Cadre", "Non cadre"]
    page = logged_client.get(f"/suivi?date={MON.isoformat()}&categorie=Cadre").text
    assert 'name="categorie"' in page and "Statut du personnel : Cadre" in page
    assert "Diallo" in page and "Sow" in page and "Ndiaye" not in page
    xlsx = load_workbook(io.BytesIO(logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}&categorie=Cadre").content))
    ws = xlsx.active
    assert ws.cell(1, 6).value == "Statut du personnel" and {row[5] for row in ws.iter_rows(min_row=2, values_only=True) if row[0] not in (None, "Moyenne")} == {"Cadre"}
    url = f"/rapports?p=perso&du={MON.isoformat()}&au={(MON + timedelta(days=4)).isoformat()}"
    rep = logged_client.get(url).text
    assert "Par statut du personnel" in rep and "Non cadre" in rep
    assert "Ba Khady" not in logged_client.get(url + "&categorie=Cadre").text
    wb = load_workbook(io.BytesIO(logged_client.get("/rapports/export.xlsx?" + url.split("?", 1)[1]).content))
    assert "Statuts du personnel" in wb.sheetnames
    admin = logged_client.get("/admin/pointage").text
    assert "Statut du personnel" in admin and "<code>Cadre</code>" in admin
    # Choix multiples : plusieurs services et plusieurs statuts du personnel.
    page = logged_client.get(f"/suivi?date={MON.isoformat()}&categorie=Cadre&categorie=Non+cadre&service=RH").text
    assert "Statut du personnel : Cadre, Non cadre" in page and "Service : RH" in page and 'class="ms-more">+1' in page
    assert "E003</td>" in page and "E005</td>" in page and "E001</td>" not in page  # E001 : Production
    both = logged_client.get(f"/suivi?date={MON.isoformat()}&service=RH&service=Production&pop=liste").text
    assert "E001</td>" in both and "E003</td>" in both
    xlsx = load_workbook(io.BytesIO(logged_client.get(
        f"/suivi/export.xlsx?date={MON.isoformat()}&service=RH&service=Production").content))
    assert {row[4] for row in xlsx.active.iter_rows(min_row=2, values_only=True)
            if row[0] not in (None, "Moyenne")} == {"RH", "Production"}
    assert dict(xlsx["Filtres"].iter_rows(values_only=True))["Service"] == "RH, Production"
    multi = logged_client.get(url + "&service=RH&service=Production&categorie=Cadre").text
    assert "Par employé" in multi and "E005</td>" not in multi and "E001</td>" in multi and "E003</td>" in multi
    # Colonne non configurée : le filtre reste visible (désactivé), avec le lien de configuration.
    import app.pointage as pt
    original = pt.categories
    pt.categories = lambda *a, **k: []
    try:
        for url in (f"/suivi?date={MON.isoformat()}", "/rapports?p=7j"):
            html = logged_client.get(url).text
            assert "Statut du personnel" in html and "Non configuré" in html and "#statut-personnel" in html, url
    finally:
        pt.categories = original


def test_badge_mails_test_mode_production_and_no_duplicates(configured, pg, logged_client, monkeypatch):
    """Mails de badge : jamais d'envoi aux employés hors production, un seul mail par pointage, modèle et essai."""
    from app import mails
    from app.models import MailLog, MailSettings, MailSubscriber

    now = datetime.now().replace(second=0, microsecond=0)

    def punch(emp_id, when):
        with pg.begin() as c:
            c.execute(text(f"INSERT INTO {SCHEMA}.punchlog (employe_id, date_pointage, heure_pointage, terminal) "
                           f"VALUES (:e, :d, :h, 'MAIL')"), {"e": emp_id, "d": when.date(), "h": when.strftime("%H%M%S")})

    try:
        # Activation refusée sans adresse de test ; production refusée sans confirmation explicite.
        base = {"enabled": "1", "mode": "test", "smtp_host": "smtp.exemple.com", "smtp_port": "587",
                "smtp_security": "starttls", "from_email": "pointage@exemple.com", "from_name": "Pointage",
                "company": "Société <Test>", "max_per_run": "50"}
        r = logged_client.post("/admin/mails/parametres", data=base, follow_redirects=True)
        assert "au moins une adresse de test" in r.text
        r = logged_client.post("/admin/mails/parametres", data={**base, "test_recipients": "rh@exemple.com"},
                               follow_redirects=True)
        assert "Paramètres des mails enregistrés" in r.text and "Mode test" in r.text
        with SessionLocal() as db:
            s = db.query(MailSettings).one()
            assert s.enabled and s.mode == "test" and s.since is not None
            s.since = now - timedelta(hours=3)  # pour le test : pointages des 3 dernières heures
            db.commit()

        # Abonnés : Awa (adresse saisie ici) et Moussa (sans adresse).
        r = logged_client.post("/admin/mails/abonnes", data={"sub": ["1", "2"], "email_1": "awa@exemple.com",
                                                               "email_2": ""}, follow_redirects=True)
        assert "2 abonné(s) enregistré(s)" in r.text
        assert "awa@exemple.com" in logged_client.get("/admin/mails/abonnes").text

        punch(1, now - timedelta(hours=2))
        punch(2, now - timedelta(hours=2, minutes=5))
        punch(3, now - timedelta(hours=2))            # Fatou : non abonnée
        punch(1, now - timedelta(hours=5))            # avant l'activation : jamais notifié
        sent = []
        result = mails.process(sender=sent.append)
        assert result["sent"] == 2 and result["mode"] == "test"
        # Mode test : tout part vers l'adresse de test, aucun mail vers un employé.
        assert {m["To"] for m in sent} == {"rh@exemple.com"}
        assert all(m["Subject"].startswith("[TEST] ") for m in sent)
        awa = next(m for m in sent if "awa@exemple.com" in m.get_body(("html",)).get_content())
        body = awa.get_body(("html",)).get_content()
        assert "MODE TEST" in body and "Bonjour Awa" in body and "Société &lt;Test&gt;" in body
        assert "awa@exemple.com" not in awa["To"]

        # Deuxième passage : rien de nouveau, aucun doublon.
        assert mails.process(sender=sent.append)["sent"] == 0 and len(sent) == 2

        # Nouveau pointage d'Awa : un seul mail, avec ses pointages du jour (sauf changement de date).
        punch(1, now - timedelta(minutes=30))
        more = []
        assert mails.process(sender=more.append)["sent"] == 1
        assert "Awa" in more[0].get_body(("html",)).get_content()

        # Passage en production : confirmation obligatoire, puis envoi réel à l'adresse de l'employé.
        prod = {**base, "test_recipients": "rh@exemple.com", "mode": "production"}
        r = logged_client.post("/admin/mails/parametres", data=prod, follow_redirects=True)
        assert "cochez la confirmation" in r.text
        r = logged_client.post("/admin/mails/parametres", data={**prod, "confirm_production": "1"}, follow_redirects=True)
        assert "Mode PRODUCTION activé" in r.text
        with SessionLocal() as db:
            s = db.query(MailSettings).one()
            assert s.mode == "production" and s.since >= now - timedelta(minutes=1)
            s.since = now - timedelta(minutes=20)
            db.commit()
        punch(1, now - timedelta(minutes=10))
        punch(2, now - timedelta(minutes=10))
        real = []
        result = mails.process(sender=real.append)
        assert result["sent"] == 1 and result["skipped"] == 1  # Moussa : aucune adresse
        assert real[0]["To"] == "awa@exemple.com" and not real[0]["Subject"].startswith("[TEST]")
        with SessionLocal() as db:
            skipped = db.query(MailLog).filter(MailLog.status == "skipped").one()
            assert skipped.emp_key == "2" and "aucune adresse" in skipped.error
        journal = logged_client.get("/admin/mails/journal").text
        assert "awa@exemple.com" in journal and "production" in journal and "ignoré" in journal

        # Modèle : aperçu avec valeurs échappées, enregistrement, modèle par défaut.
        page = logged_client.get("/admin/mails/modele").text
        assert "tpl_html" in page and "{{prenom}}" in page
        prev = logged_client.post("/admin/mails/apercu", data={"html": "<p>Bonjour {{prenom}} ({{entreprise}})</p>"}).text
        assert prev == "<p>Bonjour Aminata (Société &lt;Test&gt;)</p>"
        r = logged_client.post("/admin/mails/modele", data={"subject": "Badge {{heure}}", "html": "<p>{{prenom}}</p>"},
                               follow_redirects=True)
        assert "Modèle enregistré" in r.text
        r = logged_client.post("/admin/mails/modele", data={"reset": "1"}, follow_redirects=True)
        assert "Modèle par défaut rétabli" in r.text

        # Mail d'essai via le serveur SMTP (simulé).
        delivered = []

        class FakeSMTP:
            def __init__(self, host, port, timeout=None):
                self.host = host
            def ehlo(self): pass
            def starttls(self, context=None): pass
            def login(self, user, pwd): pass
            def send_message(self, msg): delivered.append(msg)
            def quit(self): pass

        monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)
        r = logged_client.post("/admin/mails/essai", data={"to": "moi@exemple.com"}, follow_redirects=True)
        assert "essai envoyé à moi@exemple.com" in r.text
        assert delivered[0]["To"] == "moi@exemple.com"
        assert delivered[0]["Subject"].startswith("[TEST] ")
        assert "Mails de badge" in logged_client.get("/admin/mails").text
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.punchlog WHERE terminal = 'MAIL'"))
        with SessionLocal() as db:
            db.query(MailLog).delete()
            db.query(MailSubscriber).delete()
            s = db.query(MailSettings).first()
            if s:
                s.enabled, s.mode = False, "test"
            db.commit()


def test_field_agents_excel_import(configured, pg, logged_client):
    """Agents terrain importés d'un fichier Excel ou CSV : ajout, remplacement, matricules inconnus, modèle, export."""
    import io

    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Nom employé", "Matricule", "Motif"])          # « Nom employé » ne doit pas être pris pour le matricule
    ws.append(["Ndiaye Moussa", "E002", "Commercial Nord"])
    ws.append(["Fall Ibou", "E004", "Commercial Sud"])
    ws.append(["Inconnu", "ZZZ9", "?"])
    ws.append(["Ndiaye Moussa", "E002", ""])                  # doublon
    buf = io.BytesIO()
    wb.save(buf)
    r = logged_client.post("/admin/terrain/import", files={"fichier": ("commerciaux.xlsx", buf.getvalue())},
                           data={"mode": "ajouter"}, follow_redirects=True)
    assert "2 agent(s) ajouté(s)" in r.text and "1 matricule(s) introuvable(s)" in r.text and "ZZZ9" in r.text
    assert "Commercial Nord" in r.text and "Commercial Sud" in r.text
    assert rows(pg, MON)[("E002", MON)].statut == "TERRAIN"

    # CSV sans en-tête, séparateur « ; », accents Windows : remplacement de la liste des employés.
    csv = "E007;Délégué médical\r\n".encode("cp1252")
    r = logged_client.post("/admin/terrain/import", files={"fichier": ("liste.csv", csv)},
                           data={"mode": "remplacer"}, follow_redirects=True)
    assert "1 agent(s) ajouté(s)" in r.text and "2 retiré(s)" in r.text and "Délégué médical" in r.text
    assert rows(pg, MON)[("E002", MON)].statut == "RETARD" and rows(pg, MON)[("E007", MON)].statut == "TERRAIN"

    bad = logged_client.post("/admin/terrain/import", files={"fichier": ("vieux.xls", b"...")}, follow_redirects=True)
    assert "format .xlsx" in bad.text
    empty = logged_client.post("/admin/terrain/import", files={"fichier": ("vide.csv", b"\r\n")}, follow_redirects=True)
    assert "vide" in empty.text

    model = load_workbook(io.BytesIO(logged_client.get("/admin/terrain/modele.xlsx").content)).active
    assert [c.value for c in model[1]] == ["Matricule", "Motif"]
    export = load_workbook(io.BytesIO(logged_client.get("/admin/terrain/export.xlsx").content)).active
    assert list(export.iter_rows(min_row=2, values_only=True)) == [("E007", "Délégué médical", "Diop Aminata")]
    assert "Importer une liste depuis Excel" in logged_client.get("/admin/terrain").text

    for entry in pointage.field_entries(pg, pointage.Mapping.from_json(_cfg_data())):
        logged_client.post(f"/admin/terrain/{entry['id']}/delete")


def test_average_punches_and_duration_colours(configured, logged_client):
    """Moyennes du 1er et du dernier pointage ; durée standard en vert (objectif atteint) ou rouge ; export."""
    import io

    from openpyxl import load_workbook

    # Lundi, employés de la liste au bureau : 1ers pointages 07h20, 08h10, 07h35, 07h40, 07h45 → 07h42 ;
    # derniers (2 pointages ou plus) 17h10, 16h45, 15h50, 12h00 → 15h26.
    page = logged_client.get(f"/suivi?date={MON.isoformat()}&pop=liste").text
    assert "Pointages moyens" in page and "07h42" in page and "15h26" in page
    assert 'class="avg-row"' in page and "Moyenne" in page
    # Objectif par défaut 8h : Awa (7h30) et Moussa (6h50) en rouge.
    assert page.count('class="dur dur-ko"') >= 2 and 'class="dur dur-ok"' in page  # légende
    values = {k: v for k, _, _, v in pointage.PARAMS if k != "jours_ouvres"}
    r = logged_client.post("/admin/parametres", data={**values, "objectif_duree": "07:00",
                           "jours_ouvres": ["1", "2", "3", "4", "5"], "date_effet": MON.isoformat()}, follow_redirects=True)
    assert "paramètre(s) modifié(s)" in r.text
    page = logged_client.get(f"/suivi?date={MON.isoformat()}&pop=liste&q=E001").text
    assert 'dur dur-ok"' in page.split("<tbody>")[1] and "≥ 7h00" in page
    xlsx = load_workbook(io.BytesIO(logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}&pop=liste").content))
    ws = xlsx["Suivi journalier"]
    fills = {ws.cell(i, 2).value: ws.cell(i, 11).fill.fgColor.rgb for i in range(2, ws.max_row + 1) if ws.cell(i, 11).value}
    assert fills["E001"].endswith("C6EFCE") and fills["E002"].endswith("FFC7CE")  # 7h30 ≥ 7h ; 6h50 < 7h
    # Durée effective colorée de la même façon : Awa 8h20 ≥ 7h (vert), Aminata 4h15 < 7h (rouge).
    eff = {ws.cell(i, 2).value: ws.cell(i, 12).fill.fgColor.rgb for i in range(2, ws.max_row + 1) if ws.cell(i, 12).value}
    assert eff["E001"].endswith("C6EFCE") and eff["E007"].endswith("FFC7CE")
    cells = page.split("<tbody>")[1].split("</tbody>")[0]
    assert cells.count('class="dur dur-ok"') == 2  # validée 7h30 et effective 8h20 d'Awa
    last = [c.value for c in ws[ws.max_row]]
    assert last[0] == "Moyenne" and last[6].strftime("%H:%M") == "07:42" and last[8].strftime("%H:%M") == "15:26"
    logged_client.post("/admin/parametres", data={**values, "objectif_duree": "08:00",
                       "jours_ouvres": ["1", "2", "3", "4", "5"], "date_effet": MON.isoformat()})


def test_sick_leave_workflow_and_hr(configured, pg, logged_client):
    """Arrêt maladie : déclaration par l'employé (justificatif obligatoire), validation N+1 puis RH, statut dans les
    calculs ; saisie RH validée d'office ; refus, annulation, droits d'accès au justificatif, circuit paramétrable."""
    from fastapi.testclient import TestClient

    from app.main import app
    from app.models import SickLeave, SickLeaveAction, User

    THU, FRI = date(2026, 9, 24), date(2026, 9, 25)
    pdf = b"%PDF-1.4 arret de travail"
    for data in ({"username": "moussa", "role": "lecteur", "emp_matricule": "E002"},
                 {"username": "awa", "role": "manager", "emp_matricule": "E001"},
                 {"username": "rh1", "role": "lecteur", "sick_leave_hr": "true"},
                 {"username": "fatou", "role": "lecteur", "emp_matricule": "E003"}):
        r = logged_client.post("/admin/users/save", data={**data, "password": "motdepasse1"}, follow_redirects=True)
        assert "enregistré" in r.text, r.text[:500]
    r = logged_client.post("/admin/arrets", data={"type": ["responsable", "rh"], "user": ["", ""]}, follow_redirects=True)
    assert "Responsable N+1 de l&#39;employé → Agent RH habilité" in r.text

    def client(name):
        c = TestClient(app)
        c.post("/login", data={"username": name, "password": "motdepasse1"})
        return c

    try:
        moussa, awa, rh1, fatou = client("moussa"), client("awa"), client("rh1"), client("fatou")
        page = moussa.get("/arrets").text
        assert "Déclarer un arrêt maladie" in page and "Responsable N+1 de l&#39;employé → Agent RH habilité" in page
        r = moussa.post("/arrets", data={"du": THU.isoformat(), "au": FRI.isoformat()}, follow_redirects=True)
        assert "Joignez le justificatif" in r.text
        r = moussa.post("/arrets", data={"du": THU.isoformat(), "au": FRI.isoformat()},
                        files={"justificatif": ("faux.pdf", b"pas un pdf")}, follow_redirects=True)
        assert "ne correspond pas" in r.text
        r = moussa.post("/arrets", data={"du": THU.isoformat(), "au": FRI.isoformat(), "matricule": "E005",
                                         "commentaire": "Grippe"},
                        files={"justificatif": ("arret.pdf", pdf)}, follow_redirects=True)
        assert "transmis pour validation" in r.text and "Ndiaye Moussa" in r.text  # pour lui-même, pas E005
        with SessionLocal() as db:
            leave = db.query(SickLeave).one()
            assert (leave.matricule, leave.manager_matricule, leave.status) == ("E002", "E001", "en_attente")
            lid = leave.id
        assert "Valider" not in moussa.get(f"/arrets/{lid}").text  # pas de validation de son propre arrêt
        assert moussa.post(f"/arrets/{lid}/decision", data={"decision": "valider"}, follow_redirects=True).status_code == 200
        assert rows(pg, THU)[("E002", THU)].statut == "ABSENT"  # en attente : pas encore d'effet

        # Étape 1 : le responsable N+1 (Awa) ; étape 2 : les RH.
        assert "À valider" in awa.get("/arrets").text and 'class="nav-badge"' in awa.get("/arrets").text
        assert "À valider" not in rh1.get("/arrets").text
        r = awa.post(f"/arrets/{lid}/decision", data={"decision": "valider", "commentaire": "OK"}, follow_redirects=True)
        assert "passe à l&#39;étape suivante" in r.text
        assert "À valider" in rh1.get("/arrets").text
        r = rh1.post(f"/arrets/{lid}/decision", data={"decision": "valider"}, follow_redirects=True)
        assert "Arrêt validé" in r.text
        thu = rows(pg, THU)[("E002", THU)]
        assert (thu.statut, thu.statut_libelle, hm(thu.duree_validee)) == ("ARRET_MALADIE", "Arrêt maladie", "8h00")
        assert rows(pg, FRI)[("E002", FRI)].statut == "ARRET_MALADIE"
        assert 'badge st-sick">Arrêt maladie' in logged_client.get(f"/suivi?date={THU.isoformat()}").text

        # Justificatif : l'employé, son responsable et les RH ; pas un autre employé.
        doc = moussa.get(f"/arrets/{lid}/justificatif")
        assert doc.status_code == 200 and doc.content == pdf and doc.headers["content-type"] == "application/pdf"
        assert awa.get(f"/arrets/{lid}/justificatif").status_code == 200
        assert fatou.get(f"/arrets/{lid}/justificatif").status_code == 404
        assert "Ndiaye" not in fatou.get("/arrets").text

        # Chevauchement refusé ; refus motivé obligatoire.
        r = moussa.post("/arrets", data={"du": FRI.isoformat(), "au": FRI.isoformat()},
                        files={"justificatif": ("a.png", b"\x89PNG....")}, follow_redirects=True)
        assert "couvre déjà" in r.text
        r = moussa.post("/arrets", data={"du": "2026-10-05", "au": "2026-10-06"},
                        files={"justificatif": ("b.jpg", b"\xff\xd8\xff....")}, follow_redirects=True)
        with SessionLocal() as db:
            second = db.query(SickLeave).filter(SickLeave.id != lid).one().id
        assert "motif du refus" in awa.post(f"/arrets/{second}/decision", data={"decision": "refuser"},
                                             follow_redirects=True).text
        r = awa.post(f"/arrets/{second}/decision", data={"decision": "refuser", "commentaire": "Dates erronées"},
                     follow_redirects=True)
        assert "Arrêt refusé" in r.text and "Dates erronées" in r.text

        # Saisie RH : n'importe quel employé, sans justificatif, validée d'office.
        r = rh1.post("/arrets", data={"matricule": "E005", "du": THU.isoformat(), "au": THU.isoformat()},
                     follow_redirects=True)
        assert "enregistré et validé" in r.text and rows(pg, THU)[("E005", THU)].statut == "ARRET_MALADIE"

        # Annulation d'un arrêt validé : RH seulement ; les jours redeviennent absents.
        assert "RH" in moussa.post(f"/arrets/{lid}/annuler", data={}, follow_redirects=True).text
        r = rh1.post(f"/arrets/{lid}/annuler", data={"commentaire": "Erreur"}, follow_redirects=True)
        assert "Arrêt annulé" in r.text and rows(pg, THU)[("E002", THU)].statut == "ABSENT"
        with SessionLocal() as db:
            actions = [a.action for a in db.query(SickLeaveAction).filter(SickLeaveAction.leave_id == lid)
                       .order_by(SickLeaveAction.id)]
        assert actions == ["declare", "valide", "valide", "annule"]

        # Circuit : un utilisateur désigné doit être choisi ; sans étape, validé dès la déclaration.
        bad = logged_client.post("/admin/arrets", data={"type": ["utilisateur"], "user": [""]}, follow_redirects=True)
        assert "Choisissez l&#39;utilisateur" in bad.text
        r = logged_client.post("/admin/arrets", data={}, follow_redirects=True)
        assert "aucune étape" in r.text
        r = moussa.post("/arrets", data={"du": "2026-10-12", "au": "2026-10-12"},
                        files={"justificatif": ("c.pdf", pdf)}, follow_redirects=True)
        assert "Validé" in r.text
        for c in (moussa, awa, rh1, fatou):
            c.close()
    finally:
        S = pointage.qi(pointage.Mapping.from_json(_cfg_data()).objs)
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {S}.pointage_arrets_maladie"))
        with SessionLocal() as db:
            db.query(SickLeaveAction).delete()
            db.query(SickLeave).delete()
            db.query(User).filter(User.username.in_(["moussa", "awa", "rh1", "fatou"])).delete()
            db.commit()
        logged_client.post("/admin/arrets", data={"type": ["responsable", "rh"], "user": ["", ""]})


def _planning_xlsx(days: dict) -> bytes:
    """Planning au format « PLANNING ATF » : légende, ligne MATRICULE, ligne NOM, une ligne par jour."""
    from datetime import time as t

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "PLANNING"
    ws["C1"] = "VEUILLEZ SAISIR LE CODE DE CHAQUE PLANNING HORAIRE DEVANT LE JOUR DE CHAQUE COLLABORATEUR"
    for col, value in zip("BCDFGH", ["Code", "Heure Début", "Heure Fin", "Code", "Heure Début", "Heure Fin"]):
        ws[f"{col}5"] = value
    legend = [("P1", t(6), t(14), "P5", "ASTREINTE"), ("P2", t(14), t(22), "P6", "FERIE PAYE"),
              ("P3", t(22), t(6), "P8", "FORMATION"), ("P4", "REPOS", None, "P9", "FORMATION CONTINUE")]
    for i, (c1, d1, f1, c2, l2) in enumerate(legend, start=6):
        ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"F{i}"], ws[f"G{i}"] = c1, d1, f1, c2, l2
    ws["J5"], ws["K5"], ws["L5"] = "Code", "Heure Début", "Heure Fin"
    ws["J6"], ws["K6"], ws["L6"] = "0618", t(6), t(18)
    ws["J7"], ws["K7"], ws["L7"] = "1806", t(18), t(6)
    ws.append([])
    ws.append(["MATRICULE", "E005", "X999"])
    ws.append(["NOM", "BA KHADY     ", "INCONNU"])
    ws.append(["DATE"])
    for day, (code, other) in days.items():
        ws.append([datetime(day.year, day.month, day.day), code, other])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_shift_planning_import_and_calculations(configured, pg, logged_client):
    """Horaires postés : nuit 18h-6h rattachée au jour de début, repos, retard sur l'horaire du poste, absence un
    dimanche planifié, congé du planning ; import du fichier Excel tel quel (618 = 0618, légende, inconnus)."""
    THU, FRI, SUN, MON2 = date(2026, 9, 24), date(2026, 9, 25), date(2026, 9, 27), date(2026, 9, 28)
    S = f"{SCHEMA}"
    punches = [(5, THU, "17:55"), (5, FRI, "06:05"), (5, SAT, "15:00"), (5, SAT, "22:00")]
    try:
        with pg.begin() as c:
            for emp, day, hhmm in punches:
                c.execute(text(f"INSERT INTO {S}.punchlog (employe_id, date_pointage, heure_pointage, terminal) "
                               f"VALUES (:e, :d, :h, 'PLAN')"), {"e": emp, "d": day, "h": hhmm.replace(":", "") + "00"})
        # Codes par défaut installés avec les calculs (légende du planning ATF).
        page = logged_client.get("/admin/planning")
        assert page.status_code == 200 and "1806" in page.text and "Aucun planning" in page.text

        data = _planning_xlsx({THU: (1806, "P1"), FRI: ("P4", "P1"), SAT: ("P2", "P1"), SUN: ("P1", None),
                               MON2: ("P11", None), date(2026, 9, 29): ("ZZ", None)})
        r = logged_client.post("/admin/planning/import", files={"fichier": ("PLANNING ATF.xlsx", data)},
                               follow_redirects=True)
        assert "importé du 24/09/2026 au 29/09/2026 : 1 employé(s), 5 jour(s)" in r.text, r.text[:3000]
        assert "P9" in r.text and "Nouveaux codes" in r.text          # code de la légende créé
        assert "« ZZ » (1 j)" in r.text and "X999" in r.text           # code et matricule inconnus signalés
        assert "BA" in r.text.upper() and 'pl pl-nuit">1806' in r.text
        assert "28 h" in r.text  # heures planifiées : 12 + 8 + 8 (repos et congé non comptés)
        # Période libre : longue plage, plage vide (lien vers le planning importé), plage trop longue ramenée à un an.
        page = logged_client.get("/admin/planning?du=2026-01-01&au=2026-12-31").text
        assert "Planning du 01/01/2026 au 31/12/2026" in page and "365 jour(s)" in page and "28 h" in page
        page = logged_client.get("/admin/planning?du=2007-11-21&au=2008-01-22").text
        assert "Aucun planning sur cette période" in page and "du=2026-09-24&au=2026-09-28" in page
        page = logged_client.get("/admin/planning?du=2020-01-01&au=2026-12-31").text
        assert "Période ramenée à 366 jours" in page and "Tout le planning" in page and 'id="pl_expand"' not in page

        got = rows(pg, THU, MON2)
        night = got[("E005", THU)]
        assert night.statut == "A_L_HEURE" and night.nb_pointages == 2 and night.poste == "1806 · 18h00–06h00"
        assert hm(night.premier_pointage) == "17h55" and night.dernier_pointage.date() == FRI
        assert hm(night.duree_validee) == "12h00" and hm(night.duree_effective) == "12h10"
        rest = got[("E005", FRI)]
        assert rest.statut == "REPOS" and rest.nb_pointages == 0 and not rest.jour_ouvre
        assert hm(rest.duree_validee) == "8h00" and rest.duree_effective is None  # repos crédité de 8h validées
        late = got[("E005", SAT)]  # 15h00 pour un poste de 14h (tolérance 45 min) : retard de 60 min
        assert late.statut == "RETARD" and late.retard_min == 60 and hm(late.duree_validee) == "7h00" and late.jour_ouvre
        assert got[("E005", SUN)].statut == "ABSENT" and got[("E005", SUN)].jour_ouvre
        leave = got[("E005", MON2)]
        assert leave.statut == "CONGE_ANNUEL" and hm(leave.duree_validee) == "8h00"
        assert got[("E004", SUN)].statut == "NON_OUVRE" and got[("E004", SUN)].poste is None  # non planifié : bureau

        # Détail de la nuit : le départ du lendemain est affiché ; suivi et rapports affichent le poste.
        detail = logged_client.get(f"/suivi/detail?key=5&jour={THU.isoformat()}").text
        assert "17:55" in detail and "06:05" in detail
        page = logged_client.get(f"/suivi?du={THU.isoformat()}&au={SAT.isoformat()}&q=E005").text
        assert "1806 · 18h00–06h00" in page and "+1 j" in page and "Repos (planning)" in page
        assert logged_client.get(f"/rapports?du={THU.isoformat()}&au={MON2.isoformat()}").status_code == 200

        # Codes de poste : modification, refus de suppression d'un code utilisé, export réimportable.
        r = logged_client.post("/admin/planning/postes", data={"code": "p7", "libelle": "Journée", "type": "travail",
                                                               "debut": "08:00", "fin": "08:00"}, follow_redirects=True)
        assert "heure de début et une heure de fin différentes" in r.text
        r = logged_client.post("/admin/planning/postes", data={"code": "p7", "libelle": "Journée", "type": "travail",
                                                               "debut": "08:00", "fin": "16:30", "pause": "0h30"},
                               follow_redirects=True)
        assert "Code « P7 » enregistré" in r.text
        assert "utilisé par" in logged_client.post("/admin/planning/postes/P4/delete", follow_redirects=True).text
        assert "Code « P7 » supprimé" in logged_client.post("/admin/planning/postes/P7/delete", follow_redirects=True).text
        export = logged_client.get(f"/admin/planning/export.xlsx?du={THU.isoformat()}&au={MON2.isoformat()}")
        plan = pointage.read_planning_file("x.xlsx", export.content)
        assert ("E005", THU, "1806") in plan["entries"] and ("E005", FRI, "P4") in plan["entries"]

        # Suppression : les jours reviennent à l'horaire de bureau.
        r = logged_client.post("/admin/planning/supprimer", data={"du": THU.isoformat(), "au": MON2.isoformat()},
                               follow_redirects=True)
        assert "5 jour(s) planifié(s) supprimé(s)" in r.text
        assert rows(pg, FRI)[("E005", FRI)].statut != "REPOS"
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {S}.punchlog WHERE terminal = 'PLAN'"))
            c.execute(text(f"DELETE FROM {S}.pointage_planning"))
            c.execute(text(f"DELETE FROM {S}.pointage_postes WHERE code IN ('P9', 'P7')"))


def test_parse_planning_rejects_unknown_layout():
    with pytest.raises(pointage.PointageError, match="MATRICULE"):
        pointage.parse_planning([["a", "b"], [1, 2]])
    plan = pointage.parse_planning([["MATRICULE", 590128.0], ["NOM", "KARAMOKO"], ["19/08/2026", 618]])
    assert plan["entries"] == [("590128", date(2026, 8, 19), "618")] and plan["names"] == {"590128": "KARAMOKO"}


def test_sick_leave_email_notifications(configured, pg, logged_client, monkeypatch):
    """Mails des arrêts maladie : valideur de l'étape en cours, puis l'employé à la décision ; mode test redirigé,
    auteur de l'action exclu, adresse du compte ou de la fiche employé, désactivé par défaut."""
    from fastapi.testclient import TestClient

    from app import arrets, mails
    from app.main import app
    from app.models import MailSettings, SickLeave, SickLeaveAction, User

    delivered, results = [], []

    class FakeSMTP:
        def __init__(self, host, port, timeout=None): pass
        def ehlo(self): pass
        def starttls(self, context=None): pass
        def login(self, user, pwd): pass
        def send_message(self, msg): delivered.append(msg)
        def quit(self): pass

    monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)
    monkeypatch.setattr(arrets, "start_notify", lambda *a: results.append(arrets.notify(*a)))
    for data in ({"username": "n_moussa", "emp_matricule": "E002", "email": "moussa@exemple.com"},
                 {"username": "n_awa", "emp_matricule": "E001", "email": "awa@exemple.com"},
                 {"username": "n_rh", "sick_leave_hr": "true", "email": "rh@exemple.com"},
                 {"username": "n_bad", "email": "pas-une-adresse"}):
        r = logged_client.post("/admin/users/save", data={**data, "role": "lecteur", "password": "motdepasse1"},
                               follow_redirects=True)
        assert ("invalide" if data["username"] == "n_bad" else "enregistré") in r.text
    logged_client.post("/admin/arrets", data={"type": ["responsable", "rh"], "user": ["", ""]})
    with SessionLocal() as db:
        s = mails.get_settings(db)
        s.smtp_host, s.from_email, s.mode, s.test_recipients = "smtp.test", "pointage@exemple.com", "test", "qa@exemple.com"
        db.commit()

    def client(name):
        c = TestClient(app)
        c.post("/login", data={"username": name, "password": "motdepasse1"})
        return c

    pdf = b"%PDF-1.4 arret"
    try:
        moussa, awa, rh = client("n_moussa"), client("n_awa"), client("n_rh")
        moussa.post("/arrets", data={"du": "2026-10-05", "au": "2026-10-06"}, files={"justificatif": ("a.pdf", pdf)})
        assert results[-1]["status"] == "disabled" and not delivered  # désactivé par défaut

        page = logged_client.post("/admin/arrets/notifications", data={"notify": "true"}, follow_redirects=True).text
        assert "Notifications par e-mail activées" in page and "smtp.test" in page and "qa@exemple.com" in page
        with SessionLocal() as db:
            db.query(SickLeave).delete()
            db.commit()

        # Déclaration : le responsable N+1 est prévenu ; mode test → redirigé vers l'adresse de test.
        moussa.post("/arrets", data={"du": "2026-10-05", "au": "2026-10-06", "commentaire": "Grippe"},
                    files={"justificatif": ("a.pdf", pdf)})
        assert results[-1] == {"status": "sent", "sent": 1, "to": ["awa@exemple.com"]}
        msg = delivered[-1]
        assert msg["To"] == "qa@exemple.com" and msg["Subject"] == "[TEST] Arrêt maladie à valider — Ndiaye Moussa"
        body = msg.get_body(("html",)).get_content()
        assert "MODE TEST" in body and "awa@exemple.com" in body and "Grippe" in body and "/arrets/" in body
        with SessionLocal() as db:
            lid = db.query(SickLeave).one().id

        # Étape 1 validée par Awa → les RH ; étape 2 validée par les RH → l'employé (pas l'auteur de l'action).
        awa.post(f"/arrets/{lid}/decision", data={"decision": "valider"})
        assert results[-1]["to"] == ["rh@exemple.com"]
        rh.post(f"/arrets/{lid}/decision", data={"decision": "valider"})
        assert results[-1]["to"] == ["moussa@exemple.com"] and "validé" in delivered[-1]["Subject"]

        # Production : le mail part vers la vraie adresse ; refus avec motif.
        with SessionLocal() as db:
            mails.get_settings(db).mode = "production"
            db.commit()
        moussa.post("/arrets", data={"du": "2026-10-12", "au": "2026-10-12"}, files={"justificatif": ("a.pdf", pdf)})
        assert delivered[-1]["To"] == "awa@exemple.com" and not delivered[-1]["Subject"].startswith("[TEST]")
        with SessionLocal() as db:
            lid2 = db.query(SickLeave).filter(SickLeave.id != lid).one().id
        awa.post(f"/arrets/{lid2}/decision", data={"decision": "refuser", "commentaire": "Justificatif illisible"})
        msg = delivered[-1]
        assert msg["To"] == "moussa@exemple.com" and msg["Subject"] == "Arrêt maladie refusé — Ndiaye Moussa"
        assert "Justificatif illisible" in msg.get_body(("html",)).get_content()
        for c in (moussa, awa, rh):
            c.close()
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_arrets_maladie"))
        with SessionLocal() as db:
            db.query(SickLeaveAction).delete()
            db.query(SickLeave).delete()
            db.query(User).filter(User.username.like("n\\_%", escape="\\")).delete(synchronize_session=False)
            s = db.query(MailSettings).first()
            s.smtp_host, s.from_email, s.mode, s.test_recipients = "", "", "test", ""
            arrets.set_notify(db, False, "test")
            db.commit()


def test_missions_from_authorisation_table(configured, pg, logged_client):
    """Autorisations de mission (ex. feuille Smartsheet synchronisée) : « En mission », 8h validées, même si la
    personne a badgé ; seulement les missions approuvées et les jours ouvrés."""
    import dataclasses

    with pg.begin() as c:
        c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.missions_ss"))
        c.execute(text(f"""CREATE TABLE {SCHEMA}.missions_ss (row_id bigint, matricule text, date_de_depart date,
                          date_de_retour date, statut text)"""))
        c.execute(text(f"""INSERT INTO {SCHEMA}.missions_ss VALUES
            (1, 'E002', '2026-09-21', '2026-09-22', 'Approuvée'),   -- a badgé lundi : en mission quand même
            (2, 'E005', '2026-09-21', NULL, ' APPROUVE '),          -- sans date de retour : un jour ; casse ignorée
            (3, 'E003', '2026-09-21', '2026-09-21', 'En attente'),  -- non approuvée
            (4, 'E004', '2026-09-26', '2026-09-26', 'Validée')      -- samedi : reste non ouvré"""))
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    m = dataclasses.replace(m, objects_schema="pt_test_mission", mission_table="missions_ss", mission_emp_col="matricule",
                            mission_start_col="date_de_depart", mission_end_col="date_de_retour",
                            mission_state_col="statut")
    assert m.mission_ref == "matricule" and "Approuvée" in m.mission_state_values
    with pg.begin() as c:
        c.execute(text("DROP SCHEMA IF EXISTS pt_test_mission CASCADE"))
    pointage.install(pg, m, "test")
    try:
        with pg.connect() as c:
            r = {(x.matricule, x.jour): x for x in c.execute(text(
                "SELECT * FROM pt_test_mission.f_pointage_journalier(:du, :au)"), {"du": MON, "au": SAT}).mappings()}
        moussa = r[("E002", MON)]
        assert (moussa.statut, moussa.statut_libelle) == ("MISSION", "En mission")
        assert hm(moussa.duree_validee) == "8h00" and hm(moussa.duree_effective) == "7h05"  # effective : ses pointages
        assert r[("E002", TUE)].statut == "MISSION" and r[("E002", WED)].statut != "MISSION"
        khady = r[("E005", MON)]
        assert khady.statut == "MISSION" and hm(khady.duree_validee) == "8h00" and hm(khady.duree_effective) == "8h00"
        assert r[("E005", TUE)].statut == "ABSENT"
        assert r[("E003", MON)].statut == "A_L_HEURE" and r[("E004", SAT)].statut == "NON_OUVRE"
        assert r[("E002", MON)].retard_min is None

        # Administration : section Missions avec colonnes proposées.
        admin = logged_client.get(f"/admin/pointage?conn_id={configured}&schema={SCHEMA}&punch_table=punchlog"
                                  f"&emp_table=Employes&mission_table=missions_ss").text
        section = admin.split("Table des autorisations de mission")[1]
        assert '<option value="date_de_depart" selected' in section and '<option value="date_de_retour" selected' in section
        assert '<option value="statut" selected' in section

        # Feuille Smartsheet sans matricule : lien par e-mail (colonne « contact »), plusieurs dates d'aller et de
        # retour par ligne → du premier aller au dernier retour.
        with pg.begin() as c:
            c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Email" text'))
            c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "Email" = lower("Prenom" || \'.\' || "Nom") || \'@exemple.com\''))
            c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.missions_email"))
            c.execute(text(f"""CREATE TABLE {SCHEMA}.missions_email (row_id bigint, demandeur text, demandeur_nom text,
                date_aller_1 date, date_retour_1 date, date_aller_2 date, date_retour_2 date, statut text)"""))
            c.execute(text(f"""INSERT INTO {SCHEMA}.missions_email VALUES
                (1, 'fatou.sow@exemple.com', 'Fatou Sow', '2026-09-22', '2026-09-22', '2026-09-24', '2026-09-25', 'Approuvée'),
                (2, 'ibou.fall@exemple.com', 'Ibou Fall', NULL, NULL, '2026-09-23', NULL, 'Approuvée'),
                (3, 'inconnu@exemple.com', 'Inconnu', '2026-09-22', '2026-09-22', NULL, NULL, 'Approuvée'),
                (4, 'aminata.diop@exemple.com', 'Aminata Diop', '2026-09-21', '2026-09-21', NULL, NULL, 'Refusée')"""))
        m2 = dataclasses.replace(m, email_col="Email", email_in="person", mission_table="missions_email",
                                 mission_emp_col="demandeur", mission_ref="email",
                                 mission_start_col="date_aller_1, date_aller_2",
                                 mission_end_col="date_retour_1, date_retour_2", mission_state_col="statut")
        pointage.install(pg, m2, "test")
        with pg.connect() as c:
            r = {(x.matricule, x.jour): x.statut for x in c.execute(text(
                "SELECT * FROM pt_test_mission.f_pointage_journalier(:du, :au)"), {"du": MON, "au": SAT}).mappings()}
        fatou = [r[("E003", d)] for d in (MON, TUE, WED, date(2026, 9, 24), date(2026, 9, 25))]
        assert fatou == ["A_L_HEURE", "MISSION", "MISSION", "MISSION", "MISSION"]  # 22 → 25 (écart inclus)
        assert r[("E004", WED)] == "MISSION" and r[("E004", date(2026, 9, 24))] == "ABSENT"  # seul l'aller 2 renseigné
        assert r[("E007", MON)] == "RETARD"  # mission refusée
        match = pointage.unmatched_requests(pg, m2, "mission")
        assert match == {"total": 3, "sans_employe": 1, "exemples": [("inconnu@exemple.com", 1)]}

        # Administration : dates multiples et lien par e-mail proposés automatiquement.
        admin = logged_client.get(f"/admin/pointage?conn_id={configured}&schema={SCHEMA}&punch_table=punchlog"
                                  f"&emp_table=Employes&email_col=Email&mission_table=missions_email"
                                  f"&mission_ref=matricule").text
        section = admin.split("Table des autorisations de mission")[1].split("</fieldset>")[0]
        assert '<option value="date_aller_1" selected' in section and '<option value="date_aller_2" selected' in section
        assert '<option value="date_retour_2" selected' in section and '<option value="date_retour_1" selected' in section
        assert '<option value="demandeur" selected' in section and '<option value="email" selected' in section
    finally:
        with pg.begin() as c:
            c.execute(text("DROP SCHEMA IF EXISTS pt_test_mission CASCADE"))
            c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.missions_ss"))
            c.execute(text(f"DROP TABLE IF EXISTS {SCHEMA}.missions_email"))
            c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" DROP COLUMN IF EXISTS "Email"'))


def test_manager_digests(configured, pg, logged_client, monkeypatch):
    """Résumés par mail aux responsables : équipe N-1 (sans le responsable), constats, planification, mode test,
    un seul envoi par période, journal."""
    from app import digests, mails
    from app.models import DigestLog, DigestSubscriber, MailSettings

    # Périodes et heures d'envoi.
    assert digests.period("quotidien", date(2026, 9, 22)) == (MON, MON)
    assert digests.period("hebdomadaire", date(2026, 9, 29)) == (MON, date(2026, 9, 27))
    with SessionLocal() as db:
        s = digests.get_settings(db)
        s.daily_enabled, s.daily_time, s.weekly_enabled, s.weekly_day, s.weekly_time = True, "07:30", True, 1, "08:00"
        db.commit()
        assert digests.due_kinds(s, datetime(2026, 9, 28, 7, 0)) == []
        assert digests.due_kinds(s, datetime(2026, 9, 28, 7, 45)) == ["quotidien"]
        assert digests.due_kinds(s, datetime(2026, 9, 29, 6, 0)) == ["hebdomadaire"]  # lundi manqué : rattrapé mardi

    # Contenu : équipe directe de Diallo (Ndiaye, Sow ; Gueye inactif), sans Diallo elle-même.
    engine = create_engine(POSTGRES_URL)
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    data = digests.collect(pg, m, "1", "quotidien", MON, MON)
    assert [r["matricule"] for r in data["rows"]] == ["E002", "E003"]  # retard d'abord, puis par nom
    notes = " ".join(t for _, t in digests.insights(data))
    assert "1 retard(s) : Ndiaye Moussa (40 min)" in notes and "Taux de présence de 100 %" in notes
    subject, body, _ = digests.render(data, "Diallo Awa", "ACME", "http://pointage:8000")
    assert subject == "Pointages de votre équipe — lundi 21 septembre 2026"
    assert "Point du jour — équipe de Diallo Awa" in body and "http://pointage:8000/suivi?du=2026-09-21" in body
    assert "<strong>Diallo Awa</strong>" not in body and "<strong>Ndiaye Moussa</strong>" in body
    week = digests.collect(pg, m, "1", "hebdomadaire", MON, date(2026, 9, 27))
    subject, body, text_version = digests.render(week, "Diallo Awa")
    assert subject.startswith("Bilan hebdomadaire des pointages de votre équipe — semaine du 21 septembre au 27")
    assert "Détail par collaborateur" in body
    # Moyennes par jour de présence au bureau : durée standard et durée effective de chacun.
    assert "Moy. standard / jour" in body and "Moy. effective / jour" in body
    ndiaye = next(e for e in week["employes"] if e["matricule"] == "E002")
    days = [r for r in rows(pg, MON, date(2026, 9, 27)).values()
            if r.matricule == "E002" and r.statut in ("A_L_HEURE", "RETARD")]
    assert days and abs(float(ndiaye["validee_moy_min"]) - sum(float(r.duree_validee_min) for r in days) / len(days)) < 0.01
    assert abs(float(ndiaye["effective_moy_min"]) - sum(float(r.duree_effective_min) for r in days) / len(days)) < 0.01
    expected = (f"- Ndiaye Moussa : {pointage.hhmm(timedelta(minutes=float(ndiaye['validee_moy_min'])))} / "
                f"{pointage.hhmm(timedelta(minutes=float(ndiaye['effective_moy_min'])))}")
    assert expected in text_version
    assert pointage.hhmm(timedelta(minutes=float(ndiaye["effective_moy_min"]))) in body
    engine.dispose()

    # Administration : responsables proposés, abonnements, aperçu.
    page = logged_client.get("/admin/resumes").text
    assert "Diallo Awa" in page and "Sow Fatou" in page and 'name="freq_1"' in page
    r = logged_client.post("/admin/resumes/abonnes", data={"freq_1": "les_deux", "email_1": "awa@exemple.com",
                                                          "freq_3": "hebdomadaire", "email_3": "pas-une-adresse"},
                           follow_redirects=True)
    assert "2 responsable(s) abonné(s)" in r.text and "Adresse(s) invalide(s) ignorée(s)" in r.text
    preview = logged_client.get("/admin/resumes/apercu?manager=1&kind=hebdomadaire").text
    assert "Point de la semaine — équipe de Diallo Awa" in preview

    # Envoi planifié : SMTP simulé, mode test → adresses de test ; un seul envoi par période.
    delivered = []

    class FakeSMTP:
        def __init__(self, host, port, timeout=None): pass
        def ehlo(self): pass
        def starttls(self, context=None): pass
        def login(self, user, pwd): pass
        def send_message(self, msg): delivered.append(msg)
        def quit(self): pass

    monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)
    with SessionLocal() as db:
        ms = mails.get_settings(db)
        ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "smtp.test", "pointage@exemple.com", "test", "qa@exemple.com"
        db.commit()
    try:
        result = digests.run_due(datetime(2026, 9, 28, 8, 30))  # lundi : veille = dimanche, semaine du 21 au 27
        assert result == {"status": "done", "sent": 1, "skipped": 2, "failed": 0}, result
        # Diallo : hebdomadaire envoyé, quotidien du dimanche ignoré (personne attendu) ; Sow : aucune adresse
        # valide → ignoré, comme il le serait en production.
        subjects = [msg["Subject"] for msg in delivered]
        assert len(subjects) == 1 and subjects[0].startswith("[TEST] Bilan hebdomadaire")
        assert all(msg["To"] == "qa@exemple.com" for msg in delivered)
        assert "MODE TEST" in delivered[0].get_body(("html",)).get_content()
        assert digests.run_due(datetime(2026, 9, 28, 9, 0)) == {"status": "idle"}  # déjà traité
        page = logged_client.get("/admin/resumes").text
        assert "Journal des envois" in page and "Envoyé" in page and "personne n&#39;était attendu" in page
        assert "aucune adresse e-mail" in page

        # Essai depuis l'administration, vers une adresse choisie.
        r = logged_client.post("/admin/resumes/essai", data={"manager": "1", "kind": "quotidien", "to": "moi@exemple.com"},
                               follow_redirects=True)
        assert "Résumé quotidien d&#39;essai envoyé à moi@exemple.com" in r.text
        assert delivered[-1]["To"] == "moi@exemple.com" and delivered[-1]["Subject"].startswith("[TEST] Pointages")
    finally:
        with SessionLocal() as db:
            db.query(DigestLog).delete()
            db.query(DigestSubscriber).delete()
            s = digests.get_settings(db)
            s.daily_enabled = s.weekly_enabled = False
            ms = db.query(MailSettings).first()
            ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "", "", "test", ""
            db.commit()


def test_planning_import_all_matching_sheets(configured, pg, logged_client):
    """Classeur à plusieurs feuilles nommées librement : toutes celles au format du planning sont importées,
    les autres ignorées ; toutes les lignes datées sont prises (sans filtre sur la période de paie) ; conflit signalé."""
    from openpyxl import Workbook

    def sheet(wb, title, pay, days):
        ws = wb.create_sheet(title)
        ws.append([None, "PAIE", datetime(*pay[0]), f" {pay[1][2]:02d}/{pay[1][1]:02d}/{pay[1][0]}"])
        ws.append([None, "Code", "Heure Début", "Heure Fin"])
        ws.append([None, "P1", "06:00", "14:00"])
        ws.append([])
        ws.append(["MATRICULE", "E005"])
        ws.append(["NOM", "BA KHADY"])
        ws.append(["DATE"])
        for day, code in days:
            ws.append([datetime(2026, *day), code])

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("COLLABORATEURS").append(["user_id", "name"])
    sheet(wb, "Pro Caristes Sept 2026", ((2026, 9, 1), (2026, 9, 15)), [((9, 14), "P1"), ((9, 15), "P4")])
    # Seconde feuille : 15/09 en conflit (P1 remplace P4) ; « 16/08 » hors de la période de paie : importé quand même.
    sheet(wb, " Pro Cariste  Sept-Oct 2026", ((2026, 9, 15), (2026, 10, 14)),
          [((9, 15), "P1"), ((9, 16), "P11"), ((8, 16), "P1")])
    buf = io.BytesIO()
    wb.save(buf)
    try:
        r = logged_client.post("/admin/planning/import", files={"fichier": ("PLANNING CARISTES.xlsx", buf.getvalue())},
                               follow_redirects=True)
        text_page = r.text.replace("&#39;", "'")
        assert "importé du 16/08/2026 au 16/09/2026 : 1 employé(s), 4 jour(s)" in text_page, text_page[:3000]
        assert "2 feuilles importées" in text_page and "« Pro Caristes Sept 2026 »" in text_page
        assert "Feuilles ignorées" in text_page and "« COLLABORATEURS »" in text_page
        assert "hors de la période de paie" not in text_page
        assert "1 jour(s) planifié(s) dans plusieurs feuilles" in text_page and "P4 (« Pro Caristes Sept 2026 ») remplacé par P1" in text_page
        with pg.connect() as c:
            got = dict(c.execute(text(f"SELECT jour, code FROM {SCHEMA}.pointage_planning WHERE matricule = 'E005'")).all())
        assert got == {date(2026, 9, 14): "P1", date(2026, 9, 15): "P1", date(2026, 9, 16): "P11", date(2026, 8, 16): "P1"}
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_planning"))


def test_import_size_limit(configured, logged_client, monkeypatch):
    """Fichiers importés : 50 Mo maximum (planning et agents terrain), limite affichée dans les formulaires."""
    from app.routers import admin

    assert admin.MAX_IMPORT_BYTES == 50 * 1024 * 1024
    assert ".xlsx · 50 Mo maximum" in logged_client.get("/admin/planning").text
    assert "50 Mo maximum" in logged_client.get("/admin/terrain").text
    monkeypatch.setattr(admin, "MAX_IMPORT_BYTES", 10)
    for url in ("/admin/planning/import", "/admin/terrain/import"):
        r = logged_client.post(url, files={"fichier": ("gros.xlsx", b"x" * 11)}, follow_redirects=True)
        assert "Fichier trop volumineux (50 Mo maximum)" in r.text


def test_planning_import_several_workbooks(configured, pg, logged_client):
    """Plusieurs classeurs importés en une fois : feuilles réunies, conflit entre fichiers signalé, fichier
    illisible et fichier sans planning ignorés sans bloquer les autres."""
    from openpyxl import Workbook

    def workbook(sheets: dict) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)
        for title, days in sheets.items():
            ws = wb.create_sheet(title)
            ws.append(["MATRICULE", "E005"])
            ws.append(["DATE"])
            for day, code in days:
                ws.append([datetime(2026, 9, day), code])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    files = [
        ("fichier", ("Caristes 2026.xlsx", workbook({"Sept": [(1, "P1"), (2, "P2")]}))),
        ("fichier", ("Caristes bis.xlsx", workbook({"Sept suite": [(2, "P3"), (3, "P4")]}))),
        ("fichier", ("liste.xlsx", _empty_workbook())),
        ("fichier", ("abime.xlsx", b"pas un classeur")),
    ]
    try:
        r = logged_client.post("/admin/planning/import", files=files, follow_redirects=True)
        page = r.text.replace("&#39;", "'")
        assert "Planning « 4 classeurs » importé du 01/09/2026 au 03/09/2026 : 1 employé(s), 3 jour(s)" in page, page[:3000]
        assert "« Caristes 2026.xlsx › Sept »" in page and "« Caristes bis.xlsx › Sept suite »" in page
        assert "P2 (« Caristes 2026.xlsx › Sept ») remplacé par P3 (« Caristes bis.xlsx › Sept suite »)" in page
        assert "« liste.xlsx › Liste »" in page and "« abime.xlsx » : fichier Excel illisible" in page
        with pg.connect() as c:
            got = dict(c.execute(text(f"SELECT jour, code FROM {SCHEMA}.pointage_planning WHERE matricule = 'E005'")).all())
        assert got == {date(2026, 9, 1): "P1", date(2026, 9, 2): "P3", date(2026, 9, 3): "P4"}
        assert 'name="fichier" accept=".xlsx,.xlsm" multiple' in logged_client.get("/admin/planning").text
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_planning"))


def _empty_workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Liste"
    wb.active.append(["user_id", "name"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_planning_import_with_progress(configured, pg, logged_client):
    """Import suivi en direct : démarrage, avancement lu par la page (étapes et pourcentage), fin avec les messages ;
    étapes de progression signalées par la lecture et l'enregistrement."""
    import time as _t

    from fastapi.testclient import TestClient

    from app.main import app

    data = _planning_xlsx({date(2026, 9, 24): ("1806", "P1"), date(2026, 9, 25): ("P4", "P1")})
    steps = []
    plan = pointage.read_planning_files([("a.xlsx", data), ("b.xlsx", data)], lambda p, label: steps.append((p, label)))
    assert steps[0] == (0, "Lecture du classeur « a.xlsx » (1/2)") and any("feuille « PLANNING »" in s for _, s in steps)
    assert [p for p, _ in steps] == sorted(p for p, _ in steps) and max(p for p, _ in steps) < 60
    try:
        r = logged_client.post("/admin/planning/import/demarrer", files={"fichier": ("PLANNING ATF.xlsx", data)})
        job = r.json()["job"]
        for _ in range(100):
            status = logged_client.get(f"/admin/planning/import/{job}").json()
            if status["state"] == "done":
                break
            _t.sleep(0.05)
        assert status["state"] == "done" and status["percent"] == 100 and status["label"] == "Import terminé"
        # Un autre utilisateur ne voit pas cet import.
        with TestClient(app) as other:
            other.post("/login", data={"username": "admin", "password": "secret"})
            other.cookies.clear()
            assert other.get(f"/admin/planning/import/{job}", follow_redirects=False).status_code in (303, 404)
        page = logged_client.get(f"/admin/planning/import/{job}/fin", follow_redirects=True).text
        assert "Planning « PLANNING ATF.xlsx » importé du 24/09/2026 au 25/09/2026" in page
        assert logged_client.get(f"/admin/planning/import/{job}").status_code == 404  # terminé : oublié
        assert 'id="import_progress"' in page and "Envoi des fichiers" in page
        with pg.connect() as c:
            assert c.execute(text(f"SELECT count(*) FROM {SCHEMA}.pointage_planning")).scalar() == 2
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_planning"))


def test_direction_filter_and_reports(configured, pg, logged_client):
    """Direction (colonne « Direction » de Personnel) : proposée automatiquement, filtre multiple du suivi et des
    rapports, rapport « Par direction », exports."""
    import io

    from openpyxl import load_workbook

    with pg.begin() as c:
        c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Direction" text'))
        c.execute(text(f"""UPDATE {SCHEMA}."Personnel" SET "Direction" = CASE WHEN "IDPersonnel" IN (101, 102, 106)
                          THEN 'Direction Industrielle' WHEN "IDPersonnel" IN (103, 104) THEN 'Direction RH' END"""))
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    try:
        # Suggestion automatique de la colonne dans la configuration.
        admin = logged_client.get("/admin/pointage").text
        section = admin.split('id="direction"')[1].split("</fieldset>")[0]
        assert '<option value="Direction" selected' in section
        data = {**pointage.Mapping.from_json(original).__dict__, "dir_col": "Direction", "dir_in": "person",
                "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=data, follow_redirects=True).text
        r = rows(pg, MON)
        assert r[("E001", MON)].direction == "Direction Industrielle" and r[("E005", MON)].direction is None
        with SessionLocal() as db:
            m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
        assert pointage.directions(pg, m) == ["Direction Industrielle", "Direction RH"]

        page = logged_client.get(f"/suivi?date={MON.isoformat()}&direction=Direction+RH").text
        assert 'name="direction"' in page and "Direction : Direction RH" in page
        assert "E003</td>" in page and "E004</td>" in page and "E001</td>" not in page and "E005</td>" not in page
        both = logged_client.get(f"/suivi?date={MON.isoformat()}&direction=Direction+RH"
                                 f"&direction=Direction+Industrielle&pop=liste").text
        assert "E001</td>" in both and "E003</td>" in both and "E005</td>" not in both
        xlsx = load_workbook(io.BytesIO(logged_client.get(
            f"/suivi/export.xlsx?date={MON.isoformat()}&direction=Direction+RH").content))
        ws = xlsx.active
        assert ws.cell(1, 18).value == "Direction"
        assert {row[17] for row in ws.iter_rows(min_row=2, values_only=True) if row[0] not in (None, "Moyenne")} == {"Direction RH"}
        assert dict(xlsx["Filtres"].iter_rows(values_only=True))["Direction"] == "Direction RH"

        url = f"/rapports?p=perso&du={MON.isoformat()}&au={(MON + timedelta(days=4)).isoformat()}"
        rep = logged_client.get(url).text
        assert "Par direction" in rep and "(sans direction)" in rep and "Direction Industrielle" in rep
        assert "Ba Khady" not in logged_client.get(url + "&direction=Direction+RH").text
        wb = load_workbook(io.BytesIO(logged_client.get("/rapports/export.xlsx?" + url.split("?", 1)[1]).content))
        assert "Directions" in wb.sheetnames and wb["Employés"].cell(1, 22).value == "Direction"
        assert "<code>Direction Industrielle</code>" in logged_client.get("/admin/pointage").text
    finally:
        restore = {**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=restore, follow_redirects=True).text
        with pg.begin() as c:
            c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" DROP COLUMN IF EXISTS "Direction"'))


def test_site_filter_and_reports(configured, pg, logged_client):
    """Site (Usine, Plateau) : colonne de Personnel proposée, filtre du suivi et des rapports, « Par site », exports."""
    import io

    from openpyxl import load_workbook

    with pg.begin() as c:
        c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Site" text'))
        c.execute(text(f"""UPDATE {SCHEMA}."Personnel" SET "Site" = CASE WHEN "IDPersonnel" IN (101, 102, 106, 107)
                          THEN 'Usine' ELSE 'Plateau' END"""))
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    try:
        section = logged_client.get("/admin/pointage").text.split('id="site"')[1].split("</fieldset>")[0]
        assert '<option value="Site" selected' in section and "Usine, Plateau" in section
        data = {**pointage.Mapping.from_json(original).__dict__, "site_col": "Site", "site_in": "person", "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=data, follow_redirects=True).text
        assert rows(pg, MON)[("E003", MON)].site == "Plateau"
        page = logged_client.get(f"/suivi?date={MON.isoformat()}&site=Usine").text
        assert 'name="site"' in page and "Site : Usine" in page
        assert "E001</td>" in page and "E007</td>" in page and "E003</td>" not in page
        xlsx = load_workbook(io.BytesIO(logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}&site=Plateau").content))
        assert xlsx.active.cell(1, 19).value == "Site" and dict(xlsx["Filtres"].iter_rows(values_only=True))["Site"] == "Plateau"
        url = f"/rapports?p=perso&du={MON.isoformat()}&au={(MON + timedelta(days=4)).isoformat()}"
        rep = logged_client.get(url).text
        assert "Par site" in rep and ">Usine</a>" in rep and ">Plateau</a>" in rep
        plateau = logged_client.get(url + "&site=Plateau").text
        assert "E003</td>" in plateau and "E001</td>" not in plateau
        wb = load_workbook(io.BytesIO(logged_client.get("/rapports/export.xlsx?" + url.split("?", 1)[1]).content))
        assert "Sites" in wb.sheetnames and wb["Employés"].cell(1, 23).value == "Site"
    finally:
        restore = {**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=restore, follow_redirects=True).text
        with pg.begin() as c:
            c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" DROP COLUMN IF EXISTS "Site"'))


def test_employee_accounts_see_self_and_team(configured, logged_client):
    """Comptes des employés : création en masse (identifiant = matricule, mots de passe provisoires dans un fichier
    Excel), changement obligatoire à la première connexion, chacun voit ses données et celles de son équipe."""
    import io

    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from app.main import app
    from app.models import User

    # Comptes déjà rattachés à ces matricules par d'autres tests : détachés le temps du test, puis restaurés.
    with SessionLocal() as db:
        linked = {u.id: u.emp_matricule for u in db.query(User).filter(User.emp_matricule.like("E00%"))}
        for u in db.query(User).filter(User.id.in_(linked)):
            u.emp_matricule = None
        db.commit()
    page = logged_client.get("/admin/users").text
    assert "Comptes des employés" in page and "sans compte" in page
    assert "8 caractères" in logged_client.post("/admin/users/comptes-employes", data={"mode": "commun", "mot_de_passe": "court"},
                                                follow_redirects=True).text
    try:
        r = logged_client.post("/admin/users/comptes-employes", data={"mode": "aleatoire"})
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        ws = load_workbook(io.BytesIO(r.content)).active
        creds = {row[0]: row[4] for row in ws.iter_rows(min_row=2, values_only=True)}
        assert set(creds) == {"E001", "E002", "E003", "E004", "E005", "E007"}  # E006 inactif : pas de compte
        assert all(len(p) == 10 for p in creds.values()) and len(set(creds.values())) == 6
        with SessionLocal() as db:
            awa = db.query(User).filter(User.username == "E001").one()
            assert (awa.role, awa.scope, awa.emp_matricule, awa.must_change_password, awa.full_name) == (
                "lecteur", "equipe", "E001", True, "Diallo Awa")
        assert "Tous les employés actifs ont déjà un compte" in logged_client.post(
            "/admin/users/comptes-employes", data={"mode": "aleatoire"}, follow_redirects=True).text

        def login(mat):
            c = TestClient(app)
            r = c.post("/login", data={"username": mat, "password": creds[mat]}, follow_redirects=True)
            assert "mot de passe provisoire" in r.text  # changement obligatoire
            assert c.get("/suivi", follow_redirects=False).headers["location"] == "/compte"
            r = c.post("/compte", data={"current": creds[mat], "new": "MonNouveau1", "confirm": "MonNouveau1"},
                       follow_redirects=True)
            assert "Mot de passe modifié" in r.text
            return c

        # Diallo (E001) encadre Ndiaye, Sow (et, sous Sow, Fall et Ba) : elle voit toute son équipe, pas Diop.
        awa = login("E001")
        team = awa.get(f"/suivi?date={MON.isoformat()}").text
        for mat in ("E001", "E002", "E003", "E004", "E005"):
            assert f"{mat}</td>" in team, mat
        assert "E007</td>" not in team and "Hors liste" not in team.split("<tbody>")[1]
        # Ndiaye (E002) n'encadre personne : il ne voit que lui-même, y compris dans les rapports et l'export.
        moussa = login("E002")
        own = moussa.get(f"/suivi?date={MON.isoformat()}").text
        assert "E002</td>" in own and "E001</td>" not in own and "E003</td>" not in own
        rep = moussa.get(f"/rapports?p=perso&du={MON.isoformat()}&au={MON.isoformat()}").text
        assert "Ndiaye Moussa" in rep and "Diallo Awa" not in rep.split("Par employé")[-1]
        assert moussa.get("/admin/users", follow_redirects=False).status_code in (302, 303)  # pas d'administration
        for c in (awa, moussa):
            c.close()
        # Réinitialisation par un administrateur : de nouveau provisoire.
        with SessionLocal() as db:
            uid = db.query(User).filter(User.username == "E002").one().id
        logged_client.post(f"/admin/users/{uid}/password", data={"password": "Provisoire9"})
        with SessionLocal() as db:
            assert db.get(User, uid).must_change_password is True
    finally:
        with SessionLocal() as db:
            db.query(User).filter(User.username.in_(["E001", "E002", "E003", "E004", "E005", "E007"])).delete(
                synchronize_session=False)
            for u in db.query(User).filter(User.id.in_(linked)):
                u.emp_matricule = linked[u.id]
            db.commit()


def test_staff_login_with_email_and_standard_password(configured, pg, logged_client):
    """Tout le personnel : connexion avec l'e-mail de la fiche et le mot de passe standard, compte ouvert à la première
    connexion, périmètre « son équipe » ; mot de passe personnel ensuite ; employé inactif refusé ; accès fermable."""
    from fastapi.testclient import TestClient

    from app.main import app
    from app.models import StaffAccessSettings, User

    with pg.begin() as c:
        c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Email" text'))
        c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "Email" = lower("Prenom" || \'.\' || "Nom") || \'@exemple.com\''))
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    data = {**pointage.Mapping.from_json(original).__dict__, "email_col": "Email", "email_in": "person", "conn_id": conn_id}

    def client():
        return TestClient(app)

    try:
        assert "installées" in logged_client.post("/admin/pointage", data=data, follow_redirects=True).text
        page = logged_client.get("/admin/users").text
        assert "Accès de tout le personnel" in page and "avec une adresse e-mail" in page
        r = logged_client.post("/admin/users/acces-personnel", data={"enabled": "true"}, follow_redirects=True)
        assert "Définissez le mot de passe standard" in r.text
        r = logged_client.post("/admin/users/acces-personnel", data={"enabled": "true", "mot_de_passe": "Standard2026"},
                               follow_redirects=True)
        assert "Accès du personnel ouvert" in r.text

        # Première connexion de Diallo (E001) : e-mail en majuscules accepté, compte créé, voit son équipe.
        awa = client()
        r = awa.post("/login", data={"username": "Awa.Diallo@Exemple.com", "password": "Standard2026"}, follow_redirects=True)
        assert "Suivi journalier" in r.text
        with SessionLocal() as db:
            u = db.query(User).filter(User.username == "awa.diallo@exemple.com").one()
            assert (u.role, u.scope, u.emp_matricule, u.auto_account, u.must_change_password) == ("lecteur", "equipe", "E001", True, False)
        team = awa.get(f"/suivi?date={MON.isoformat()}").text
        assert "E002</td>" in team and "E004</td>" in team and "E007</td>" not in team
        # Mauvais mot de passe, adresse inconnue, employé inactif (Gueye, E006) : refusés.
        for user, pwd in (("awa.diallo@exemple.com", "mauvais!"), ("inconnu@exemple.com", "Standard2026"),
                          ("ousmane.gueye@exemple.com", "Standard2026")):
            assert "Identifiants incorrects" in client().post("/login", data={"username": user, "password": pwd},
                                                              follow_redirects=True).text, user
        # Mot de passe personnel : le mot de passe standard ne vaut plus pour elle.
        r = awa.post("/compte", data={"current": "Standard2026", "new": "AwaPerso99", "confirm": "AwaPerso99"},
                     follow_redirects=True)
        assert "Mot de passe modifié" in r.text
        assert "Identifiants incorrects" in client().post("/login", data={"username": "awa.diallo@exemple.com",
                                                                          "password": "Standard2026"}, follow_redirects=True).text
        assert "Suivi journalier" in client().post("/login", data={"username": "awa.diallo@exemple.com",
                                                                   "password": "AwaPerso99"}, follow_redirects=True).text
        # Nouveau mot de passe standard (et mot de passe personnel exigé) : vaut pour ceux qui n'en ont pas choisi.
        moussa = client()
        assert "Suivi journalier" in moussa.post("/login", data={"username": "moussa.ndiaye@exemple.com",
                                                                 "password": "Standard2026"}, follow_redirects=True).text
        logged_client.post("/admin/users/acces-personnel", data={"enabled": "true", "mot_de_passe": "Nouveau2027",
                                                                 "force_change": "true"})
        assert "Identifiants incorrects" in client().post("/login", data={"username": "moussa.ndiaye@exemple.com",
                                                                          "password": "Standard2026"}, follow_redirects=True).text
        r = client().post("/login", data={"username": "moussa.ndiaye@exemple.com", "password": "Nouveau2027"}, follow_redirects=True)
        assert "mot de passe provisoire" in r.text  # mot de passe personnel exigé
        # Compte créé auparavant par matricule (même e-mail), sans mot de passe personnel : le standard vaut aussi.
        with SessionLocal() as db:
            db.add(User(username="E003", email="fatou.sow@exemple.com", emp_matricule="E003", role="lecteur",
                        scope="equipe", password_hash="x", active=True))
            db.commit()
        fatou = client()
        fatou.post("/login", data={"username": "fatou.sow@exemple.com", "password": "Nouveau2027"})
        assert fatou.get("/compte").status_code == 200 and "E003" in fatou.get("/compte").text
        with SessionLocal() as db:
            db.query(User).filter(User.username == "E003").delete()
            db.commit()
        # Accès fermé : plus de connexion par le mot de passe standard ; l'administrateur et Awa (personnel) passent.
        logged_client.post("/admin/users/acces-personnel", data={"mot_de_passe": ""})
        assert "Identifiants incorrects" in client().post("/login", data={"username": "fatou.sow@exemple.com",
                                                                          "password": "Nouveau2027"}, follow_redirects=True).text
        assert "Suivi journalier" in client().post("/login", data={"username": "awa.diallo@exemple.com",
                                                                   "password": "AwaPerso99"}, follow_redirects=True).text
        assert client().post("/login", data={"username": "admin", "password": "secret"}, follow_redirects=False).status_code == 303
    finally:
        with SessionLocal() as db:
            db.query(User).filter(User.auto_account.is_(True)).delete(synchronize_session=False)
            db.query(StaffAccessSettings).delete()
            db.commit()
        restore = {**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id}
        assert "installées" in logged_client.post("/admin/pointage", data=restore, follow_redirects=True).text
        with pg.begin() as c:
            c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" DROP COLUMN IF EXISTS "Email"'))


def test_punch_reminders(configured, pg, logged_client, monkeypatch):
    """Rappels de pointage : arrivée non pointée (heure de bureau ou début de poste + délai), départ non pointé, oubli
    de la veille (copie au responsable), filtres, garde-fou de synchronisation, un seul rappel par journée."""
    from app import mails, rappels
    from app.models import MailSettings, ReminderLog

    with pg.begin() as c:
        c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Email" text'))
        c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "Email" = lower("Prenom" || \'.\' || "Nom") || \'@exemple.com\''))
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    data = {**pointage.Mapping.from_json(original).__dict__, "email_col": "Email", "email_in": "person", "conn_id": conn_id}
    delivered = []

    class FakeSMTP:
        def __init__(self, host, port, timeout=None): pass
        def ehlo(self): pass
        def starttls(self, context=None): pass
        def login(self, user, pwd): pass
        def send_message(self, msg): delivered.append(msg)
        def quit(self): pass

    monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)

    def configure(**values):
        with SessionLocal() as db:
            s = rappels.get_settings(db)
            for k, v in {"arrival_enabled": False, "departure_enabled": False, "forgot_enabled": False, "alert_enabled": False,
                         "predeparture_enabled": False, "alert_cc_manager": False, "texts": "{}", "audience": "tous",
                         "services": "[]",
                         "sites": "[]", "directions": "[]", "excluded": "", "forgot_cc_manager": False, **values}.items():
                setattr(s, k, v)
            db.commit()

    try:
        assert "installées" in logged_client.post("/admin/pointage", data=data, follow_redirects=True).text
        with SessionLocal() as db:
            ms = mails.get_settings(db)
            ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "smtp.test", "pointage@exemple.com", "test", "qa@exemple.com"
            db.commit()

        # Paramétrage depuis l'administration (heures au format libre « 8h20 », délais des postes).
        r = logged_client.post("/admin/rappels/parametres", data={"arrival_time": "25:00", "departure_time": "17:45",
                                                                 "forgot_time": "09:00"}, follow_redirects=True)
        assert "Heure du rappel d&#39;arrivée invalide" in r.text
        r = logged_client.post("/admin/rappels/parametres", data={
            "arrival_enabled": "true", "arrival_time": "8h20", "arrival_delay": "15", "departure_time": "17:45",
            "departure_delay": "30", "forgot_time": "09:00", "max_lag_minutes": "60", "message": "Contact : rh@exemple.com"},
            follow_redirects=True)
        assert "Paramètres des rappels enregistrés" in r.text
        with SessionLocal() as db:
            s = rappels.get_settings(db)
            assert (s.arrival_enabled, s.arrival_time, s.departure_enabled) == (True, "08:20", False)

        # Arrivée : à 8h20, seule Ba Khady (E005, aucun pointage le lundi) est attendue sans avoir badgé.
        configure(arrival_enabled=True, arrival_time="08:20")  # destinataires : tous les employés (abonnés : plus bas)
        assert rappels.run_due(datetime(2026, 9, 21, 8, 10)) == {"status": "idle"}  # heure non atteinte
        assert rappels.run_due(datetime(2026, 9, 21, 8, 30)) == {"status": "done", "sent": 1, "skipped": 0, "failed": 0}
        msg = delivered[-1]
        assert msg["To"] == "qa@exemple.com" and msg["Subject"] == "[TEST] Rappel : votre arrivée n'est pas encore pointée"
        body = msg.get_body(("html",)).get_content()
        assert "Bonjour Khady" in body and "khady.ba@exemple.com" in body and "Contact : rh@exemple.com" in body
        # Présenté comme le mail de confirmation de badge : bandeau, case « Arrivée / Date », pointages de la journée.
        assert "Rappel de pointage</p>" in body and "Non pointée" in body and "Vos pointages de la journée" in body
        assert "Aucun pointage enregistré" in body and "Matricule E005 · RH" in body
        assert rappels.run_due(datetime(2026, 9, 21, 8, 45)) == {"status": "idle"}  # un seul rappel par journée
        assert rappels.run_due(datetime(2026, 9, 21, 11, 30)) == {"status": "idle"}

        # Départ : à 17h45, Fall Ibou (E004) n'a badgé que son arrivée (7h40).
        configure(departure_enabled=True)
        assert rappels.run_due(datetime(2026, 9, 21, 17, 50)) == {"status": "done", "sent": 1, "skipped": 0, "failed": 0}
        assert delivered[-1]["Subject"] == "[TEST] Rappel : pensez à pointer votre départ"
        assert "Votre dernier pointage aujourd'hui date de 07h40" in delivered[-1].get_body(("html",)).get_content()

        # Oubli de la veille, en production : à l'employé, responsable N+1 (Sow Fatou) en copie.
        configure(forgot_enabled=True, forgot_cc_manager=True)
        with SessionLocal() as db:
            db.query(MailSettings).one().mode = "production"
            db.commit()
        assert rappels.run_due(datetime(2026, 9, 22, 9, 5)) == {"status": "done", "sent": 1, "skipped": 0, "failed": 0}
        msg = delivered[-1]
        assert (msg["To"], msg["Cc"]) == ("ibou.fall@exemple.com", "fatou.sow@exemple.com")
        assert msg["Subject"] == "Pointage incomplet le lundi 21 septembre 2026"
        assert "Un seul pointage a été enregistré le lundi 21 septembre 2026, à 07h40" in msg.get_body(("html",)).get_content()
        with SessionLocal() as db:
            db.query(MailSettings).one().mode = "test"
            db.query(ReminderLog).delete()
            db.commit()

        # Alerte du matin (2e niveau) à 9h30 : toujours aucun pointage pour Ba Khady ; responsable (Sow Fatou) en copie.
        configure(alert_enabled=True, alert_time="09:30", alert_cc_manager=True)
        with SessionLocal() as db:
            db.query(MailSettings).one().mode = "production"
            db.commit()
        assert rappels.run_due(datetime(2026, 9, 21, 9, 0)) == {"status": "idle"}
        assert rappels.run_due(datetime(2026, 9, 21, 9, 35)) == {"status": "done", "sent": 1, "skipped": 0, "failed": 0}
        msg = delivered[-1]
        assert (msg["To"], msg["Cc"], msg["Subject"]) == ("khady.ba@exemple.com", "fatou.sow@exemple.com",
                                                         "Alerte : aucun pointage enregistré ce matin")
        with SessionLocal() as db:
            db.query(MailSettings).one().mode = "test"
            db.commit()

        # Avertissement du soir, texte personnalisé depuis l'administration : à 16h45, tous les employés encore
        # « entrés » (nombre impair de pointages) : Diallo Awa (3 pointages) et Fall Ibou (1 seul).
        r = logged_client.post("/admin/rappels/parametres", data={
            "predeparture_enabled": "true", "predeparture_time": "16h45", "predeparture_before": "15", "audience": "tous",
            "subject_avant_depart": "{prenom}, pensez à badger ce soir",
            "body_avant_depart": "Arrivée enregistrée à {arrivee} ({date}).\n\nBonne soirée {prenom} {nom} !"},
            follow_redirects=True)
        assert "Paramètres des rappels enregistrés" in r.text and "personnalisé" in r.text
        before = len(delivered)
        assert rappels.run_due(datetime(2026, 9, 21, 16, 50)) == {"status": "done", "sent": 2, "skipped": 0, "failed": 0}
        sent = {m["Subject"]: m.get_body(("html",)).get_content() for m in delivered[before:]}
        assert set(sent) == {"[TEST] Awa, pensez à badger ce soir", "[TEST] Ibou, pensez à badger ce soir"}
        body = sent["[TEST] Awa, pensez à badger ce soir"]
        assert "Arrivée enregistrée à 07h20 (lundi 21 septembre 2026)." in body and "Bonne soirée Awa Diallo !" in body
        assert "Awa, pensez à badger ce soir</p>" in body and "Bonjour Awa,</h1>" in body
        with SessionLocal() as db:
            db.query(ReminderLog).delete()
            db.commit()

        # Destinataires par défaut : seulement les employés cochés dans Mails de badge → Abonnés, à l'adresse saisie
        # pour eux (comme la confirmation de pointage). Fall Ibou abonné, Ba Khady non.
        from app.models import MailSubscriber

        with SessionLocal() as db:
            db.query(MailSubscriber).delete()
            db.add(MailSubscriber(emp_key="4", matricule="E004", name="Fall Ibou", email="ibou.perso@exemple.com"))
            db.query(MailSettings).one().mode = "production"
            db.commit()
        configure(arrival_enabled=True, departure_enabled=True, audience="abonnes")
        assert rappels.run_due(datetime(2026, 9, 21, 8, 30)) == {"status": "idle"}  # Ba Khady n'est pas abonnée
        assert rappels.run_due(datetime(2026, 9, 21, 17, 50)) == {"status": "done", "sent": 1, "skipped": 0, "failed": 0}
        assert delivered[-1]["To"] == "ibou.perso@exemple.com"
        page = logged_client.get("/admin/rappels").text
        assert "Employés abonnés aux mails de badge" in page and "(1 abonné(s))" in page
        with SessionLocal() as db:
            db.query(MailSubscriber).delete()
            db.query(MailSettings).one().mode = "test"
            db.query(ReminderLog).delete()
            db.commit()
        r = logged_client.post("/admin/rappels/parametres", data={"audience": "abonnes"}, follow_redirects=True)
        assert "Aucun employé n'est abonné aux mails de badge" in r.text

        # Filtres : matricule exclu, service non retenu.
        configure(departure_enabled=True, excluded="E004")
        assert rappels.run_due(datetime(2026, 9, 21, 17, 50)) == {"status": "idle"}
        configure(departure_enabled=True, services='["Production"]')  # Fall Ibou est au service RH
        assert rappels.run_due(datetime(2026, 9, 21, 17, 50)) == {"status": "idle"}

        # Horaires postés : poste de 14h à 22h planifié pour Ba Khady → rappel 15 min après le début du poste.
        with pg.begin() as c:
            c.execute(text(f"INSERT INTO {SCHEMA}.pointage_postes (code, libelle, type, debut, fin) "
                           f"VALUES ('RAPP', 'Après-midi', 'travail', '14:00', '22:00') ON CONFLICT (code) DO NOTHING"))
            c.execute(text(f"INSERT INTO {SCHEMA}.pointage_planning (matricule, jour, code) VALUES ('E005', :d, 'RAPP')"),
                      {"d": MON})
        configure(arrival_enabled=True, arrival_time="08:20", arrival_delay=15)
        with SessionLocal() as db:
            m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
            s = rappels.get_settings(db)
        rows = rappels.day_rows(pg, m, MON, MON)
        due = {(k, r["matricule"]): w for k, r, w in rappels.candidates(rows, s, datetime(2026, 9, 21, 8, 30), ["arrivee"],
                                                                          timed=False)}
        assert due == {("arrivee", "E005"): datetime(2026, 9, 21, 14, 15)}
        assert rappels.run_due(datetime(2026, 9, 21, 8, 30)) == {"status": "idle"}
        assert rappels.run_due(datetime(2026, 9, 21, 14, 20))["sent"] == 1
        assert "(poste RAPP (14h00 – 22h00))" in delivered[-1].get_body(("html",)).get_content()
        assert rappels.run_due(datetime(2026, 9, 21, 17, 20)) == {"status": "idle"}  # pas de doublon

        # Garde-fou : pointages non synchronisés depuis 2 h → aucun rappel.
        with SessionLocal() as db:
            db.query(ReminderLog).delete()
            db.commit()
        monkeypatch.setattr(rappels, "last_sync", lambda db, m: rappels.utcnow() - timedelta(hours=2))
        result = rappels.run_due(datetime(2026, 9, 21, 14, 20))
        assert result["status"] == "stale" and "non synchronisés depuis 120 min" in result["reason"]
        monkeypatch.undo()
        monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)

        # Administration : situation, aperçu, essai, journal ; accessible au rôle RH.
        page = logged_client.get("/admin/rappels").text
        assert "Situation du jour" in page and "Rappels et heures d'envoi" in page and 'value="08:20"' in page
        assert "Pensez à pointer votre départ" in logged_client.get("/admin/rappels/apercu?kind=depart").text
        r = logged_client.post("/admin/rappels/essai", data={"kind": "oubli", "to": "moi@exemple.com"}, follow_redirects=True)
        assert "Rappel d&#39;essai « Oubli de la veille » envoyé à moi@exemple.com" in r.text
        assert delivered[-1]["To"] == "moi@exemple.com" and delivered[-1]["Subject"].startswith("[TEST] Pointage incomplet")
        assert "Rappels de pointage" in logged_client.get("/admin").text
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_planning WHERE code = 'RAPP'"))
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_postes WHERE code = 'RAPP'"))
        with SessionLocal() as db:
            db.query(ReminderLog).delete()
            s = rappels.get_settings(db)
            s.arrival_enabled = s.departure_enabled = s.forgot_enabled = s.alert_enabled = s.predeparture_enabled = False
            s.services, s.excluded, s.message, s.texts = "[]", "", "", "{}"
            ms = db.query(MailSettings).first()
            ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "", "", "test", ""
            db.commit()
        logged_client.post("/admin/pointage", data={**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id})


def test_hide_validated_duration(configured, logged_client):
    """Durée standard du suivi : règle d'affichage (RH et admin / personne) et choix personnel, écran et export."""
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from app import auth
    from app.main import app

    url = f"/suivi?date={MON.isoformat()}"
    page = logged_client.get(url).text
    assert "Durée standard" in page and "Masquer la durée standard" in page
    # Choix personnel : masquée à l'écran et dans l'export, puis réaffichée.
    logged_client.get("/suivi/affichage?validee=0", headers={"referer": f"http://testserver{url}"})
    page = logged_client.get(url).text
    assert "Durée standard" not in page and "Afficher la durée standard" in page and "Durée effective moyenne" in page
    ws = load_workbook(io.BytesIO(logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}").content))["Suivi journalier"]
    headers = [c.value for c in ws[1]]
    assert "Durée standard" not in headers and "Durée standard (min)" not in headers
    assert headers[10:12] == ["Durée effective", "Durée effective (min)"] and headers[12] == "Responsable"
    logged_client.get("/suivi/affichage?validee=1")
    assert "Durée standard" in logged_client.get(url).text

    # Règle « RH et administrateurs » : un lecteur ne la voit plus (ni bouton pour la réafficher), l'admin si.
    r = logged_client.post("/admin/parametres/affichage", data={"validee_display": "rh"}, follow_redirects=True)
    assert "visible seulement pour les rh et les administrateurs" in r.text
    with SessionLocal() as db:
        db.query(User).filter_by(username="lect_vd").delete()
        db.add(User(username="lect_vd", role="lecteur", password_hash=auth.hash_password("motdepasse7")))
        db.commit()
    try:
        client = TestClient(app)
        client.post("/login", data={"username": "lect_vd", "password": "motdepasse7"})
        page = client.get(url).text
        assert "Durée standard" not in page and "Afficher la durée standard" not in page
        ws = load_workbook(io.BytesIO(client.get(f"/suivi/export.xlsx?date={MON.isoformat()}").content))["Suivi journalier"]
        assert "Durée standard" not in [c.value for c in ws[1]]
        assert "Durée standard" in logged_client.get(url).text
        # Rapports : heures standard réservées aux RH et aux administrateurs (écran, Excel, PDF, tri).
        q = f"p=perso&du={MON.isoformat()}&au={TUE.isoformat()}"
        assert "Heures standard" not in client.get("/rapports?" + q + "&tri=heures").text
        assert "Heures standard" in logged_client.get("/rapports?" + q).text
        wb = load_workbook(io.BytesIO(client.get("/rapports/export.xlsx?" + q).content))
        assert not {"Heures standard", "Moyenne heures / jour au bureau"} & {c.value for c in wb["Employés"][1]}
        assert "Heures standard" in [c.value for c in load_workbook(io.BytesIO(
            logged_client.get("/rapports/export.xlsx?" + q).content))["Employés"][1]]
        assert "(STANDARD) Tj" not in _pdf_text(client.get("/rapports/export.pdf?" + q).content)
        # L'ancien réglage « visible pour tous » n'existe plus : retour à « RH et administrateurs ».
        logged_client.post("/admin/parametres/affichage", data={"validee_display": "tous"})
        assert "Durée standard" not in client.get(url).text
        # Règle « masquée pour tous » : l'administrateur ne la voit plus non plus.
        logged_client.post("/admin/parametres/affichage", data={"validee_display": "aucun"})
        assert "Durée standard" not in logged_client.get(url).text
    finally:
        logged_client.post("/admin/parametres/affichage", data={"validee_display": "rh"})
        with SessionLocal() as db:
            db.query(User).filter_by(username="lect_vd").delete()
            db.commit()


def test_shift_workers_filter(configured, pg, logged_client):
    """Filtre « Horaires » : agents en horaires postés (au moins un jour planifié sur la période) seulement, ou
    seulement les autres, dans le suivi, les rapports et leurs exports."""
    from openpyxl import load_workbook

    with pg.begin() as c:
        c.execute(text(f"INSERT INTO {SCHEMA}.pointage_postes (code, libelle, type, debut, fin) "
                       f"VALUES ('PF', 'Matin', 'travail', '06:00', '14:00') ON CONFLICT (code) DO NOTHING"))
        c.execute(text(f"INSERT INTO {SCHEMA}.pointage_planning (matricule, jour, code) VALUES "
                       f"('E004', :d, 'PF'), ('E005', :d, 'PF')"), {"d": MON})
    try:
        page = logged_client.get(f"/suivi?date={MON.isoformat()}&horaires=postes").text
        assert "E004</td>" in page and "E005</td>" in page and "E001</td>" not in page and "E002</td>" not in page
        assert "Horaires : Horaires postés (planning)" in page
        page = logged_client.get(f"/suivi?date={MON.isoformat()}&horaires=bureau").text
        assert "E001</td>" in page and "E004</td>" not in page and "E005</td>" not in page
        # Mardi : le planning ne couvre que le lundi → plus personne n'est « posté » sur cette journée.
        assert "E004</td>" not in logged_client.get(f"/suivi?date={TUE.isoformat()}&horaires=postes").text

        def report_matricules(horaires):
            url = f"/rapports/export.xlsx?p=perso&du={MON.isoformat()}&au={TUE.isoformat()}&horaires={horaires}"
            wb = load_workbook(io.BytesIO(logged_client.get(url).content))
            ws = wb["Employés"]
            filtres = {r[0].value: r[1].value for r in wb["Filtres"].iter_rows()}
            return {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}, filtres["Horaires"]

        assert report_matricules("postes") == ({"E004", "E005"}, "Horaires postés (planning)")
        bureau, label = report_matricules("bureau")
        assert {"E001", "E002", "E003"} <= bureau and not {"E004", "E005"} & bureau
        assert label == "Horaire de bureau (sans planning)"
        assert {"E001", "E004"} <= report_matricules("")[0]
        page = logged_client.get(f"/rapports?p=perso&du={MON.isoformat()}&au={TUE.isoformat()}&horaires=postes").text
        assert '<option value="postes" selected>' in page and "Fall Ibou" in page
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_planning WHERE code = 'PF'"))
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_postes WHERE code = 'PF'"))


def test_maternity_leave_and_breastfeeding(configured, pg, logged_client):
    """Congé de maternité (statut dédié, durée attribuée, même si l'employée badge) et allaitement : les heures
    d'allaitement (1 h, 2 h ou 3 h ; paramètre par défaut 1 h) complètent le temps réellement travaillé jusqu'à
    l'objectif (8h) et décalent d'autant le seuil de retard (heure prise en début de journée)."""
    from datetime import timedelta as td

    def day_rows(day=MON):
        with pg.connect() as c:
            return {r.matricule: r for r in c.execute(text(
                f"SELECT matricule, statut, duree_validee, duree_effective, allaitement, retard_min "
                f"FROM {SCHEMA}.f_pointage_journalier(:d, :d)"), {"d": day})}

    before = day_rows()
    assert before["E001"].duree_validee == td(hours=7, minutes=30) and before["E005"].statut == "ABSENT"
    try:
        assert "Maternité et allaitement" in logged_client.get("/admin/maternite").text
        # Congé de maternité : Ba Khady (absente lundi) et Ndiaye Moussa (a badgé lundi) ; saisie via la liste « matricule — nom ».
        for mat in ("E005 — Ba Khady", "E002"):
            r = logged_client.post("/admin/maternite", data={"matricule": mat, "type": "maternite", "du": MON.isoformat(),
                                                             "au": TUE.isoformat(), "commentaire": "certificat"},
                                   follow_redirects=True)
            assert "enregistré(e) du 21/09/2026 au 22/09/2026" in r.text, r.text[:2000]
        # Allaitement : Diallo Awa 1 h (par défaut), Sow Fatou 2 h, Diop Aminata 1 h.
        for mat, hours in (("E001", ""), ("E003", "02:00"), ("E007", "01:00")):
            r = logged_client.post("/admin/maternite", data={"matricule": mat, "type": "allaitement", "du": MON.isoformat(),
                                                             "au": MON.isoformat(), "heures": hours}, follow_redirects=True)
            assert "Période d&#39;allaitement" in r.text and "enregistré(e)" in r.text
        # Contrôles de saisie.
        for data, message in (
                ({"matricule": "E005", "type": "maternite", "du": TUE.isoformat(), "au": WED.isoformat()}, "déjà enregistré"),
                ({"matricule": "E007", "type": "maternite", "du": TUE.isoformat(), "au": MON.isoformat()}, "doit suivre"),
                ({"matricule": "X999", "type": "maternite", "du": MON.isoformat(), "au": MON.isoformat()}, "Aucun employé"),
                ({"matricule": "E004", "type": "allaitement", "du": MON.isoformat(), "au": MON.isoformat(),
                  "heures": "04:00"}, "1 h, 2 h ou 3 h")):
            assert message in logged_client.post("/admin/maternite", data=data, follow_redirects=True).text.replace("&#39;", "'")
        with pytest.raises(pointage.PointageError, match="01:00, 02:00 ou 03:00"):
            pointage.validate_param("heures_allaitement", "04:00")

        rows = day_rows()
        assert rows["E005"].statut == "CONGE_MATERNITE" and rows["E005"].duree_validee == td(hours=8)
        assert rows["E002"].statut == "CONGE_MATERNITE"  # même si elle a badgé
        assert day_rows(WED)["E005"].statut == "NON_OUVRE"  # mercredi férié : pas un jour de congé
        # Temps réellement travaillé + heures d'allaitement, plafonné à l'objectif : Diallo 8h20 réelles + 1 h → 8h ;
        # Sow 6h45 réelles (7h35 → 15h50, pause déduite) + 2 h → 8h ; durée effective inchangée.
        assert (rows["E001"].duree_validee, rows["E001"].allaitement) == (td(hours=8), td(hours=1))
        assert rows["E001"].duree_effective == before["E001"].duree_effective
        assert (rows["E003"].duree_validee, rows["E003"].allaitement) == (td(hours=8), td(hours=2))
        # Seuil de retard décalé : Diop arrive à 7h45 (= seuil, donc « En retard » d'habitude) → à l'heure avec 1 h
        # d'allaitement ; 4h15 travaillées (7h45 → 12h) + 1 h = 5h15 validées.
        assert (before["E007"].statut, before["E007"].retard_min) == ("RETARD", 15)
        assert (rows["E007"].statut, rows["E007"].retard_min) == ("A_L_HEURE", None)
        assert rows["E007"].duree_validee == td(hours=5, minutes=15)
        assert rows["E004"].allaitement is None and rows["E004"].duree_validee == before["E004"].duree_validee

        # 1 h seulement pour Sow Fatou : 6h45 réelles + 1 h = 7h45 (pas d'arrondi à 8h).
        with pg.begin() as c:
            c.execute(text(f"UPDATE {SCHEMA}.pointage_maternite SET heures = interval '1 hour' WHERE matricule = 'E003'"))
        assert day_rows()["E003"].duree_validee == td(hours=7, minutes=45)

        page = logged_client.get(f"/suivi?date={MON.isoformat()}").text
        assert "Congé maternité" in page and "Maternité <b>2</b>" in page and "Allaitement 1h00" in page
        # Rapports : le congé de maternité est un congé (ni absence ni journée disponible).
        from app import rapports
        with SessionLocal() as db:
            m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
        kpi = rapports.build(pg, m, pointage.Filters(du=MON, au=MON), "nom")["kpi"]
        assert kpi["conges"] >= 2 and kpi["absences"] == before_absences(before) - 1

        # Suppression : retour au calcul normal.
        page = logged_client.get("/admin/maternite").text
        with pg.connect() as c:
            ids = [i for (i,) in c.execute(text(f"SELECT id FROM {SCHEMA}.pointage_maternite"))]
        for i in ids:
            logged_client.post(f"/admin/maternite/{i}/delete")
        assert day_rows()["E005"].statut == "ABSENT" and day_rows()["E001"].duree_validee == td(hours=7, minutes=30)
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_maternite"))


def test_maternity_requests_workflow(configured, pg, logged_client):
    """Demandes de maternité et d'allaitement faites par l'employée : circuit paramétré (N+1 puis RH), effet sur le
    suivi seulement une fois validée, chevauchement, refus motivé, annulation (RH), saisie RH validée d'office,
    suppression de la période par les RH → demande annulée."""
    from datetime import timedelta as td

    from fastapi.testclient import TestClient

    from app import maternite
    from app.main import app
    from app.models import MaternityRequest, MaternityRequestAction, MaternitySettings, User

    def allaitement(mat="E003", day=MON):
        with pg.connect() as c:
            return c.execute(text(f"SELECT allaitement FROM {SCHEMA}.f_pointage_journalier(:d, :d) WHERE matricule = :m"),
                             {"d": day, "m": mat}).scalar()

    for data in ({"username": "mat_fatou", "role": "lecteur", "emp_matricule": "E003"},
                 {"username": "mat_awa", "role": "manager", "emp_matricule": "E001"},
                 {"username": "mat_rh", "role": "lecteur", "sick_leave_hr": "true"},
                 {"username": "mat_moussa", "role": "lecteur", "emp_matricule": "E002"}):
        r = logged_client.post("/admin/users/save", data={**data, "password": "motdepasse1"}, follow_redirects=True)
        assert "enregistré" in r.text, r.text[:500]

    def client(name):
        c = TestClient(app)
        c.post("/login", data={"username": name, "password": "motdepasse1"})
        return c

    try:
        # Circuit défini dans les paramètres (Ressources humaines → Maternité : périodes et circuit).
        r = logged_client.post("/admin/maternite/circuit", data={"type": ["responsable", "rh"], "user": ["", ""]},
                               follow_redirects=True)
        assert "Circuit enregistré : Responsable N+1 de l&#39;employé → Agent RH habilité" in r.text
        fatou, awa, rh, moussa = client("mat_fatou"), client("mat_awa"), client("mat_rh"), client("mat_moussa")
        page = fatou.get("/maternite").text
        assert "Faire une demande" in page and "Responsable N+1 de l&#39;employé → Agent RH habilité" in page

        # Demande d'allaitement (2 h) : pour elle-même uniquement, sans effet tant qu'elle n'est pas validée.
        r = fatou.post("/maternite", data={"type": "allaitement", "du": MON.isoformat(), "au": MON.isoformat(),
                                           "heures": "02:00", "matricule": "E005"}, follow_redirects=True)
        assert "transmise pour validation" in r.text and "Sow Fatou" in r.text
        with SessionLocal() as db:
            req = db.query(MaternityRequest).one()
            assert (req.matricule, req.manager_matricule, req.status, req.hours) == ("E003", "E001", "en_attente", "02:00")
            rid = req.id
        assert allaitement() is None
        assert "Valider" not in fatou.get(f"/maternite/{rid}").text  # pas de validation de sa propre demande
        assert "Sow Fatou" not in moussa.get("/maternite").text  # ni visible d'un autre employé
        assert moussa.get(f"/maternite/{rid}", follow_redirects=True).text.count("Demande introuvable") == 1

        # Étape 1 : le responsable N+1 ; étape 2 : les RH → période appliquée au suivi.
        page = awa.get("/maternite").text
        assert "À valider" in page and 'title="1 demande(s) à valider"' in page
        assert "passe à l&#39;étape suivante" in awa.post(f"/maternite/{rid}/decision", data={"decision": "valider"},
                                                           follow_redirects=True).text
        assert allaitement() is None
        r = rh.post(f"/maternite/{rid}/decision", data={"decision": "valider", "commentaire": "OK"}, follow_redirects=True)
        assert "Demande validée" in r.text and "Appliquée au suivi" in r.text
        assert allaitement() == td(hours=2)
        with SessionLocal() as db:
            pg_id = db.get(MaternityRequest, rid).pg_id
        assert pg_id and f"Demande n°{rid}" in logged_client.get("/admin/maternite").text

        # Chevauchement ; demande de congé de maternité avec certificat, refus motivé obligatoire.
        r = fatou.post("/maternite", data={"type": "allaitement", "du": MON.isoformat(), "au": TUE.isoformat()},
                       follow_redirects=True)
        assert "couvre déjà" in r.text
        r = fatou.post("/maternite", data={"type": "maternite", "du": "2026-10-05", "au": "2026-12-31"},
                       files={"certificat": ("certificat.png", b"\x89PNG....")}, follow_redirects=True)
        assert "transmise pour validation" in r.text
        with SessionLocal() as db:
            second = db.query(MaternityRequest).filter(MaternityRequest.id != rid).one().id
        assert fatou.get(f"/maternite/{second}/certificat").content == b"\x89PNG...."
        assert moussa.get(f"/maternite/{second}/certificat").status_code == 404
        assert "motif du refus" in awa.post(f"/maternite/{second}/decision", data={"decision": "refuser"},
                                             follow_redirects=True).text
        r = awa.post(f"/maternite/{second}/decision", data={"decision": "refuser", "commentaire": "Dates à revoir"},
                     follow_redirects=True)
        assert "Demande refusée" in r.text
        with SessionLocal() as db:
            assert db.get(MaternityRequest, second).status == "refuse"
            subject, _, text_body = maternite.build_notification(db.get(MaternityRequest, second), "refuse", "http://p", "",
                                                                 "Dates à revoir")
        assert subject.startswith("Congé de maternité refusé(e)") and "Dates à revoir" in text_body

        # Annulation d'une demande validée : réservée aux RH, la période est retirée du suivi.
        assert "que par les RH" in fatou.post(f"/maternite/{rid}/annuler", follow_redirects=True).text
        assert "Demande annulée" in rh.post(f"/maternite/{rid}/annuler", data={"commentaire": "erreur"},
                                            follow_redirects=True).text
        assert allaitement() is None

        # Saisie par un agent RH : validée d'office ; suppression de la période par les RH → demande annulée.
        r = rh.post("/maternite", data={"type": "allaitement", "du": MON.isoformat(), "au": MON.isoformat(),
                                        "matricule": "E007"}, follow_redirects=True)
        assert "enregistré(e) et validé(e)" in r.text and allaitement("E007") == td(hours=1)
        with SessionLocal() as db:
            third = db.query(MaternityRequest).filter(MaternityRequest.matricule == "E007").one()
            assert (third.status, third.source) == ("valide", "rh")
            third_id, third_pg = third.id, third.pg_id
        r = logged_client.post(f"/admin/maternite/{third_pg}/delete", follow_redirects=True)
        assert f"La demande n°{third_id} d&#39;origine est annulée" in r.text
        with SessionLocal() as db:
            assert db.get(MaternityRequest, third_id).status == "annule"
        assert allaitement("E007") is None
    finally:
        with pg.begin() as c:
            c.execute(text(f"DELETE FROM {SCHEMA}.pointage_maternite"))
        with SessionLocal() as db:
            db.query(MaternityRequestAction).delete()
            db.query(MaternityRequest).delete()
            db.query(MaternitySettings).delete()
            db.query(User).filter(User.username.in_(["mat_fatou", "mat_awa", "mat_rh", "mat_moussa"])).delete()
            db.commit()


def before_absences(rows) -> int:
    return sum(1 for r in rows.values() if r.statut == "ABSENT")


def test_pending_leave_reminders(configured, pg, logged_client, monkeypatch):
    """Relances des demandes de congé en attente (table HR Smart) : prochain validateur par son matricule, un mail par
    validateur regroupant ses demandes, jours déjà comptés absents, copie RH, espacement, délai de grâce, suivi."""
    from app import mails, validations
    from app.models import ApprovalReminderItem, ApprovalReminderLog, MailSettings

    with pg.begin() as c:
        c.execute(text(f'ALTER TABLE {SCHEMA}."Personnel" ADD COLUMN IF NOT EXISTS "Email" text'))
        c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "Email" = lower("Prenom" || \'.\' || "Nom") || \'@exemple.com\''))
        c.execute(text(f'DROP TABLE IF EXISTS {SCHEMA}."DemandeConge"'))
        c.execute(text(f"""CREATE TABLE {SCHEMA}."DemandeConge" ("Matricule" text, "DateDebut" timestamp, "DateFin" timestamp,
            "Etat" text, "Document" text, "ProchainvalidateurMatricule" text, "EtapeValidation" text, "DateDemande" date)"""))
        c.execute(text(f"""INSERT INTO {SCHEMA}."DemandeConge" VALUES
            ('E005', '2026-09-21', '2026-09-22', 'En Validation', 'CONGE', 'E001', 'N+1', '2026-09-18'),
            ('E003', '2026-09-22', '2026-09-22', 'Validé Par N1', 'CONGE', 'E001', 'RH', '2026-09-18'),
            ('E001', '2026-09-21', '2026-09-21', ' en validation ', 'CONGE', 'E003', 'N+1', '2026-09-18'),
            ('E004', '2026-09-25', '2026-09-25', 'En Validation', 'CONGE', 'X999', 'N+1', '2026-09-18'),
            ('E007', '2026-09-22', '2026-09-22', 'Approuvée', 'CONGE', NULL, NULL, '2026-09-18'),
            ('E002', '2026-09-22', '2026-09-22', 'Refusée', 'CONGE', 'E001', NULL, '2026-09-18')"""))
    with SessionLocal() as db:
        cfg = db.query(PointageConfig).one()
        original, conn_id = cfg.data, cfg.conn_id
    data = {**pointage.Mapping.from_json(original).__dict__, "email_col": "Email", "email_in": "person", "conn_id": conn_id,
            "leave_table": "DemandeConge", "leave_emp_col": "Matricule", "leave_ref": "matricule",
            "leave_start_col": "DateDebut", "leave_end_col": "DateFin", "leave_state_col": "Etat",
            "leave_state_values": "Approuvée", "leave_type_col": "Document"}
    delivered = []

    class FakeSMTP:
        def __init__(self, host, port, timeout=None): pass
        def ehlo(self): pass
        def starttls(self, context=None): pass
        def login(self, user, pwd): pass
        def send_message(self, msg): delivered.append(msg)
        def quit(self): pass

    monkeypatch.setattr(mails.smtplib, "SMTP", FakeSMTP)
    settings = {"enabled": "true", "send_time": "08h30", "weekdays": ["1", "2", "3", "4", "5"], "app_name": "HR Smart",
                "app_url": "https://ico-hrsmart.ea.holcim.net/HRSmartV3", "leave_enabled": "true",
                "leave_pending": "En Validation, Validé Par N1", "leave_validator_col": "ProchainvalidateurMatricule",
                "leave_step_col": "EtapeValidation", "leave_created_col": "DateDemande", "repeat_days": "2",
                "horizon_days": "14", "past_days": "31", "cc_addresses": "rh@exemple.com", "cc_after_absent": "2",
                "show_in_suivi": "true"}
    TUE9 = datetime(2026, 9, 22, 9, 0)
    try:
        assert "installées" in logged_client.post("/admin/pointage", data=data, follow_redirects=True).text
        with SessionLocal() as db:
            ms = mails.get_settings(db)
            ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "smtp.test", "pointage@exemple.com", "production", ""
            db.commit()
        r = logged_client.post("/admin/validations/parametres", data={**settings, "leave_pending": ""}, follow_redirects=True)
        assert "indiquez au moins un statut" in r.text
        r = logged_client.post("/admin/validations/parametres", data=settings, follow_redirects=True)
        assert "Paramètres des relances enregistrés" in r.text

        # Demandes en attente : statuts paramétrés (casse et espaces ignorés), validateur retrouvé par son matricule.
        with SessionLocal() as db:
            s, m = validations.get_settings(db), pointage.Mapping.from_json(db.query(PointageConfig).one().data)
        reqs = {r["matricule"]: r for r in validations.pending_requests(pg, m, s, MON - timedelta(days=31),
                                                                        MON + timedelta(days=14), today=TUE)}
        assert set(reqs) == {"E005", "E003", "E001", "E004"}  # approuvée et refusée ignorées
        assert (reqs["E005"]["validator"]["email"], reqs["E005"]["absents"], reqs["E005"]["jours"]) == (
            "awa.diallo@exemple.com", 2, 2)
        assert reqs["E001"]["absents"] == 0 and reqs["E001"]["validator"]["nom"] == "Sow Fatou"
        assert reqs["E004"]["validator"]["found"] is False and reqs["E005"]["etape"] == "N+1"
        assert reqs["E005"]["attente"] == 4

        # Envoi : avant l'heure, rien ; à 9h, un mail par validateur.
        assert validations.run_due(datetime(2026, 9, 22, 8, 0)) == {"status": "idle"}
        assert validations.run_due(TUE9) == {"status": "done", "sent": 3, "skipped": 0, "failed": 0}
        by_to = {msg["To"]: msg for msg in delivered}
        awa = by_to["awa.diallo@exemple.com"]
        assert awa["Subject"] == "2 demandes de congé en attente de votre validation"
        assert awa["Cc"] == "rh@exemple.com"  # 2 jours déjà comptés absents ≥ seuil de copie
        body = awa.get_body(("html",)).get_content()
        assert "Ba Khady" in body and "Sow Fatou" in body and "Valider dans HR Smart" in body
        assert "https://ico-hrsmart.ea.holcim.net/HRSmartV3" in body and "3 jour(s) déjà comptés absents" in body
        assert "N+1" in body and "en attente depuis 4 j" in body and "Bonjour Awa,</h1>" in body
        assert by_to["fatou.sow@exemple.com"]["Cc"] is None  # demande sans jour compté absent : pas de copie
        assert "Fall Ibou" in by_to["rh@exemple.com"].get_body(("html",)).get_content()  # validateur inconnu → RH

        # Un seul envoi par jour ; mercredi : espacement de 2 jours non atteint ; jeudi : nouvelle relance.
        assert validations.run_due(datetime(2026, 9, 22, 11, 0)) == {"status": "idle"}
        assert validations.run_due(datetime(2026, 9, 23, 9, 0)) == {"status": "idle"}
        assert validations.run_due(datetime(2026, 9, 24, 9, 0))["sent"] == 3
        with SessionLocal() as db:
            assert {i.count for i in db.query(ApprovalReminderItem)} == {2}
            db.query(ApprovalReminderItem).delete()
            db.query(ApprovalReminderLog).delete()
            s = validations.get_settings(db)
            s.grace_days = 10  # demandes faites le 18/09 : pas avant le 28/09
            db.commit()
        assert validations.run_due(TUE9) == {"status": "idle"}
        with SessionLocal() as db:
            validations.get_settings(db).grace_days = 0
            db.commit()

        # Administration : situation par validateur, aperçu ; suivi : absence signalée « en attente de validation ».
        monkeypatch.setattr(validations, "local_now", lambda: TUE9)  # « aujourd'hui » = date des données de test
        page = logged_client.get("/admin/validations").text
        assert "Demandes en attente" in page and "matricule introuvable" in page and "ProchainvalidateurMatricule" in page
        assert "Demandes à valider" in logged_client.get("/admin/validations/apercu?validateur=E001").text
        page = logged_client.get(f"/suivi?date={TUE.isoformat()}").text
        assert "Congé (CONGE) en attente de validation (En Validation)" in page
        assert "Congé (CONGE) en attente de validation (Validé Par N1)" in page
    finally:
        with SessionLocal() as db:
            db.query(ApprovalReminderItem).delete()
            db.query(ApprovalReminderLog).delete()
            validations.get_settings(db).enabled = False
            ms = db.query(MailSettings).first()
            ms.smtp_host, ms.from_email, ms.mode, ms.test_recipients = "", "", "test", ""
            db.commit()
        logged_client.post("/admin/pointage", data={**pointage.Mapping.from_json(original).__dict__, "conn_id": conn_id})
        with pg.begin() as c:
            c.execute(text(f'DROP TABLE IF EXISTS {SCHEMA}."DemandeConge"'))


def test_hors_liste_list(configured, pg, logged_client):
    """Liste des badges hors liste dans le suivi : badge, jours, pointages, premier et dernier pointage, fiche proche."""
    with SessionLocal() as db:
        m = pointage.Mapping.from_json(db.query(PointageConfig).one().data)
    rows = pointage.hors_liste(pg, m, MON, TUE)
    assert [(r["emp_key"], r["jours"], r["pointages"], r["suggestion_matricule"]) for r in rows] == [("99", 1, 2, None)]
    assert rows[0]["premier"] == at(MON, "09:00") and rows[0]["dernier"] == at(MON, "17:00")
    page = logged_client.get(f"/suivi?date={MON.isoformat()}").text
    assert "1 badge(s) hors liste" in page and '<td class="mono">99</td>' in page and "aucune fiche proche" in page
    assert 'id="hors-liste" open' not in page
    page = logged_client.get(f"/suivi?date={MON.isoformat()}&pop=hors").text
    assert 'id="hors-liste" open' in page
    assert "badge(s) hors liste" not in logged_client.get(f"/suivi?date={TUE.isoformat()}").text  # personne mardi


def test_default_staff_categories(configured, pg, logged_client):
    """Statuts du personnel cochés par défaut (Cadre et Non Cadre) dans le suivi, les rapports et l'export ; tout
    décocher = tout le personnel ; réglage dans Paramètres horaires ; sans effet sur les badges hors liste."""
    from openpyxl import load_workbook

    with pg.begin() as c:
        c.execute(text(f"INSERT INTO {SCHEMA}.statuts VALUES (3, 'Stagiaire') ON CONFLICT DO NOTHING"))
        c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "IDStatut" = 3 WHERE "IDPersonnel" = 107'))  # Diop Aminata
    url = f"/suivi?date={MON.isoformat()}"
    try:
        page = logged_client.get(url).text
        assert "E001</td>" in page and "E007</td>" not in page  # stagiaire masquée par défaut
        assert 'name="categorie_set" value="1"' in page and "Statut du personnel : Cadre, Non cadre" in page
        page = logged_client.get(url + "&categorie_set=1").text  # tout décoché : tout le personnel
        assert "E007</td>" in page and "Statut du personnel :" not in page
        assert "E007</td>" in logged_client.get(url + "&categorie=Stagiaire").text
        assert "99</td>" in logged_client.get(url + "&pop=hors").text  # hors liste : pas de statut, pas de filtre
        ws = load_workbook(io.BytesIO(logged_client.get(f"/suivi/export.xlsx?date={MON.isoformat()}").content))["Suivi journalier"]
        assert "E007" not in {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
        rap = logged_client.get(f"/rapports/export.xlsx?p=perso&du={MON.isoformat()}&au={MON.isoformat()}")
        assert "E007" not in {c.value for c in load_workbook(io.BytesIO(rap.content))["Employés"]["A"]}

        # Réglage : plus aucun statut par défaut → tout le personnel.
        page = logged_client.get("/admin/parametres").text
        assert "Statuts du personnel cochés par défaut" in page and 'value="Stagiaire" >' in page.replace("  ", " ")
        r = logged_client.post("/admin/parametres/affichage", data={"validee_display": "rh", "categories_form": "1"},
                               follow_redirects=True)
        assert "statuts du personnel par défaut : tous" in r.text
        assert "E007</td>" in logged_client.get(url).text
        logged_client.post("/admin/parametres/affichage", data={"validee_display": "rh", "categories_form": "1",
                                                                 "default_categories": ["Stagiaire"]})
        page = logged_client.get(url).text
        assert "E007</td>" in page and "E001</td>" not in page
    finally:
        logged_client.post("/admin/parametres/affichage", data={"validee_display": "rh", "categories_form": "1",
                                                                 "default_categories": ["Cadre", "Non cadre"]})
        with pg.begin() as c:
            c.execute(text(f'UPDATE {SCHEMA}."Personnel" SET "IDStatut" = 2 WHERE "IDPersonnel" = 107'))
            c.execute(text(f"DELETE FROM {SCHEMA}.statuts WHERE id = 3"))
